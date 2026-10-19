#!/usr/bin/env python3
"""Convierte un Markdown simple en un Word con la marca del proyecto.

Uso:
    python herramientas/md_a_docx.py entrada.md salida.docx --proyecto proyectos/<slug>

Soporta: títulos (#, ##, ###), párrafos, viñetas (- o *), listas numeradas (1.),
tablas Markdown, **negrita** y *cursiva*. En tablas, las celdas que dicen Verde,
Amarillo o Rojo se colorean como semáforo. Toma logo, colores, fuente, tamaño de
página y pie de página desde proyecto.yaml (sección formato).

Autoría limpia: normaliza caracteres invisibles, omite comentarios HTML, deja como autor a
autoria.nombre (o responsable.nombre) con fecha actual, elimina partes heredadas de la plantilla y usa el
modo de compatibilidad actual de Word. Si el texto tiene marcadores [PENDIENTE…] no genera el documento,
salvo con --borrador (uso interno; verificar_autoria.py lo marcará igual).
"""
import argparse
import re
import sys
from pathlib import Path

import _config  # noqa: E402  (proyecto.yaml con herencia de programa)

try:
    import yaml
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("Faltan dependencias: pip install python-docx pyyaml")

import autoria  # noqa: E402  (módulo del harness en herramientas/)

SEMAFORO = {"verde": "C6EFCE", "amarillo": "FFEB9C", "rojo": "FFC7CE"}
PAGINAS = {"carta": (Cm(21.59), Cm(27.94)), "a4": (Cm(21.0), Cm(29.7))}


def cargar_cfg(ruta_proyecto):
    if not ruta_proyecto:
        return {}, None
    raiz = Path(ruta_proyecto)
    if not (raiz / "proyecto.yaml").exists():
        return {}, raiz
    return _config.cargar(raiz), raiz


def rgb(hexa, defecto="1F3864"):
    h = str(hexa or defecto).lstrip("#")
    if not re.fullmatch(r"[0-9A-Fa-f]{6}", h):
        h = defecto
    return RGBColor.from_string(h.upper()), h.upper()


def sombrear(celda, hexa):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    tcPr.append(shd)


def agregar_inline(parrafo, texto, color=None, negrita_base=False):
    partes = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            run = parrafo.add_run(parte[1:-1])
            run.italic = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = parrafo.add_run(parte[1:-1])
        else:
            run = parrafo.add_run(parte)
        if negrita_base:
            run.bold = True
        if color is not None:
            run.font.color.rgb = color


def campo_pagina(parrafo):
    run = parrafo.add_run()
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def configurar(doc, cfg, raiz):
    fmt = cfg.get("formato", {}) or {}
    proy = cfg.get("proyecto", {}) or {}
    fuente = fmt.get("fuente", "Arial")
    color, _ = rgb(fmt.get("color_primario"))

    estilo = doc.styles["Normal"]
    estilo.font.name = fuente
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), fuente)
    estilo.font.size = Pt(10.5)
    for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
        st = doc.styles[f"Heading {nivel}"]
        st.font.name = fuente
        rfonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:eastAsia"), fuente)
        st.font.size = Pt(tam)
        st.font.bold = True
        st.font.color.rgb = color

    ancho, alto = PAGINAS.get(str(fmt.get("tamano_pagina", "carta")).lower(), PAGINAS["carta"])
    sec = doc.sections[0]
    sec.page_width, sec.page_height = ancho, alto
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)

    # Encabezado: logo a la izquierda, nombre del proyecto a la derecha
    enc = sec.header.paragraphs[0]
    logo = (raiz / fmt.get("logo", "marca/logo.png")) if raiz else None
    if logo and logo.exists():
        enc.add_run().add_picture(str(logo), height=Cm(1.2))
        enc.add_run("\t\t")
    nombre = proy.get("nombre", "")
    if cfg.get("_programa") and nombre and "PENDIENTE" not in str(nombre).upper():
        nombre = f"{cfg['_programa'].get('nombre', '')} · {nombre}"  # subproyecto: programa y subproyecto en el encabezado
    if nombre and "PENDIENTE" not in str(nombre).upper():
        r = enc.add_run(str(nombre))
        r.font.size = Pt(9)
        r.font.color.rgb = color

    # Pie: texto de confidencialidad + número de página
    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cliente = proy.get("cliente", "")
    texto_pie = fmt.get("pie_de_pagina", "Documento confidencial")
    if cliente and "PENDIENTE" not in str(cliente).upper():
        texto_pie = f"{texto_pie} - {cliente}"
    r = pie.add_run(f"{texto_pie}   |   Página ")
    r.font.size = Pt(8)
    campo_pagina(pie)
    return color, fmt


def agregar_tabla(doc, filas, color_hex):
    encabezado = filas[0]
    cuerpo = filas[1:]
    tabla = doc.add_table(rows=1, cols=len(encabezado))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, texto in enumerate(encabezado):
        celda = tabla.rows[0].cells[j]
        celda.text = ""
        agregar_inline(celda.paragraphs[0], texto, color=RGBColor(0xFF, 0xFF, 0xFF), negrita_base=True)
        sombrear(celda, color_hex)
    for fila in cuerpo:
        celdas = tabla.add_row().cells
        for j in range(len(encabezado)):
            texto = fila[j] if j < len(fila) else ""
            celdas[j].text = ""
            agregar_inline(celdas[j].paragraphs[0], texto)
            clave = re.sub(r"[^a-záéíóú]", "", texto.lower())
            if clave in SEMAFORO:
                sombrear(celdas[j], SEMAFORO[clave])
    for fila in tabla.rows:
        for celda in fila.cells:
            for par in celda.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                for run in par.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def dividir_fila(linea):
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def convertir(md, doc, color, color_hex):
    lineas = md.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        if not linea.strip() or linea.strip().startswith("<!--"):
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", linea)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
            i += 1
            continue
        if linea.strip().startswith("|"):
            bloque = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", lineas[i].strip()):
                    bloque.append(dividir_fila(lineas[i]))
                i += 1
            if bloque:
                agregar_tabla(doc, bloque, color_hex)
            continue
        m = re.match(r"^\s*[-*]\s+(.*)", linea)
        if m:
            agregar_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", linea)
        if m:
            agregar_inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        if linea.strip().startswith(">"):
            p = doc.add_paragraph()
            agregar_inline(p, linea.strip().lstrip(">").strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        # párrafo: junta líneas seguidas
        texto = [linea.strip()]
        i += 1
        while i < len(lineas) and lineas[i].strip() and not re.match(r"^(#|\||\s*[-*]\s|\s*\d+[.)]\s|>)", lineas[i]):
            texto.append(lineas[i].strip())
            i += 1
        agregar_inline(doc.add_paragraph(), " ".join(texto))


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("entrada")
    p.add_argument("salida")
    p.add_argument("--proyecto", help="carpeta del proyecto (para marca, formato y autoría)")
    p.add_argument("--borrador", action="store_true", help="permitir marcadores [PENDIENTE] (documento interno)")
    args = p.parse_args()

    md = Path(args.entrada).read_text(encoding="utf-8")
    md = autoria.normalizar_texto(re.sub(r"<!--.*?-->", "", md, flags=re.S))
    pendientes = [n for n, linea in enumerate(md.splitlines(), 1) if "[PENDIENTE" in linea.upper()]
    if pendientes and not args.borrador:
        sys.exit(f"El texto tiene {len(pendientes)} marcador(es) [PENDIENTE] (líneas {', '.join(map(str, pendientes[:10]))}). "
                 "Complétalos antes de emitir, o usa --borrador para un documento interno.")

    cfg, raiz = cargar_cfg(args.proyecto)
    doc = Document()
    color, fmt = configurar(doc, cfg, raiz)
    _, color_hex = rgb(fmt.get("color_primario"))
    convertir(md, doc, color, color_hex)
    Path(args.salida).parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.salida)
    aut = autoria.cargar_autoria(raiz)
    titulo = next((m.group(1).strip() for m in re.finditer(r"^#\s+(.+)$", md, re.M)), "")
    autoria.limpiar_ooxml(args.salida, aut.get("nombre", ""), titulo, aut.get("organizacion", ""))
    print(f"Documento generado: {args.salida}")


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    main()

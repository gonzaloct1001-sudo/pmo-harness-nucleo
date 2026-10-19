#!/usr/bin/env python3
"""Genera un modelo de evaluación económica en Excel (con fórmulas) a partir de un archivo de supuestos.

Uso:
    python herramientas/modelo_economico.py <supuestos.yaml> [--salida modelo.xlsx] [--proyecto proyectos/<slug>]

Hojas: Supuestos (celdas de entrada en azul), Flujo (flujo de caja anual con fórmulas), Indicadores (VAN,
TIR, período de recuperación, VAN/inversión) y Sensibilidad (VAN ante variaciones de ingresos, costos,
inversión y tasa). Todo se calcula con fórmulas que referencian la hoja Supuestos: al cambiar un supuesto en
Excel, el modelo se recalcula. El script también calcula los indicadores en Python y los imprime como control.

Convenciones del modelo (quedan escritas en la hoja Supuestos):
- Año 0 = inversión inicial; flujos al final de cada año.
- Impuesto = tasa × utilidad antes de impuesto cuando es positiva (sin arrastre de pérdidas: simplificación).
- Capital de trabajo requerido en el año t = % × ingresos del año t+1; se recupera al final del horizonte.
- Valor residual al final del horizonte: valor libro (inversión − depreciación acumulada), un monto dado, o ninguno.
- Evaluación social: usar tasa social y precios sociales según la metodología oficial aplicable (tasa de impuesto 0).
"""
import argparse
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import autoria  # noqa: E402  (módulo del harness en herramientas/)

AZUL = Font(name="Arial", color="0000FF", size=10)
NORMAL = Font(name="Arial", size=10)
NEGRITA = Font(name="Arial", size=10, bold=True)
TITULO = Font(name="Arial", size=13, bold=True, color="1F3864")
ENC = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONDO_ENC = PatternFill("solid", fgColor="1F3864")
FONDO_TOTAL = PatternFill("solid", fgColor="D9E1F2")
MONEDA = '#,##0;(#,##0);"-"'
PCT = "0.0%"


def calcular(s, fi=0.0, fc=0.0, finv=0.0, tasa=None):
    """Réplica en Python del modelo (control y sensibilidad)."""
    n = int(s["horizonte_anios"])
    t_imp = float(s.get("tasa_impuesto", 0))
    r = float(s["tasa_descuento"] if tasa is None else tasa)
    ing = [sum(x["monto_anual"] * (1 + x.get("crecimiento_anual", 0)) ** (t - x.get("anio_inicio", 1))
               for x in s.get("ingresos", []) if t >= x.get("anio_inicio", 1)) * (1 + fi) for t in range(n + 1)]
    cos = [sum(x["monto_anual"] * (1 + x.get("crecimiento_anual", 0)) ** (t - x.get("anio_inicio", 1))
               for x in s.get("costos", []) if t >= x.get("anio_inicio", 1)) * (1 + fc) for t in range(n + 1)]
    dep = [sum(d["monto"] / d["vida_util"] for d in s.get("depreciacion", []) if 1 <= t <= d["vida_util"]) * (1 + finv)
           for t in range(n + 1)]
    inv = [sum(x["monto"] for x in s.get("inversiones", []) if x.get("anio", 0) == t) * (1 + finv) for t in range(n + 1)]
    pct = float((s.get("capital_trabajo") or {}).get("porcentaje_ingresos", 0))
    ct_req = [pct * ing[t + 1] if t < n else 0 for t in range(n + 1)]
    dct = [ct_req[t] - (ct_req[t - 1] if t else 0) for t in range(n + 1)]
    vr_cfg = s.get("valor_residual") or {"metodo": "ninguno"}
    if vr_cfg.get("metodo") == "valor_libro":
        vr_final = sum(inv) - sum(dep[1:])
    elif vr_cfg.get("metodo") == "monto":
        vr_final = float(vr_cfg.get("monto", 0))
    else:
        vr_final = 0.0
    flujo = []
    for t in range(n + 1):
        uai = ing[t] - cos[t] - dep[t]
        imp = max(0.0, uai) * t_imp
        flujo.append(uai - imp + dep[t] - inv[t] - dct[t] + (vr_final if t == n else 0))
    van = sum(f / (1 + r) ** t for t, f in enumerate(flujo))
    return flujo, van


def tir(flujo):
    lo, hi = -0.99, 10.0
    f = lambda r: sum(x / (1 + r) ** t for t, x in enumerate(flujo))
    if f(lo) * f(hi) > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        if f(lo) * f(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("supuestos")
    ap.add_argument("--salida")
    ap.add_argument("--proyecto", help="carpeta del proyecto (autoría del archivo); por defecto se busca hacia arriba")
    a = ap.parse_args()
    ruta = Path(a.supuestos)
    s = yaml.safe_load(ruta.read_text(encoding="utf-8"))
    n = int(s["horizonte_anios"])
    salida = Path(a.salida) if a.salida else ruta.with_name("modelo-economico.xlsx")

    wb = Workbook()
    # ------------------------- Supuestos -------------------------
    ws = wb.active
    ws.title = "Supuestos"
    ws["A1"] = f"Supuestos — {s.get('nombre', 'Modelo económico')}"
    ws["A1"].font = TITULO
    ws["A2"] = "Celdas en azul: supuestos (editables). El resto del libro se calcula con fórmulas."
    ws["A2"].font = Font(name="Arial", italic=True, size=9)
    ref = {}
    fila = 4
    for clave, etiqueta, valor, fmt in (("tasa", "Tasa de descuento", s["tasa_descuento"], PCT),
                                        ("horizonte", "Horizonte (años)", n, "0"),
                                        ("impuesto", "Tasa de impuesto", s.get("tasa_impuesto", 0), PCT),
                                        ("ct", "Capital de trabajo (% ingresos del año siguiente)",
                                         (s.get("capital_trabajo") or {}).get("porcentaje_ingresos", 0), PCT)):
        ws.cell(fila, 1, etiqueta).font = NORMAL
        c = ws.cell(fila, 2, valor)
        c.font, c.number_format = AZUL, fmt
        ref[clave] = f"Supuestos!$B${fila}"
        fila += 1
    ws.cell(fila, 1, "Moneda").font = NORMAL
    ws.cell(fila, 2, s.get("moneda", "CLP")).font = AZUL
    fila += 1
    ws.cell(fila, 1, "Tipo de evaluación").font = NORMAL
    ws.cell(fila, 2, s.get("tipo_evaluacion", "privada")).font = AZUL
    fila += 2

    def tabla(titulo, columnas, filas_datos, campos, formatos):
        nonlocal fila
        ws.cell(fila, 1, titulo).font = NEGRITA
        fila += 1
        for j, c in enumerate(columnas, 1):
            celda = ws.cell(fila, j, c)
            celda.font, celda.fill = ENC, FONDO_ENC
        fila += 1
        refs = []
        for d in filas_datos:
            r = {}
            for j, (campo, fmt) in enumerate(zip(campos, formatos), 1):
                celda = ws.cell(fila, j, d.get(campo, 0 if j > 1 else ""))
                celda.font = AZUL if j > 1 else NORMAL
                if fmt:
                    celda.number_format = fmt
                r[campo] = f"Supuestos!${L(j)}${fila}"
            refs.append(r)
            fila += 1
        fila += 1
        return refs

    r_inv = tabla("Inversiones", ["Concepto", "Año", "Monto"], s.get("inversiones", []),
                  ["concepto", "anio", "monto"], [None, "0", MONEDA])
    r_ing = tabla("Ingresos", ["Concepto", "Año inicio", "Monto anual", "Crecimiento anual"], s.get("ingresos", []),
                  ["concepto", "anio_inicio", "monto_anual", "crecimiento_anual"], [None, "0", MONEDA, PCT])
    r_cos = tabla("Costos", ["Concepto", "Año inicio", "Monto anual", "Crecimiento anual"], s.get("costos", []),
                  ["concepto", "anio_inicio", "monto_anual", "crecimiento_anual"], [None, "0", MONEDA, PCT])
    r_dep = tabla("Depreciación", ["Activo", "Monto depreciable", "Vida útil (años)"], s.get("depreciacion", []),
                  ["activo", "monto", "vida_util"], [None, MONEDA, "0"])
    vr = s.get("valor_residual") or {"metodo": "ninguno"}
    ws.cell(fila, 1, "Valor residual: método (valor_libro | monto | ninguno)").font = NORMAL
    ws.cell(fila, 2, vr.get("metodo", "ninguno")).font = AZUL
    ref["vr_metodo"] = f"Supuestos!$B${fila}"
    fila += 1
    ws.cell(fila, 1, "Valor residual: monto (si el método es 'monto')").font = NORMAL
    c = ws.cell(fila, 2, float(vr.get("monto", 0)))
    c.font, c.number_format = AZUL, MONEDA
    ref["vr_monto"] = f"Supuestos!$B${fila}"
    fila += 2
    for nota in (__doc__.split("Convenciones del modelo (quedan escritas en la hoja Supuestos):")[1].strip().splitlines()):
        ws.cell(fila, 1, nota.strip()).font = Font(name="Arial", size=9, italic=True)
        fila += 1
    ws.column_dimensions["A"].width = 52
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    # ------------------------- Flujo -------------------------
    fl = wb.create_sheet("Flujo")
    fl["A1"] = "Flujo de caja"
    fl["A1"].font = TITULO
    fl.cell(3, 1, "Concepto").font = ENC
    fl.cell(3, 1).fill = FONDO_ENC
    for t in range(n + 1):
        c = fl.cell(3, 2 + t, t)
        c.font, c.fill, c.alignment = ENC, FONDO_ENC, Alignment(horizontal="center")
    col = lambda t: L(2 + t)
    ult = col(n)
    fila_f = [4]
    filas = {}

    def linea(etiqueta, formula_t, clave=None, negrita=False, total=False):
        f = fila_f[0]
        fl.cell(f, 1, etiqueta).font = NEGRITA if negrita else NORMAL
        for t in range(n + 1):
            c = fl.cell(f, 2 + t, formula_t(t, f))
            c.number_format = MONEDA
            c.font = NEGRITA if negrita else NORMAL
            if total:
                c.fill = FONDO_TOTAL
        if total:
            fl.cell(f, 1).fill = FONDO_TOTAL
        if clave:
            filas[clave] = f
        fila_f[0] += 1
        return f

    anio = lambda t: f"{col(t)}$3"
    ing_f = [linea(f"Ingreso: {d.get('concepto', '')}",
                   lambda t, f, r=r: f"=IF({anio(t)}>={r['anio_inicio']},{r['monto_anual']}*(1+{r['crecimiento_anual']})^({anio(t)}-{r['anio_inicio']}),0)")
             for d, r in zip(s.get("ingresos", []), r_ing)]
    linea("Total ingresos", lambda t, f: f"=SUM({col(t)}{ing_f[0]}:{col(t)}{ing_f[-1]})" if ing_f else "=0", "ing", True, True)
    cos_f = [linea(f"Costo: {d.get('concepto', '')}",
                   lambda t, f, r=r: f"=IF({anio(t)}>={r['anio_inicio']},{r['monto_anual']}*(1+{r['crecimiento_anual']})^({anio(t)}-{r['anio_inicio']}),0)")
             for d, r in zip(s.get("costos", []), r_cos)]
    linea("Total costos", lambda t, f: f"=SUM({col(t)}{cos_f[0]}:{col(t)}{cos_f[-1]})" if cos_f else "=0", "cos", True, True)
    dep_f = [linea(f"Depreciación: {d.get('activo', '')}",
                   lambda t, f, r=r: f"=IF(AND({anio(t)}>=1,{anio(t)}<={r['vida_util']}),{r['monto']}/{r['vida_util']},0)")
             for d, r in zip(s.get("depreciacion", []), r_dep)]
    linea("Total depreciación", lambda t, f: f"=SUM({col(t)}{dep_f[0]}:{col(t)}{dep_f[-1]})" if dep_f else "=0", "dep", True, True)
    linea("Utilidad antes de impuesto", lambda t, f: f"={col(t)}{filas['ing']}-{col(t)}{filas['cos']}-{col(t)}{filas['dep']}", "uai", True)
    linea("Impuesto", lambda t, f: f"=MAX(0,{col(t)}{filas['uai']})*{ref['impuesto']}", "imp")
    linea("Utilidad neta", lambda t, f: f"={col(t)}{filas['uai']}-{col(t)}{filas['imp']}", "un", True)
    linea("(+) Depreciación", lambda t, f: f"={col(t)}{filas['dep']}", "dep2")
    inv_f = [linea(f"Inversión: {d.get('concepto', '')}", lambda t, f, r=r: f"=IF({anio(t)}={r['anio']},{r['monto']},0)")
             for d, r in zip(s.get("inversiones", []), r_inv)]
    linea("(−) Total inversión", lambda t, f: f"=SUM({col(t)}{inv_f[0]}:{col(t)}{inv_f[-1]})" if inv_f else "=0", "inv", True, True)
    linea("Capital de trabajo requerido",
          lambda t, f: (f"={ref['ct']}*{col(t + 1)}{filas['ing']}" if t < n else "=0"), "ctr")
    linea("(−) Variación capital de trabajo",
          lambda t, f: (f"={col(t)}{filas['ctr']}" if t == 0 else f"={col(t)}{filas['ctr']}-{col(t - 1)}{filas['ctr']}"), "dct")
    linea("(+) Valor residual",
          lambda t, f: (f"=IF({ref['vr_metodo']}=\"valor_libro\",SUM($B${filas['inv']}:${ult}${filas['inv']})-SUM($C${filas['dep']}:${ult}${filas['dep']}),"
                        f"IF({ref['vr_metodo']}=\"monto\",{ref['vr_monto']},0))" if t == n else "=0"), "vr")
    linea("FLUJO NETO", lambda t, f: f"={col(t)}{filas['un']}+{col(t)}{filas['dep2']}-{col(t)}{filas['inv']}-{col(t)}{filas['dct']}+{col(t)}{filas['vr']}",
          "fn", True, True)
    linea("Flujo acumulado", lambda t, f: (f"={col(t)}{filas['fn']}" if t == 0 else f"={col(t - 1)}{f}+{col(t)}{filas['fn']}"), "acum")
    fl.column_dimensions["A"].width = 38
    for t in range(n + 1):
        fl.column_dimensions[col(t)].width = 14
    fl.freeze_panes = "B4"
    rango_fn = f"Flujo!$B${filas['fn']}:${ult}${filas['fn']}"
    rango_fn1 = f"Flujo!$C${filas['fn']}:${ult}${filas['fn']}"

    # ------------------------- Indicadores -------------------------
    ind = wb.create_sheet("Indicadores")
    ind["A1"] = "Indicadores"
    ind["A1"].font = TITULO
    items = [("VAN", f"=Flujo!$B${filas['fn']}+NPV({ref['tasa']},{rango_fn1})", MONEDA),
             ("TIR", f"=IFERROR(IRR({rango_fn}),\"No calculable\")", PCT),
             ("Período de recuperación (años, flujo acumulado ≥ 0)",
              f"=IF(COUNTIF(Flujo!$B${filas['acum']}:${ult}${filas['acum']},\"<0\")>{n},\"No recupera\","
              f"COUNTIF(Flujo!$B${filas['acum']}:${ult}${filas['acum']},\"<0\"))", "0"),
             ("Inversión total", f"=SUM(Flujo!$B${filas['inv']}:${ult}${filas['inv']})", MONEDA),
             ("VAN / inversión", "=IFERROR(B3/B6,0)", "0.00"),
             ("Tasa de descuento", f"={ref['tasa']}", PCT)]
    for i, (et, fo, fmt) in enumerate(items, 3):
        ind.cell(i, 1, et).font = NEGRITA
        c = ind.cell(i, 2, fo)
        c.number_format, c.font = fmt, NEGRITA
    ind.column_dimensions["A"].width = 50
    ind.column_dimensions["B"].width = 20

    # ------------------------- Sensibilidad -------------------------
    se = wb.create_sheet("Sensibilidad")
    se["A1"] = "Sensibilidad del VAN"
    se["A1"].font = TITULO
    se["A2"] = "Cada fila recalcula el flujo con los factores indicados (variación sobre el caso base)."
    se["A2"].font = Font(name="Arial", italic=True, size=9)
    cab = ["Escenario", "Δ ingresos", "Δ costos", "Δ inversión"] + [f"Año {t}" for t in range(n + 1)] + ["VAN"]
    for j, c in enumerate(cab, 1):
        celda = se.cell(4, j, c)
        celda.font, celda.fill = ENC, FONDO_ENC
    sens = s.get("sensibilidad") or {}
    variaciones = sens.get("variaciones", [-0.2, -0.1, 0.1, 0.2])
    variables = sens.get("variables", ["ingresos", "costos", "inversion"])
    escenarios = [("Base", 0, 0, 0)]
    for v in variables:
        for x in variaciones:
            escenarios.append((f"{v} {x:+.0%}", x if v == "ingresos" else 0, x if v == "costos" else 0, x if v == "inversion" else 0))
    F = lambda clave, t: f"Flujo!{col(t)}${filas[clave]}"
    fila_s = 5
    for nombre, fi, fc, finv in escenarios:
        se.cell(fila_s, 1, nombre).font = NORMAL
        for j, v in enumerate((fi, fc, finv), 2):
            c = se.cell(fila_s, j, v)
            c.number_format, c.font = PCT, AZUL
        fi_c, fc_c, finv_c = f"$B{fila_s}", f"$C{fila_s}", f"$D{fila_s}"
        for t in range(n + 1):
            uai = f"({F('ing', t)}*(1+{fi_c})-{F('cos', t)}*(1+{fc_c})-{F('dep', t)}*(1+{finv_c}))"
            vr_t = (f"+{F('vr', t)}*IF({ref['vr_metodo']}=\"valor_libro\",1+{finv_c},1)" if t == n else "")
            formula = (f"={uai}-MAX(0,{uai})*{ref['impuesto']}+{F('dep', t)}*(1+{finv_c})"
                       f"-{F('inv', t)}*(1+{finv_c})-{F('dct', t)}*(1+{fi_c}){vr_t}")
            c = se.cell(fila_s, 5 + t, formula)
            c.number_format = MONEDA
        c0, c1, cn = L(5), L(6), L(5 + n)
        van = se.cell(fila_s, 6 + n, f"={c0}{fila_s}+NPV({ref['tasa']},{c1}{fila_s}:{cn}{fila_s})")
        van.number_format, van.font = MONEDA, NEGRITA
        fila_s += 1
    fila_s += 1
    se.cell(fila_s, 1, "Sensibilidad a la tasa de descuento").font = NEGRITA
    fila_s += 1
    base_tasa = float(s["tasa_descuento"])
    for d in (-0.02, -0.01, 0.0, 0.01, 0.02):
        se.cell(fila_s, 1, f"Tasa {base_tasa + d:.1%}").font = NORMAL
        c = se.cell(fila_s, 2, f"={ref['tasa']}+({d})")
        c.number_format = PCT
        v = se.cell(fila_s, 3, f"=Flujo!$B${filas['fn']}+NPV(B{fila_s},{rango_fn1})")
        v.number_format = MONEDA
        fila_s += 1
    se.column_dimensions["A"].width = 22
    wb.save(salida)
    raiz = Path(a.proyecto) if a.proyecto else next((d for d in ruta.resolve().parents if (d / "proyecto.yaml").exists()), None)
    aut = autoria.cargar_autoria(raiz)
    autoria.limpiar_ooxml(salida, aut.get("nombre", ""), str(s.get("nombre") or "Evaluación económica"), aut.get("organizacion", ""))

    # ------------------------- Control en Python -------------------------
    flujo, van = calcular(s)
    t_ = tir(flujo)
    acum, rec = 0, None
    for i, f in enumerate(flujo):
        acum += f
        if acum >= 0 and rec is None:
            rec = i
    print(f"Modelo generado: {salida}")
    print(f"VAN ({float(s['tasa_descuento']):.1%}): {van:,.0f} {s.get('moneda', '')}".replace(",", "."))
    print(f"TIR: {t_:.1%}" if t_ is not None else "TIR: no calculable")
    print(f"Recuperación: {'no recupera' if rec is None else str(rec) + ' años'}")
    print("Sensibilidad (VAN):")
    for v in variables:
        vals = []
        for x in variaciones:
            kw = {"fi": x} if v == "ingresos" else {"fc": x} if v == "costos" else {"finv": x}
            vals.append(f"{x:+.0%}: {calcular(s, **kw)[1]:,.0f}".replace(",", "."))
        print(f"  {v}: " + " | ".join(vals))


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    main()

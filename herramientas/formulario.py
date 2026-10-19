#!/usr/bin/env python3
"""Formulario de una plataforma de postulación (CORFO, fondos.gob.cl u otra) cuando no hay un documento que llenar.

Uso:
    python herramientas/formulario.py <proyecto> --reconstruir <archivo> [<archivo> ...] [--color 0000FF]
    python herramientas/formulario.py <proyecto>                    # control de campos (lo usa postulacion.py)
    python herramientas/formulario.py <proyecto> --limpiar          # corrige caracteres invisibles y letras de otro alfabeto
    python herramientas/formulario.py <proyecto> --word             # Word de trabajo: pregunta, instrucción en azul, conteo, respuesta
    python herramientas/formulario.py <proyecto> --hoja-carga [--abrir]   # página para copiar y pegar campo por campo

El formulario vive en `postulacion/formulario.yaml`: un campo por pregunta de la plataforma, con su sección, si es
obligatorio, su límite de caracteres (incluidos los espacios), su ayuda y la respuesta (corta, en `respuesta`, o larga,
en el archivo que indica `archivo`, por ejemplo `propuesta/campos/C05.md`).

- --reconstruir arma un BORRADOR (`postulacion/formulario.borrador.yaml`) desde los formularios vacíos o impresos del
  portal (PDF o texto: «NOMBRE DEL CAMPO (*)», límite en la línea siguiente, luego la ayuda) o desde un Word donde las
  instrucciones de la plataforma están en color («1. Título *», ayuda, contador «0 / 250»). Nunca sobrescribe el
  formulario vigente: el borrador se revisa contra la plataforma y luego se guarda como `formulario.yaml`.
- El control cuenta caracteres como la plataforma (texto plano, sin marcas de formato) y avisa campos que exceden su
  límite, obligatorios vacíos y caracteres raros pegados desde la web (invisibles o de otro alfabeto, como una «з»
  cirílica en «Miniзterio»), que algunas plataformas rechazan o cuentan distinto.
"""
import argparse
import html
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

import yaml

ARCHIVO = Path("postulacion") / "formulario.yaml"
BORRADOR = Path("postulacion") / "formulario.borrador.yaml"
INVISIBLES = re.compile("[​-‏  ‪-‮⁠-⁤﻿­]")
# Letras de otros alfabetos idénticas a una latina (las inequívocas se corrigen solas)
HOMOGLIFOS = {"а": "a", "е": "e", "о": "o", "р": "p", "с": "c", "у": "y", "х": "x", "і": "i", "ј": "j", "ѕ": "s",
              "А": "A", "В": "B", "Е": "E", "К": "K", "М": "M", "Н": "H", "О": "O", "Р": "P", "С": "C", "Т": "T", "Х": "X",
              "ο": "o", "Ο": "O", "α": "a", "Α": "A", "Ε": "E", "Ι": "I", "Κ": "K", "Μ": "M", "Ν": "N", "Τ": "T", "Χ": "X"}
# Restos de la impresión desde el navegador: marcas de página, URL con «1/4», fecha y hora del encabezado, «×»
ARTEFACTO = re.compile(r"^(-+\s*P[áa]gina\s+\d+\s*-+|https?://\S+(\s+\d+\s*/\s*\d+)?|\d{1,2}/\d{1,2}/\d{2,4},?\s+\d{1,2}:\d{2}.*|"
                       r"×|<!--.*-->|#\s.*\.(pdf|docx?))$", re.I)


def _es_artefacto(linea):
    """Resto de impresión; «1/4» es un número de página, pero «0 / 250» es el contador de caracteres de un campo."""
    pagina = re.fullmatch(r"(\d+)\s*/\s*(\d+)", linea)
    return bool(ARTEFACTO.match(linea)) or bool(pagina and int(pagina.group(2)) < 50)


# ----------------------------------------------------------------------------------------- lectura y texto plano
def cargar(raiz, ruta=ARCHIVO):
    p = Path(raiz) / ruta
    return (yaml.safe_load(p.read_text(encoding="utf-8")) or {}) if p.exists() else {}


def texto_plano(md):
    """Lo que se pega en la plataforma: sin encabezado YAML, sin marcas de Markdown, sin espacios sobrantes."""
    t = re.sub(r"\A---\n.*?\n---\n", "", md or "", count=1, flags=re.S)
    t = re.sub(r"<!--.*?-->", "", t, flags=re.S)
    t = re.sub(r"^#{1,6}\s+", "", t, flags=re.M)
    t = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), t)
    t = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"\1", t)
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    t = t.replace(" ", " ")
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def respuesta(raiz, campo):
    if campo.get("archivo"):
        p = Path(raiz) / campo["archivo"]
        return texto_plano(p.read_text(encoding="utf-8")) if p.exists() else ""
    return texto_plano(str(campo.get("respuesta") or ""))


def raros(texto):
    """Caracteres que conviene revisar: invisibles y letras de alfabetos no latinos (con el contexto donde aparecen)."""
    hallazgos = []
    for m in re.finditer(r"\S*[^\x00-ɏ‐-‧‰-⁞€™\n\t ]\S*|\S*" + INVISIBLES.pattern + r"\S*", texto):
        palabra = m.group(0)
        malos = [c for c in palabra if INVISIBLES.match(c) or (ord(c) > 0x24F and unicodedata.category(c).startswith("L"))]
        if malos:
            nombres = ", ".join(sorted({unicodedata.name(c, hex(ord(c))).title() for c in malos}))
            hallazgos.append(f"«{INVISIBLES.sub('·', palabra)[:40]}» ({nombres})")
    return hallazgos


def limpiar_texto(texto):
    t = INVISIBLES.sub("", texto).replace(" ", " ")
    return "".join(HOMOGLIFOS.get(c, c) for c in t)


# ----------------------------------------------------------------------------------------------------- control
def controlar(raiz, hoy=None):
    """Métricas y hallazgos del formulario. Sin formulario.yaml devuelve ceros (no bloquea las compuertas)."""
    f = cargar(raiz)
    campos = f.get("campos") or []
    m = {"formulario_campos": len(campos), "formulario_campos_excedidos": 0, "formulario_obligatorios_vacios": 0,
         "formulario_caracteres_raros": 0, "formulario_campos_listos": 0}
    filas, crit, adv = [], [], []
    for c in campos:
        t = respuesta(raiz, c)
        n = len(t)
        lim = c.get("limite")
        lim = int(re.sub(r"\D", "", str(lim))) if lim not in (None, "") and re.sub(r"\D", "", str(lim)) else None
        oblig = str(c.get("obligatorio", "")).lower() in ("si", "sí", "true", "1", "x")
        estado = "ok"
        if not t:
            estado = "vacío"
            if oblig:
                m["formulario_obligatorios_vacios"] += 1
        elif lim and n > lim:
            estado = f"excede en {n - lim}"
            m["formulario_campos_excedidos"] += 1
            adv.append(f"Campo {c.get('id')} «{str(c.get('pregunta'))[:60]}»: {n} caracteres, el límite es {lim} (sobran {n - lim}).")
        elif lim and n + t.count("\n") > lim:
            adv.append(f"Campo {c.get('id')}: {n} caracteres cabe, pero no si la plataforma cuenta cada salto de línea como 2 "
                       f"({n + t.count(chr(10))} > {lim}). Deja margen o une párrafos.")
        r = raros(t) + raros(str(c.get("pregunta") or ""))
        if r:
            m["formulario_caracteres_raros"] += len(r)
            adv.append(f"Campo {c.get('id')}: caracteres raros {', '.join(r[:3])} → `formulario.py <proyecto> --limpiar` o corrígelos a mano.")
        if t and estado == "ok":
            m["formulario_campos_listos"] += 1
        filas.append({"id": c.get("id"), "seccion": c.get("seccion", ""), "pregunta": c.get("pregunta", ""), "obligatorio": oblig,
                      "limite": lim, "usados": n, "estado": estado, "texto": t, "ayuda": c.get("ayuda", ""),
                      "opciones": c.get("opciones") or [], "tipo": c.get("tipo", "texto")})
    if m["formulario_obligatorios_vacios"]:
        vacios = [x["id"] for x in filas if x["estado"] == "vacío" and x["obligatorio"]]
        adv.append(f"Obligatorios sin respuesta: {', '.join(vacios[:15])}{'…' if len(vacios) > 15 else ''}.")
    return m, filas, crit, adv, f


# -------------------------------------------------------------------------------------------- reconstrucción
def _es_titulo(linea):
    letras = [c for c in linea if c.isalpha()]
    return len(letras) >= 4 and len(linea) <= 160 and sum(c.isupper() for c in letras) / len(letras) >= 0.85


def _limite(linea):
    m = re.fullmatch(r"(?:\d[\d.,]*\s*/\s*)?(\d[\d.,]*)", linea.strip())  # «1,500» o el contador «0 / 250»
    return int(re.sub(r"[.,]", "", m.group(1))) if m else None


def desde_texto(texto, seccion):
    """Campos desde el texto de un formulario impreso del portal (títulos en mayúsculas o preguntas numeradas)."""
    lineas = [l.strip() for l in texto.splitlines()]
    lineas = [l for l in lineas if l and not _es_artefacto(l) and "fondos.gob.cl" not in l.lower()]
    campos, actual, titulo_sec = [], None, ""
    for l in lineas:
        numerada = re.match(r"^\d{1,2}\.\s+\S.{3,}", l) and (l.rstrip().endswith("*") or "?" in l or len(l) < 120)
        if _es_titulo(l) or numerada:
            # título en dos líneas: «¿... DE 60» + «AÑOS?»
            if actual and not actual["_cerrado"] and _es_titulo(l) and len(l) < 40 and not actual["ayuda"] and actual["limite"] is None:
                actual["pregunta"] += " " + l
                actual["_cerrado"] = l.endswith("?") or "(*)" in l
                actual["obligatorio"] = actual["obligatorio"] or "(*)" in l
                continue
            actual = {"seccion": f"{seccion} · {titulo_sec}" if titulo_sec else seccion, "pregunta": l, "obligatorio": "(*)" in l or l.endswith("*"),
                      "limite": None, "ayuda": "", "_cerrado": l.endswith("?") or "(*)" in l or l.endswith("*")}
            campos.append(actual)
            continue
        if actual is None:
            if not titulo_sec and len(l) < 80:
                titulo_sec = l
            continue
        lim = _limite(l)
        if lim is not None and actual["limite"] is None and (not actual["ayuda"] or lim >= 50):
            actual["limite"] = lim  # al imprimir, el límite a veces queda después de la ayuda (salto de página)
            actual["_cerrado"] = True
            continue
        actual["_cerrado"] = True
        actual["ayuda"] = (actual["ayuda"] + " " + l).strip()
    return campos


def desde_word(ruta, color):
    """Campos desde un Word donde las instrucciones de la plataforma están en color y las respuestas u opciones no."""
    from docx import Document
    doc = Document(str(ruta))
    campos, actual, seccion = [], None, Path(ruta).stem
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        coloreado = sum(len(r.text) for r in p.runs if r.font.color is not None and r.font.color.type is not None
                        and r.font.color.rgb is not None and (str(r.font.color.rgb).upper() == color if color else str(r.font.color.rgb) != "000000"))
        es_instr = coloreado > len(t) / 2
        if (p.style.name or "").lower().startswith(("heading", "título")) and not es_instr:
            seccion = t
            continue
        primera, *resto = t.split("\n")
        numerada = re.match(r"^\d{1,3}[.)]\s+\S", primera)
        if numerada and (es_instr or primera.rstrip().endswith("*")):  # a veces una pregunta quedó sin color
            actual = {"seccion": seccion, "pregunta": primera.rstrip(" *").strip(), "obligatorio": primera.rstrip().endswith("*"),
                      "limite": None, "ayuda": " ".join(x.strip() for x in resto if x.strip()), "opciones": [], "respuesta": ""}
            campos.append(actual)
            continue
        if actual is None:
            continue
        if es_instr:
            lim = _limite(t)
            if lim is not None:
                actual["limite"] = lim
            else:
                actual["ayuda"] = (actual["ayuda"] + " " + t).strip()
        elif actual["limite"] is None and len(t) <= 120:
            actual["opciones"].append(t)  # sin contador de caracteres: son las alternativas de una selección
        else:
            actual["respuesta"] = (actual["respuesta"] + "\n\n" + t).strip()
    return campos


def reconstruir(raiz, archivos, color):
    import _nucleo
    campos = []
    for a in archivos:
        ruta = _nucleo.ruta_larga(a)  # rutas de más de 260 caracteres en Windows (OneDrive, carpetas profundas)
        if not ruta.exists():
            sys.exit(f"No existe {ruta}")
        if ruta.suffix.lower() == ".docx":
            campos += desde_word(ruta, color)
        elif ruta.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            texto = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(ruta)).pages)
            campos += desde_texto(texto, ruta.stem)
        else:
            campos += desde_texto(ruta.read_text(encoding="utf-8", errors="replace"), ruta.stem)
    salida = []
    for i, c in enumerate(campos, 1):
        pregunta = re.sub(r"\s*\(\*\)\s*$|\s*\*\s*$", "", c["pregunta"]).strip()
        x = {"id": f"C{i:02d}", "seccion": c["seccion"], "pregunta": pregunta, "obligatorio": "si" if c["obligatorio"] else "no",
             "limite": c["limite"] or "", "tipo": "seleccion" if c.get("opciones") else "texto", "ayuda": c["ayuda"][:1500]}
        if c.get("opciones"):
            x["opciones"] = c["opciones"]
        x["respuesta"] = c.get("respuesta", "")
        x["archivo"] = ""
        salida.append(x)
    return salida


def _yaml(datos):
    class Dumper(yaml.SafeDumper):
        pass
    Dumper.add_representer(str, lambda d, s: d.represent_scalar("tag:yaml.org,2002:str", s, style="|" if "\n" in s else None))
    return yaml.dump(datos, Dumper=Dumper, allow_unicode=True, sort_keys=False, width=1000)


# ---------------------------------------------------------------------------------------------------- salidas
def generar_word(raiz, filas, f, destino):
    from docx import Document
    from docx.shared import Pt, RGBColor
    import autoria
    import _config
    cfg = _config.cargar(raiz)
    nombre = (cfg.get("proyecto") or {}).get("nombre", Path(raiz).name)
    doc = Document()
    doc.add_heading(f"Formulario de postulación — {nombre}", level=1)
    if f.get("plataforma"):
        doc.add_paragraph(f"Plataforma: {f['plataforma']} · Versión de trabajo del {date.today():%d-%m-%Y}")
    seccion = None
    for x in filas:
        if x["seccion"] != seccion:
            seccion = x["seccion"]
            doc.add_heading(str(seccion), level=2)
        p = doc.add_paragraph()
        r = p.add_run(f"{x['id']}. {x['pregunta']}{' *' if x['obligatorio'] else ''}")
        r.bold = True
        if x["ayuda"]:
            ra = doc.add_paragraph().add_run(str(x["ayuda"]))
            ra.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)
            ra.font.size = Pt(9)
        if x["opciones"]:
            for o in x["opciones"]:
                doc.add_paragraph(f"{'☒' if x['texto'] and x['texto'].strip() == str(o).strip() else '☐'} {o}")
        lim = f"{x['usados']:,} / {x['limite']:,} caracteres".replace(",", ".") if x["limite"] else f"{x['usados']:,} caracteres".replace(",", ".")
        rc = doc.add_paragraph().add_run(f"{lim} · {x['estado']}")
        rc.font.size = Pt(8)
        rc.font.color.rgb = RGBColor(0xC0, 0x00, 0x00) if x["estado"].startswith("excede") else RGBColor(0x59, 0x59, 0x59)
        if x["texto"] and not x["opciones"]:
            for parrafo in x["texto"].split("\n\n"):
                doc.add_paragraph(parrafo)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    aut = autoria.cargar_autoria(raiz)
    autoria.limpiar_ooxml(destino, aut.get("nombre", ""), f"Formulario de postulación — {nombre}", aut.get("organizacion", ""))


def generar_hoja_carga(raiz, filas, f, destino):
    import _config
    cfg = _config.cargar(raiz)
    nombre = html.escape(str((cfg.get("proyecto") or {}).get("nombre", Path(raiz).name)))
    listos = sum(1 for x in filas if x["estado"] == "ok" and x["texto"])
    excedidos = sum(1 for x in filas if x["estado"].startswith("excede"))
    vacios = sum(1 for x in filas if x["estado"] == "vacío" and x["obligatorio"])
    tarjetas, seccion = [], None
    for i, x in enumerate(filas):
        if x["seccion"] != seccion:
            seccion = x["seccion"]
            tarjetas.append(f'<h2>{html.escape(str(seccion))}</h2>')
        estado = x["estado"]
        clase = "mal" if estado.startswith("excede") or (estado == "vacío" and x["obligatorio"]) else ("vacio" if estado == "vacío" else "bien")
        conteo = f"{x['usados']} / {x['limite']}" if x["limite"] else f"{x['usados']} car."
        ayuda = f'<p class="ayuda">{html.escape(str(x["ayuda"]))}</p>' if x["ayuda"] else ""
        tarjetas.append(
            f'<section class="campo {clase}"><div class="cab"><span class="id">{html.escape(str(x["id"]))}</span>'
            f'<span class="preg">{html.escape(str(x["pregunta"]))}{" <b class=ob>*</b>" if x["obligatorio"] else ""}</span>'
            f'<span class="cont">{conteo} · {html.escape(estado)}</span></div>{ayuda}'
            f'<textarea id="t{i}" readonly rows="{min(14, max(2, x["texto"].count(chr(10)) + len(x["texto"]) // 110 + 1))}">'
            f'{html.escape(x["texto"])}</textarea>'
            f'<button onclick="copiar({i}, this)" {"disabled" if not x["texto"] else ""}>Copiar</button></section>')
    plataforma = html.escape(str(f.get("plataforma") or "Plataforma de postulación"))
    cuerpo = "".join(tarjetas)
    pagina = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Hoja de carga — {nombre}</title><style>
:root{{--fondo:#f6f7f9;--tarjeta:#fff;--texto:#1d2330;--suave:#5b6474;--borde:#dde1e7;--bien:#1f7a4d;--mal:#b42318;--acento:#2457c5}}
@media (prefers-color-scheme: dark){{:root{{--fondo:#14171c;--tarjeta:#1d2129;--texto:#e7eaf0;--suave:#9aa3b2;--borde:#2c323d;--bien:#4cc38a;--mal:#ff6b5b;--acento:#7aa2ff}}}}
body{{margin:0;background:var(--fondo);color:var(--texto);font:15px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}}
main{{max-width:900px;margin:0 auto;padding:20px 16px 60px}} h1{{font-size:22px;margin:0 0 4px}} h2{{font-size:16px;margin:28px 0 8px;color:var(--suave)}}
.resumen{{display:flex;gap:12px;flex-wrap:wrap;margin:12px 0 4px}} .kpi{{background:var(--tarjeta);border:1px solid var(--borde);border-radius:10px;padding:8px 14px}}
.campo{{background:var(--tarjeta);border:1px solid var(--borde);border-left:4px solid var(--bien);border-radius:10px;padding:12px;margin:10px 0}}
.campo.mal{{border-left-color:var(--mal)}} .campo.vacio{{border-left-color:var(--borde)}}
.cab{{display:flex;gap:8px;align-items:baseline;flex-wrap:wrap}} .id{{font-weight:700;color:var(--acento)}} .preg{{flex:1;font-weight:600}}
.cont{{font-size:13px;color:var(--suave)}} .mal .cont{{color:var(--mal);font-weight:600}} .ob{{color:var(--mal)}}
.ayuda{{font-size:13px;color:var(--suave);margin:6px 0}}
textarea{{width:100%;box-sizing:border-box;background:var(--fondo);color:var(--texto);border:1px solid var(--borde);border-radius:8px;padding:8px;font:14px/1.45 inherit;resize:vertical}}
button{{margin-top:6px;background:var(--acento);color:#fff;border:0;border-radius:8px;padding:6px 16px;font-weight:600;cursor:pointer}}
button:disabled{{opacity:.35;cursor:default}}
</style></head><body><main>
<h1>Hoja de carga — {nombre}</h1>
<div style="color:var(--suave)">{plataforma} · generada el {date.today():%d-%m-%Y}. Copia cada campo y pégalo en la plataforma.</div>
<div class="resumen"><span class="kpi">Campos: {len(filas)}</span><span class="kpi">Listos: {listos}</span>
<span class="kpi" style="color:var(--mal)">Exceden el límite: {excedidos}</span><span class="kpi">Obligatorios vacíos: {vacios}</span></div>
{cuerpo}
</main><script>
function copiar(i, b){{const t=document.getElementById('t'+i);const ok=()=>{{b.textContent='Copiado ✓';setTimeout(()=>b.textContent='Copiar',1600)}};
if(navigator.clipboard&&window.isSecureContext){{navigator.clipboard.writeText(t.value).then(ok)}}else{{t.removeAttribute('readonly');t.select();document.execCommand('copy');t.setAttribute('readonly','');ok()}}}}
</script></body></html>"""
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(pagina, encoding="utf-8")


# -------------------------------------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("proyecto")
    ap.add_argument("--reconstruir", nargs="+", metavar="ARCHIVO", help="formularios vacíos o impresos del portal, o un Word con las instrucciones en color")
    ap.add_argument("--color", default="", help="color hexadecimal de las instrucciones en el Word (por defecto, cualquier color distinto del negro)")
    ap.add_argument("--limpiar", action="store_true", help="corregir caracteres invisibles y letras de otro alfabeto en las respuestas")
    ap.add_argument("--word", action="store_true", help="generar el Word de trabajo")
    ap.add_argument("--hoja-carga", action="store_true", help="generar la página para copiar y pegar en la plataforma")
    ap.add_argument("--abrir", action="store_true", help="abrir la hoja de carga en el navegador")
    a = ap.parse_args()
    raiz = Path(a.proyecto)

    if a.reconstruir:
        campos = reconstruir(raiz, a.reconstruir, a.color.upper())
        datos = {"plataforma": "", "reconstruido_desde": [Path(x).name for x in a.reconstruir], "campos": campos}
        (raiz / BORRADOR).parent.mkdir(parents=True, exist_ok=True)
        (raiz / BORRADOR).write_text("# BORRADOR reconstruido automáticamente: revísalo contra la plataforma (preguntas, límites,\n"
                                     "# obligatorios, opciones) y guárdalo como postulacion/formulario.yaml.\n" + _yaml(datos), encoding="utf-8")
        con_lim = sum(1 for c in campos if c["limite"])
        print(f"Borrador: {(raiz / BORRADOR).as_posix()} · campos: {len(campos)} · con límite de caracteres: {con_lim} · "
              f"obligatorios: {sum(1 for c in campos if c['obligatorio'] == 'si')} · con respuesta previa: {sum(1 for c in campos if c['respuesta'])}")
        for c in campos[:12]:
            limite = f" · {c['limite']} car." if c["limite"] else ""
            marca = " *" if c["obligatorio"] == "si" else ""
            print(f"  - {c['id']} [{c['seccion'][:30]}] {c['pregunta'][:70]}{marca}{limite}")
        if len(campos) > 12:
            print(f"  - … y {len(campos) - 12} más")
        return 0

    if not (raiz / ARCHIVO).exists():
        sys.exit(f"No existe {ARCHIVO.as_posix()}: reconstrúyelo con --reconstruir <formularios del portal> y revisa el borrador.")
    if a.limpiar:
        f = cargar(raiz)
        cambios = 0
        for c in f.get("campos") or []:
            if c.get("archivo") and (raiz / c["archivo"]).exists():
                p = raiz / c["archivo"]
                t = p.read_text(encoding="utf-8")
                n = limpiar_texto(t)
                if n != t:
                    p.write_text(n, encoding="utf-8")
                    cambios += 1
            elif c.get("respuesta"):
                n = limpiar_texto(str(c["respuesta"]))
                if n != c["respuesta"]:
                    c["respuesta"] = n
                    cambios += 1
        (raiz / ARCHIVO).write_text(_yaml(f), encoding="utf-8")
        print(f"Campos corregidos: {cambios}. Lo que no es una letra latina inequívoca (por ejemplo «з») queda para revisar a mano.")

    m, filas, crit, adv, f = controlar(raiz)
    if a.word:
        destino = raiz / "propuesta" / "salida" / f"formulario-trabajo_{date.today().isoformat()}.docx"
        generar_word(raiz, filas, f, destino)
        print(f"Word de trabajo: {destino.as_posix()}")
    if a.hoja_carga:
        destino = raiz / "propuesta" / "salida" / f"hoja-de-carga_{date.today().isoformat()}.html"
        generar_hoja_carga(raiz, filas, f, destino)
        print(f"Hoja de carga: {destino.as_posix()}")
        if a.abrir:
            import webbrowser
            webbrowser.open(destino.resolve().as_uri())
    print(f"\n# Formulario — {f.get('plataforma') or 'plataforma'}\n")
    print(f"- Campos: {m['formulario_campos']} · listos: {m['formulario_campos_listos']} · exceden su límite: {m['formulario_campos_excedidos']} · "
          f"obligatorios vacíos: {m['formulario_obligatorios_vacios']} · caracteres raros: {m['formulario_caracteres_raros']}")
    print("\n| Campo | Pregunta | Límite | Usados | Estado |\n|---|---|---|---|---|")
    for x in filas:
        print(f"| {x['id']} | {str(x['pregunta'])[:60]}{' *' if x['obligatorio'] else ''} | {x['limite'] or '—'} | {x['usados']} | {x['estado']} |")
    for titulo, lista in (("Crítico", crit), ("Advertencias", adv)):
        if lista:
            print(f"\n## {titulo}\n" + "\n".join(f"- {x}" for x in lista))
    return 0


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    sys.exit(main())

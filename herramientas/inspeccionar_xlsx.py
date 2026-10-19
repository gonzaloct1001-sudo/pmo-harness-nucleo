#!/usr/bin/env python3
"""Describe la estructura de un Excel para decidir cómo actualizarlo sin romperlo.

Uso:
    python herramientas/inspeccionar_xlsx.py archivo.xlsx [--celdas 40]

Muestra por hoja: rango usado, celdas combinadas, protección, cantidad de fórmulas,
ejemplos de fórmulas y las celdas con valores fijos (candidatas a celdas de entrada),
destacando las que tienen fuente azul o relleno amarillo (convención habitual de inputs).
También lista los nombres definidos y los vínculos a otros archivos, y advierte si la planilla tiene elementos que se
pierden al guardarla con openpyxl (listas desplegables modernas, formas, macros): en ese caso se llena con
`llenar_xlsx.py`, que escribe solo los valores y deja todo lo demás intacto.
"""
import argparse
import sys
import warnings
import zipfile


def riesgos_al_guardar(ruta):
    """Elementos de la planilla que openpyxl no conserva al guardar (y que una entidad podría exigir)."""
    riesgos = []
    try:
        with zipfile.ZipFile(ruta) as z:
            nombres = z.namelist()
            hojas = "".join(z.read(n).decode("utf-8", "ignore") for n in nombres if n.startswith("xl/worksheets/sheet"))
            dibujos = "".join(z.read(n).decode("utf-8", "ignore") for n in nombres if n.startswith("xl/drawings/drawing"))
    except (zipfile.BadZipFile, OSError):
        return riesgos
    if "x14:dataValidation" in hojas:
        riesgos.append(f"{hojas.count('x14:dataValidation ')} lista(s) desplegable(s) modernas")
    if "<xdr:sp" in dibujos:
        riesgos.append(f"{dibujos.count('<xdr:sp')} forma(s) o cuadro(s) de texto")
    if "x14:conditionalFormatting" in hojas:
        riesgos.append("formatos condicionales modernos")
    if any("vbaProject" in n for n in nombres):
        riesgos.append("macros")
    return riesgos

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Falta openpyxl: pip install openpyxl")


def color_de(obj):
    try:
        c = obj.rgb if obj is not None else None
        return c if isinstance(c, str) else None
    except Exception:
        return None


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("archivo")
    p.add_argument("--celdas", type=int, default=40, help="máximo de celdas de entrada a listar por hoja")
    args = p.parse_args()

    keep_vba = args.archivo.lower().endswith(".xlsm")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # los avisos de openpyxl se traducen abajo, en lenguaje simple
        wb = load_workbook(args.archivo, keep_vba=keep_vba)
    print(f"# Estructura de {args.archivo}\n")
    riesgos = riesgos_al_guardar(args.archivo)
    if riesgos:
        print(f"**Atención — no guardar con openpyxl:** tiene {', '.join(riesgos)}, que se borrarían. Para llenarla usa "
              "`python herramientas/llenar_xlsx.py` (escribe solo los valores y deja todo lo demás intacto).\n")
    if getattr(wb, "_external_links", None):
        print(f"**Atención:** tiene {len(wb._external_links)} vínculo(s) a otros archivos. "
              "No guardar con openpyxl sin respaldar esos valores.\n")
    nombres = list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else []
    if nombres:
        print("Nombres definidos: " + ", ".join(nombres) + "\n")

    for ws in wb.worksheets:
        formulas, entradas = [], []
        for fila in ws.iter_rows():
            for c in fila:
                if c.value is None:
                    azul = (color_de(c.font.color) or "").upper().endswith("0000FF")
                    amarillo = (color_de(c.fill.fgColor) or "").upper().endswith("FFFF00")
                    if azul or amarillo:
                        entradas.append((c.coordinate, "(vacía)", " (formato de entrada)"))
                    continue
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append((c.coordinate, c.value))
                elif not isinstance(c.value, str) or (color_de(c.font.color) or "").upper().endswith("0000FF"):
                    azul = (color_de(c.font.color) or "").upper().endswith("0000FF")
                    amarillo = (color_de(c.fill.fgColor) or "").upper().endswith("FFFF00")
                    marca = " (fuente azul)" if azul else ""
                    marca += " (relleno amarillo)" if amarillo else ""
                    entradas.append((c.coordinate, c.value, marca))
        print(f"## Hoja: {ws.title}")
        print(f"- Rango usado: {ws.dimensions} · Estado: {ws.sheet_state}")
        print(f"- Protegida: {'sí' if ws.protection.sheet else 'no'} · Celdas combinadas: {len(ws.merged_cells.ranges)}")
        print(f"- Fórmulas: {len(formulas)}")
        for coord, f in formulas[:8]:
            print(f"  - {coord}: `{f}`")
        if len(formulas) > 8:
            print(f"  - … y {len(formulas) - 8} más")
        print(f"- Celdas de entrada probables (valores fijos o con formato de entrada): {len(entradas)}")
        for coord, v, marca in entradas[: args.celdas]:
            print(f"  - {coord} = {v}{marca}")
        # encabezados probables: primera fila con 3+ textos
        for fila in ws.iter_rows(max_row=min(ws.max_row, 15)):
            textos = [c for c in fila if isinstance(c.value, str) and not c.value.startswith("=")]
            if len(textos) >= 3:
                print(f"- Encabezados probables (fila {fila[0].row}): " +
                      " | ".join(f"{c.coordinate}={c.value}" for c in textos))
                break
        print()


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    main()

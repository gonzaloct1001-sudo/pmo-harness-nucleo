#!/usr/bin/env python3
"""Status de gestión: todos los proyectos de esta empresa en una vista (o de varias empresas, con un panel).

Uso:
    python herramientas/status.py                               # en el chat: semáforos, lo que viene, decisiones, cambios
    python herramientas/status.py --breve                       # una línea por proyecto (lo usa el estado al iniciar sesión)
    python herramientas/status.py --formato html --abrir        # reporte web (también md, xlsx o json: «html,xlsx»)
    python herramientas/status.py --hoy AAAA-MM-DD              # fecha de corte
    python herramientas/status.py --panel <carpeta-panel> [--formato html --abrir]    # tablero único de varias empresas
    python herramientas/status.py --crear-panel <carpeta> --empresas <carpeta-empresa> [<carpeta-empresa> ...]

Lee cada proyecto con los mismos cálculos de sus scripts (postulacion.py, validar_proyecto.py y evaluar_metricas.py,
en modo --json), el traspaso, los vencimientos y la agenda; no modifica ningún dato. Los reportes quedan en `status/`
de la carpeta de la empresa (o en la carpeta del panel), junto con un registro diario en `status/historial/` para
mostrar qué cambió desde el status anterior.

Semáforo (reglas fijas, iguales para cualquier asistente):
- Postulación — rojo: cierre vencido sin entregar, o a 7 días o menos del cierre con requisitos excluyentes pendientes,
  fallas fatales, documentos vencidos, presupuesto que no cumple o formulario incompleto; hito del calendario vencido.
  Amarillo: alguno de esos pendientes con más plazo, o a 14 días o menos con menos del 80% de los requisitos cumplidos.
  Gris: presentada y cerrada (adjudicada, no adjudicada, desistida o No-Go).
- Ejecución y otros — rojo: hito vencido, término fuera del plazo contractual, riesgo alto sin plan, regla crítica
  incumplida o atraso rojo. Amarillo: tareas atrasadas, acuerdos o problemas vencidos, regla de advertencia incumplida,
  atraso amarillo o cambios sin resolver.
- Todos: un vencimiento propio (seguimiento/vencimientos.csv) ya vencido es rojo; a 7 días o menos, amarillo.

Privacidad: el status de cada empresa vive en su carpeta (su repositorio privado). Un tablero que junte varias
empresas va en una carpeta de panel aparte, con su propio repositorio privado, para no mezclar datos entre empresas.
"""
import argparse
import csv
import html
import json
import os
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import yaml

H = Path(__file__).resolve().parent
RAIZ = H.parent
ICONO = {"rojo": "🔴", "amarillo": "🟡", "verde": "🟢", "gris": "⚪"}
ORDEN = {"rojo": 0, "amarillo": 1, "verde": 2, "gris": 3}
TIPOS = {"postulacion": "Postulación", "programa": "Programa", "ejecucion": "Ejecución", "obra": "Obra", "agil": "Ágil",
         "cartera": "Cartera", "evaluacion": "Evaluación", "estudio": "Estudio"}
ABIERTO = ("", "abierto", "abierta", "pendiente", "en curso", "en_curso", "en progreso")
SIN_MOVIMIENTO_DIAS = 21


# ---------------------------------------------------------------------------------------------- utilidades
def fecha(v):
    if isinstance(v, date):
        return v if not isinstance(v, datetime) else v.date()
    try:
        return datetime.strptime(str(v).strip()[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None


def leer_csv(ruta):
    if not ruta.exists():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def ejecutar_json(script, proyecto, raiz_empresa, hoy):
    try:
        r = subprocess.run([sys.executable, str(H / script), str(proyecto), "--json", "--hoy", hoy.isoformat()], cwd=raiz_empresa,
                           capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180)
        return json.loads(r.stdout) if r.returncode == 0 and r.stdout.strip().startswith("{") else None
    except (subprocess.TimeoutExpired, json.JSONDecodeError, OSError):
        return None


def traspaso(carpeta):
    ruta = carpeta / "seguimiento" / "traspaso.md"
    if not ruta.exists():
        return "", [], None
    t = ruta.read_text(encoding="utf-8", errors="replace")
    campo = lambda n: (re.search(rf"^\s*{n}:[ \t]*(.*)$", t, re.M) or [None, ""])[1]  # noqa: E731
    proximo = re.sub(r"\s+#.*$", "", campo("proximo_paso") or "").strip()
    sec = re.search(r"## Decisiones pendientes del usuario\s*\n(.*?)(\n## |\Z)", t, re.S)
    decisiones = [x.strip()[1:].strip() for x in (sec.group(1) if sec else "").splitlines() if x.strip().startswith("-") and x.strip()[1:].strip()]
    return proximo, decisiones, fecha(re.sub(r"\s+#.*$", "", campo("fecha") or ""))


def ultimo_movimiento(carpeta, raiz_empresa):
    try:
        r = subprocess.run(["git", "log", "-1", "--format=%cs", "--", str(carpeta)], cwd=raiz_empresa, capture_output=True,
                           text=True, timeout=15)
        if r.returncode == 0 and r.stdout.strip():
            git_f = fecha(r.stdout.strip())
        else:
            git_f = None
    except (OSError, subprocess.TimeoutExpired):
        git_f = None
    mt = [p.stat().st_mtime for p in carpeta.rglob("*") if p.is_file() and "_entrada" not in p.parts and "status" not in p.parts]
    loc = date.fromtimestamp(max(mt)) if mt else None
    return max([x for x in (git_f, loc) if x], default=None)


def dinero(v, moneda="CLP"):
    try:
        return f"{moneda} {float(v):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return "—"


# ------------------------------------------------------------------------------------------ evaluación
def evaluar_proyecto(carpeta, raiz_empresa, hoy):
    cfg = yaml.safe_load((carpeta / "proyecto.yaml").read_text(encoding="utf-8")) or {}
    arqs = cfg.get("arquetipos") or []
    nombre = str((cfg.get("proyecto") or {}).get("nombre") or carpeta.name)
    if nombre.upper() == "PENDIENTE":
        nombre = carpeta.name
    tipo = " + ".join(TIPOS.get(a, a) for a in arqs if a in TIPOS) or "Proyecto"
    p = {"id": carpeta.name, "nombre": nombre, "tipo": tipo, "etapa": "", "semaforo": "verde", "motivos": [], "indicador": "",
         "proximo": None, "proximo_paso": "", "decisiones": [], "sin_movimiento": None, "agenda": [], "postulacion": None,
         "error": ""}
    rojo, amarillo = p["motivos"], []

    v = ejecutar_json("validar_proyecto.py", carpeta, raiz_empresa, hoy) or {}
    e = ejecutar_json("evaluar_metricas.py", carpeta, raiz_empresa, hoy) or {}
    if not v and not e:
        p["error"] = "no se pudo evaluar (revisa sus archivos con validar_proyecto.py)"
    c = v.get("conteos") or {}

    if "postulacion" in arqs:
        pj = ejecutar_json("postulacion.py", carpeta, raiz_empresa, hoy) or {}
        m = pj.get("metricas") or {}
        resultado = str(pj.get("resultado") or "pendiente").lower()
        etapa_id = str(pj.get("etapa") or "")
        p["etapa"] = pj.get("etapa_nombre") or etapa_id
        p["postulacion"] = {"resultado": resultado, "decision_go": str(pj.get("decision_go") or "pendiente").lower(),
                            "monto": pj.get("monto_solicitado"), "moneda": pj.get("moneda") or "CLP", "etapa": etapa_id,
                            "monto_adjudicado": ((cfg.get("postulacion") or {}).get("monto_adjudicado"))}
        d = m.get("dias_al_cierre")
        pendientes = [(m.get("excluyentes_pendientes"), "requisito(s) excluyente(s) pendiente(s)"),
                      (m.get("fallas_fatales_abiertas"), "falla(s) fatal(es) abierta(s)"),
                      (m.get("contenidos_vencidos"), "documento(s) vencido(s) al cierre"),
                      (m.get("controles_economicos_incumplidos"), "control(es) del presupuesto que no cumple(n)"),
                      (m.get("formulario_campos_excedidos"), "campo(s) del formulario sobre su límite"),
                      (m.get("formulario_obligatorios_vacios"), "campo(s) obligatorio(s) del formulario vacío(s)")]
        pendientes = [f"{n} {t}" for n, t in pendientes if n]
        cerrada = resultado not in ("", "pendiente") or p["postulacion"]["decision_go"] == "no_go" or etapa_id == "cierre"
        if cerrada:
            p["semaforo"] = "gris"
            p["etapa"] = "No-Go" if p["postulacion"]["decision_go"] == "no_go" else f"Cerrada: {resultado.replace('_', ' ')}"
        elif etapa_id == "post_entrega":
            p["etapa"] = "Presentada: esperando resultado"
        else:
            if d is not None and d < 0:
                rojo.append(f"el cierre ({pj.get('cierre')}) pasó y no figura entregada")
            elif d is not None and d <= 7 and pendientes:
                rojo.append(f"cierre en {d} día(s) con {', '.join(pendientes)}")
            elif pendientes:
                amarillo.append(", ".join(pendientes))
            if d is not None and 0 <= d <= 14 and m.get("requisitos_total") and (m.get("cumplimiento_pct") or 0) < 0.8:
                amarillo.append(f"cierre en {d} día(s) con {m.get('cumplimiento_pct', 0):.0%} de los requisitos cumplidos")
            if any(x.startswith("Hito vencido") for x in pj.get("criticos") or []):
                rojo.append("hito del calendario vencido")
        if m.get("requisitos_total"):
            p["indicador"] = f"{m.get('cumplimiento_pct', 0):.0%} de {m['requisitos_total']} requisitos"
        elif m.get("barrido_total"):
            p["indicador"] = f"{m.get('barrido_sin_revisar', 0)} de {m['barrido_total']} candidatos por revisar"
        if not cerrada and p["postulacion"]["decision_go"] == "pendiente" and etapa_id in ("recepcion", "go_no_go"):
            p["decisiones"].append("Decidir si se postula (Go/No-Go)")
        if not cerrada and pj.get("compuerta_cumplida") and etapa_id not in ("post_entrega", "cierre"):
            p["decisiones"].append(f"Aprobar el paso de «{p['etapa']}» a la etapa siguiente")
        if pj.get("cierre") and not cerrada and etapa_id != "post_entrega":
            p["agenda"].append({"fecha": fecha(pj["cierre"]), "texto": f"Cierre de la postulación{' ' + str(pj['hora_cierre']) if pj.get('hora_cierre') else ''}",
                                "tipo": "cierre"})
        for x in leer_csv(carpeta / "postulacion" / "calendario.csv"):
            f = fecha(x.get("fecha"))
            if f and (x.get("estado") or "pendiente").strip().lower() not in ("cumplido", "no_aplica") and (x.get("tipo") or "").lower() != "cierre":
                p["agenda"].append({"fecha": f, "texto": x.get("hito") or "hito de la postulación", "tipo": x.get("tipo") or "hito"})
    else:
        p["etapa"] = "En ejecución" if any(a in arqs for a in ("ejecucion", "programa", "obra", "agil")) else tipo
        tot = v.get("avance_total")
        if tot:
            p["indicador"] = f"avance {tot['real']:.0f}% (plan {tot['plan']:.0f}%)"
        estados = [x.get("estado") for x in v.get("avance") or []]
        if c.get("hitos_vencidos"):
            rojo.append(f"{c['hitos_vencidos']} hito(s) vencido(s)")
        if c.get("fuera_de_plazo_contractual"):
            rojo.append("hay tareas que terminan después del plazo contractual")
        if "Rojo" in estados:
            rojo.append("atraso relevante frente al plan")
        elif "Amarillo" in estados:
            amarillo.append("atraso moderado frente al plan")
        if c.get("tareas_atrasadas"):
            amarillo.append(f"{c['tareas_atrasadas']} tarea(s) atrasada(s)")
        if c.get("cambios_sin_resolver"):
            amarillo.append(f"{c['cambios_sin_resolver']} cambio(s) sin resolver")
            p["decisiones"].append(f"Resolver {c['cambios_sin_resolver']} solicitud(es) de cambio")
        rojas = [x for x in (e.get("metricas_proyecto") or []) if x.get("semaforo") == "Rojo"]
        if rojas and not p["indicador"]:
            p["indicador"] = f"{len(rojas)} indicador(es) en rojo"
        for x in leer_csv(carpeta / "cronograma" / "cronograma.csv"):
            f = fecha(x.get("fin_plan"))
            try:
                av = float(x.get("avance_pct") or 0)
            except ValueError:
                av = 0
            if f and av < 100 and str(x.get("es_hito", "")).strip().lower() in ("si", "sí", "x", "1", "true") and f >= hoy:
                p["agenda"].append({"fecha": f, "texto": f"Hito: {x.get('tarea')}", "tipo": "hito"})
    # comunes a todos los proyectos
    if c.get("riesgos_altos_sin_plan"):
        rojo.append(f"{c['riesgos_altos_sin_plan']} riesgo(s) alto(s) sin plan")
    if e.get("reglas_criticas"):
        rojo.append(f"{e['reglas_criticas']} regla(s) crítica(s) incumplida(s)")
    if e.get("reglas_advertencia"):
        amarillo.append(f"{e['reglas_advertencia']} regla(s) de advertencia incumplida(s)")
    if c.get("acuerdos_vencidos") or c.get("issues_vencidos"):
        amarillo.append(f"{(c.get('acuerdos_vencidos') or 0) + (c.get('issues_vencidos') or 0)} acuerdo(s) o problema(s) vencido(s)")
    if e.get("candidatos"):
        p["decisiones"].append(f"{len(e['candidatos'])} ítem(s) pueden avanzar de etapa")
    for archivo, campo_f, texto_k in (("seguimiento/acuerdos.csv", "fecha_compromiso", "acuerdo"), ("riesgos/issues.csv", "fecha_compromiso", "titulo")):
        for x in leer_csv(carpeta / archivo):
            f = fecha(x.get(campo_f))
            if f and f >= hoy and (x.get("estado") or "").strip().lower() in ABIERTO:
                p["agenda"].append({"fecha": f, "texto": f"{'Acuerdo' if 'acuerdo' in archivo else 'Problema'}: {x.get(texto_k)}", "tipo": "compromiso"})
    for x in leer_csv(carpeta / "seguimiento" / "vencimientos.csv"):
        f = fecha(x.get("fecha"))
        if not f or (x.get("estado") or "").strip().lower() not in ABIERTO:
            continue
        if f < hoy:
            rojo.append(f"venció «{x.get('compromiso')}» ({f:%d-%m})")
        else:
            if (f - hoy).days <= 7:
                amarillo.append(f"«{x.get('compromiso')}» vence en {(f - hoy).days} día(s)")
            p["agenda"].append({"fecha": f, "texto": x.get("compromiso") or "vencimiento", "tipo": x.get("tipo") or "vencimiento"})

    proximo, dec_traspaso, _ = traspaso(carpeta)
    p["proximo_paso"] = proximo
    p["decisiones"] += dec_traspaso
    if p["semaforo"] != "gris":
        p["semaforo"] = "rojo" if rojo else ("amarillo" if amarillo else "verde")
        p["motivos"] = rojo + amarillo
        mov = ultimo_movimiento(carpeta, raiz_empresa)
        if mov and (hoy - mov).days >= SIN_MOVIMIENTO_DIAS:
            p["sin_movimiento"] = (hoy - mov).days
    p["agenda"] = sorted([x for x in p["agenda"] if x["fecha"]], key=lambda x: x["fecha"])
    futuros = [x for x in p["agenda"] if x["fecha"] >= hoy]
    p["proximo"] = futuros[0] if futuros else None
    if p["error"] and p["semaforo"] == "verde":
        p["semaforo"] = "amarillo"
        p["motivos"].append(p["error"])
    return p


def nombre_empresa(raiz):
    ey = raiz / "empresa" / "empresa.yaml"
    if ey.exists():
        n = str(((yaml.safe_load(ey.read_text(encoding="utf-8")) or {}).get("empresa") or {}).get("nombre") or "")
        if n and n.upper() != "PENDIENTE":
            return n
    return raiz.name


def evaluar_empresa(raiz, hoy):
    proyectos = sorted(p for p in (raiz / "proyectos").glob("*") if (p / "proyecto.yaml").exists()) if (raiz / "proyectos").exists() else []
    ps = [evaluar_proyecto(p, raiz, hoy) for p in proyectos]
    ps.sort(key=lambda p: (ORDEN[p["semaforo"]], (p["proximo"] or {}).get("fecha") or date.max))
    post = [p["postulacion"] for p in ps if p["postulacion"]]
    pipe = {"en_preparacion": sum(1 for x in post if x["resultado"] == "pendiente" and x["decision_go"] != "no_go" and x["etapa"] not in ("post_entrega", "cierre")),
            "presentadas": sum(1 for x in post if x["resultado"] == "pendiente" and x["etapa"] == "post_entrega"),
            "adjudicadas": sum(1 for x in post if x["resultado"] == "adjudicada"),
            "no_adjudicadas": sum(1 for x in post if x["resultado"] in ("no_adjudicada", "inadmisible")),
            "desistidas": sum(1 for x in post if x["resultado"] == "desistida" or x["decision_go"] == "no_go")}
    num = lambda v: float(v) if isinstance(v, (int, float)) or re.fullmatch(r"\d+(\.\d+)?", str(v or "")) else 0.0  # noqa: E731
    pipe["monto_en_juego"] = sum(num(x["monto"]) for x in post if x["resultado"] == "pendiente" and x["decision_go"] != "no_go")
    pipe["monto_adjudicado"] = sum(num(x["monto_adjudicado"] or x["monto"]) for x in post if x["resultado"] == "adjudicada")
    resueltas = pipe["adjudicadas"] + pipe["no_adjudicadas"]
    pipe["tasa"] = pipe["adjudicadas"] / resueltas if resueltas else None
    pipe["moneda"] = next((x["moneda"] for x in post), "CLP")
    return {"empresa": nombre_empresa(raiz), "raiz": str(raiz), "fecha": hoy.isoformat(), "proyectos": ps, "pipeline": pipe if post else None}


# --------------------------------------------------------------------------------------- historial
def historial_dir(destino):
    return destino / "historial"


def guardar_historial(destino, datos):
    d = historial_dir(destino)
    d.mkdir(parents=True, exist_ok=True)
    resumen = {e["empresa"] + "/" + p["id"]: p["semaforo"] for e in datos for p in e["proyectos"]}
    (d / f"{datos[0]['fecha'] if datos else date.today().isoformat()}.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=1), encoding="utf-8")


def anterior(destino, hoy):
    d = historial_dir(destino)
    previos = sorted(x for x in d.glob("*.json") if fecha(x.stem) and fecha(x.stem) < hoy) if d.exists() else []
    return (fecha(previos[-1].stem), json.loads(previos[-1].read_text(encoding="utf-8"))) if previos else (None, {})


def cambios(datos, previo):
    out = []
    for e in datos:
        for p in e["proyectos"]:
            k = e["empresa"] + "/" + p["id"]
            antes = previo.get(k)
            if antes is None:
                if previo:
                    out.append((p, "nuevo"))
            elif antes != p["semaforo"] and p["semaforo"] != "gris":
                out.append((p, "mejoró" if ORDEN[p["semaforo"]] > ORDEN.get(antes, 9) else "empeoró"))
                p["cambio"] = out[-1][1]
    return out


# -------------------------------------------------------------------------------------------- salidas
def agenda_total(datos, hoy, dias):
    items = []
    for e in datos:
        for p in e["proyectos"]:
            for x in p["agenda"]:
                if hoy <= x["fecha"] <= hoy + timedelta(days=dias):
                    items.append({**x, "empresa": e["empresa"], "proyecto": p["nombre"]})
    return sorted(items, key=lambda x: x["fecha"])


def cuando(f, hoy):
    d = (f - hoy).days
    return "hoy" if d == 0 else ("mañana" if d == 1 else f"en {d} días")


def texto_chat(datos, hoy, fecha_prev, cambios_l, breve=False, varias=False):
    ps = [(e, p) for e in datos for p in e["proyectos"]]
    if breve:
        if not ps:
            return []
        lineas = []
        for e, p in ps:
            prox = f"{p['proximo']['texto']} {cuando(p['proximo']['fecha'], hoy)}" if p["proximo"] else "sin vencimientos próximos"
            mot = f" · {p['motivos'][0]}" if p["motivos"] else ""
            paso = f" · próximo paso: {p['proximo_paso']}" if p["proximo_paso"] else ""
            lineas.append(f"{ICONO[p['semaforo']]} {p['nombre']} ({p['etapa'] or p['tipo']}): {prox}{mot}{paso}")
        dec = sum(len(p["decisiones"]) for _, p in ps)
        if dec:
            lineas.append(f"Decisiones pendientes del usuario: {dec}")
        return lineas
    conteo = {k: sum(1 for _, p in ps if p["semaforo"] == k) for k in ICONO}
    titulo = "Tablero general" if varias else f"Status — {datos[0]['empresa']}" if datos else "Status"
    out = [f"# {titulo} · {hoy:%d-%m-%Y}", "",
           " · ".join(f"{ICONO[k]} {v}" for k, v in conteo.items()) + f" · próximos 7 días: {len(agenda_total(datos, hoy, 7))}"
           f" · decisiones pendientes: {sum(len(p['decisiones']) for _, p in ps)}", ""]
    if not ps:
        return out + ["No hay proyectos todavía."]
    out += ["| 🚦 | " + ("Empresa | " if varias else "") + "Proyecto | Etapa | Próximo vencimiento | Avance | Alerta principal | Próximo paso |",
            "|---|" + ("---|" if varias else "") + "---|---|---|---|---|---|"]
    for e, p in ps:
        prox = f"{p['proximo']['texto']} ({p['proximo']['fecha']:%d-%m}, {cuando(p['proximo']['fecha'], hoy)})" if p["proximo"] else "—"
        out.append(f"| {ICONO[p['semaforo']]} | " + (f"{e['empresa']} | " if varias else "") +
                   f"{p['nombre']} | {p['etapa'] or p['tipo']} | {prox} | {p['indicador'] or '—'} | {(p['motivos'] or ['—'])[0]} | {p['proximo_paso'] or '—'} |")
    ag = agenda_total(datos, hoy, 14)
    if ag:
        out += ["", "## Lo que viene (14 días)"] + [f"- {x['fecha']:%d-%m} ({cuando(x['fecha'], hoy)}) · {x['proyecto']}: {x['texto']}" for x in ag]
    dec = [(p, d) for _, p in ps for d in p["decisiones"]]
    if dec:
        out += ["", "## Requiere tu decisión"] + [f"- {p['nombre']}: {d}" for p, d in dec]
    quietos = [p for _, p in ps if p["sin_movimiento"]]
    if quietos:
        out += ["", "## Sin movimiento"] + [f"- {p['nombre']}: {p['sin_movimiento']} días sin cambios" for p in quietos]
    if cambios_l:
        out += ["", f"## Cambios desde el {fecha_prev:%d-%m-%Y}"] + [f"- {p['nombre']}: {c} ({ICONO[p['semaforo']]})" for p, c in cambios_l]
    for e in datos:
        pp = e["pipeline"]
        if pp:
            tasa = f" · tasa de adjudicación {pp['tasa']:.0%}" if pp["tasa"] is not None else ""
            out += ["", f"## Postulaciones{' — ' + e['empresa'] if varias else ''}",
                    f"- En preparación: {pp['en_preparacion']} · presentadas: {pp['presentadas']} · adjudicadas: {pp['adjudicadas']} · "
                    f"no adjudicadas: {pp['no_adjudicadas']} · desistidas: {pp['desistidas']}{tasa}",
                    f"- En juego: {dinero(pp['monto_en_juego'], pp['moneda'])} · adjudicado: {dinero(pp['monto_adjudicado'], pp['moneda'])}"]
    return out


def html_status(datos, hoy, fecha_prev, cambios_l, varias=False):
    esc = lambda x: html.escape(str(x))  # noqa: E731
    ps = [(e, p) for e in datos for p in e["proyectos"]]
    conteo = {k: sum(1 for _, p in ps if p["semaforo"] == k) for k in ICONO}
    titulo = "Tablero general" if varias else f"Status de proyectos — {datos[0]['empresa']}" if datos else "Status"
    filas = []
    for e, p in ps:
        prox = f"{esc(p['proximo']['texto'])}<br><small>{p['proximo']['fecha']:%d-%m-%Y} · {cuando(p['proximo']['fecha'], hoy)}</small>" if p["proximo"] else "—"
        cambio = {"mejoró": '<span class="sube">▲ mejoró</span>', "empeoró": '<span class="baja">▼ empeoró</span>'}.get(p.get("cambio"), "")
        filas.append(f'<tr><td><span class="pto {p["semaforo"]}"></span></td>' + (f"<td>{esc(e['empresa'])}</td>" if varias else "") +
                     f'<td><b>{esc(p["nombre"])}</b><br><small>{esc(p["tipo"])}</small> {cambio}</td><td>{esc(p["etapa"] or "—")}</td><td>{prox}</td>'
                     f'<td>{esc(p["indicador"] or "—")}</td><td>{"<br>".join(esc(m) for m in p["motivos"][:3]) or "—"}</td>'
                     f'<td>{esc(p["proximo_paso"] or "—")}</td></tr>')

    def lista(titulo_s, items):
        return f"<section><h2>{titulo_s}</h2><ul>{''.join(f'<li>{x}</li>' for x in items)}</ul></section>" if items else ""

    agenda_html = ""
    for dias, etiqueta in ((30, "Próximos 30 días"), (60, "31 a 60 días"), (90, "61 a 90 días")):
        desde = {30: 0, 60: 31, 90: 61}[dias]
        items = [x for x in agenda_total(datos, hoy, dias) if (x["fecha"] - hoy).days >= desde]
        if items:
            agenda_html += f"<h3>{etiqueta}</h3><ul>" + "".join(
                f"<li><b>{x['fecha']:%d-%m}</b> · {esc(x['proyecto'])}{' (' + esc(x['empresa']) + ')' if varias else ''}: {esc(x['texto'])}</li>" for x in items) + "</ul>"
    decisiones = [f"<b>{esc(p['nombre'])}</b>: {esc(d)}" for _, p in ps for d in p["decisiones"]]
    quietos = [f"<b>{esc(p['nombre'])}</b>: {p['sin_movimiento']} días sin cambios" for _, p in ps if p["sin_movimiento"]]
    cambios_h = [f"<b>{esc(p['nombre'])}</b>: {c}" for p, c in cambios_l]
    pipes = ""
    for e in datos:
        pp = e["pipeline"]
        if pp:
            tasa = f"{pp['tasa']:.0%}" if pp["tasa"] is not None else "—"
            pipes += (f'<div class="kpis">' + (f"<span class='kpi tit'>{esc(e['empresa'])}</span>" if varias else "") +
                      f"<span class='kpi'>En preparación <b>{pp['en_preparacion']}</b></span><span class='kpi'>Presentadas <b>{pp['presentadas']}</b></span>"
                      f"<span class='kpi'>Adjudicadas <b>{pp['adjudicadas']}</b></span><span class='kpi'>No adjudicadas <b>{pp['no_adjudicadas']}</b></span>"
                      f"<span class='kpi'>Tasa de adjudicación <b>{tasa}</b></span><span class='kpi'>En juego <b>{dinero(pp['monto_en_juego'], pp['moneda'])}</b></span>"
                      f"<span class='kpi'>Adjudicado <b>{dinero(pp['monto_adjudicado'], pp['moneda'])}</b></span></div>")
    comparado = f" · comparado con el {fecha_prev:%d-%m-%Y}" if fecha_prev else ""
    th_empresa = "<th>Empresa</th>" if varias else ""
    cuerpo = "".join(filas) or '<tr><td colspan="8">No hay proyectos todavía.</td></tr>'
    sec_decisiones = lista("Requiere tu decisión", decisiones)
    sec_agenda = f"<section><h2>Agenda</h2>{agenda_html}</section>" if agenda_html else ""
    sec_postulaciones = f"<section><h2>Postulaciones</h2>{pipes}</section>" if pipes else ""
    sec_cambios = lista("Cambios desde el status anterior", cambios_h)
    sec_quietos = lista("Sin movimiento", quietos)
    return f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(titulo)}</title><style>
:root{{--fondo:#f5f6f8;--tarjeta:#fff;--texto:#1c2230;--suave:#5d6676;--borde:#dfe3e9;--rojo:#d92d20;--amarillo:#e0a100;--verde:#17a34a;--gris:#98a2b3;--acento:#2457c5}}
@media (prefers-color-scheme: dark){{:root{{--fondo:#13161b;--tarjeta:#1c2028;--texto:#e6e9ef;--suave:#9aa3b2;--borde:#2b313c;--acento:#7aa2ff}}}}
body{{margin:0;background:var(--fondo);color:var(--texto);font:15px/1.5 system-ui,-apple-system,Segoe UI,Roboto,sans-serif}}
main{{max-width:1180px;margin:0 auto;padding:22px 16px 60px}} h1{{font-size:24px;margin:0}} h2{{font-size:17px;margin:26px 0 8px}} h3{{font-size:14px;color:var(--suave);margin:14px 0 4px}}
.sub{{color:var(--suave);margin:2px 0 16px}} .kpis{{display:flex;gap:10px;flex-wrap:wrap;margin:8px 0}}
.kpi{{background:var(--tarjeta);border:1px solid var(--borde);border-radius:10px;padding:8px 14px}} .kpi b{{font-size:18px;margin-left:4px}} .tit{{font-weight:700}}
.tabla{{overflow-x:auto;background:var(--tarjeta);border:1px solid var(--borde);border-radius:12px}}
table{{border-collapse:collapse;width:100%;min-width:860px}} th,td{{text-align:left;vertical-align:top;padding:10px 12px;border-bottom:1px solid var(--borde)}}
th{{font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--suave)}} tr:last-child td{{border-bottom:0}} small{{color:var(--suave)}}
.pto{{display:inline-block;width:14px;height:14px;border-radius:50%;margin-top:4px}} .rojo{{background:var(--rojo)}} .amarillo{{background:var(--amarillo)}} .verde{{background:var(--verde)}} .gris{{background:var(--gris)}}
.sube{{color:var(--verde);font-size:12px;font-weight:600}} .baja{{color:var(--rojo);font-size:12px;font-weight:600}}
section ul{{background:var(--tarjeta);border:1px solid var(--borde);border-radius:12px;padding:12px 12px 12px 30px;margin:0}} li{{margin:3px 0}}
</style></head><body><main>
<h1>{esc(titulo)}</h1><div class="sub">Corte al {hoy:%d-%m-%Y}{comparado}</div>
<div class="kpis"><span class="kpi"><span class="pto rojo"></span> Rojo <b>{conteo['rojo']}</b></span><span class="kpi"><span class="pto amarillo"></span> Amarillo <b>{conteo['amarillo']}</b></span>
<span class="kpi"><span class="pto verde"></span> Verde <b>{conteo['verde']}</b></span><span class="kpi"><span class="pto gris"></span> Cerrados <b>{conteo['gris']}</b></span>
<span class="kpi">Próximos 7 días <b>{len(agenda_total(datos, hoy, 7))}</b></span><span class="kpi">Decisiones pendientes <b>{len(decisiones)}</b></span></div>
<h2>Proyectos</h2><div class="tabla"><table><thead><tr><th></th>{th_empresa}<th>Proyecto</th><th>Etapa</th><th>Próximo vencimiento</th><th>Avance</th><th>Alertas</th><th>Próximo paso</th></tr></thead>
<tbody>{cuerpo}</tbody></table></div>
{sec_decisiones}{sec_agenda}{sec_postulaciones}{sec_cambios}{sec_quietos}
</main></body></html>"""


def xlsx_status(datos, hoy, ruta, varias=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import autoria
    wb = Workbook()
    ws = wb.active
    ws.title = "Proyectos"
    cab = (["Empresa"] if varias else []) + ["Semáforo", "Proyecto", "Tipo", "Etapa", "Próximo vencimiento", "Fecha", "Avance", "Alertas", "Próximo paso", "Decisiones"]
    ws.append(cab)
    color = {"rojo": "F8D7D4", "amarillo": "FCEFC7", "verde": "D6F0DF", "gris": "E4E7EC"}
    for e in datos:
        for p in e["proyectos"]:
            ws.append(([e["empresa"]] if varias else []) + [p["semaforo"].capitalize(), p["nombre"], p["tipo"], p["etapa"],
                      (p["proximo"] or {}).get("texto", ""), (p["proximo"] or {}).get("fecha"), p["indicador"],
                      "; ".join(p["motivos"]), p["proximo_paso"], "; ".join(p["decisiones"])])
            ws.cell(ws.max_row, 2 if varias else 1).fill = PatternFill("solid", fgColor=color[p["semaforo"]])
    ag = wb.create_sheet("Agenda")
    ag.append((["Empresa"] if varias else []) + ["Fecha", "Proyecto", "Compromiso", "Tipo"])
    for x in agenda_total(datos, hoy, 90):
        ag.append(([x["empresa"]] if varias else []) + [x["fecha"], x["proyecto"], x["texto"], x["tipo"]])
    for hoja in (ws, ag):
        for c in hoja[1]:
            c.font = Font(bold=True)
        hoja.auto_filter.ref = hoja.dimensions
        hoja.freeze_panes = "A2"
        for col in hoja.columns:
            hoja.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    # autoría de la empresa (se hereda en cualquiera de sus proyectos)
    yamls = [] if varias or not datos else sorted((Path(datos[0]["raiz"]) / "proyectos").glob("*/proyecto.yaml"))
    aut = autoria.cargar_autoria(yamls[0].parent) if yamls else {}
    autoria.limpiar_ooxml(ruta, aut.get("nombre", ""), "Status de proyectos", aut.get("organizacion", ""))


def escribir_salidas(datos, hoy, destino, formatos, fecha_prev, cambios_l, varias, abrir):
    rutas = []
    for fmt in formatos:
        ruta = destino / f"status-{hoy.isoformat()}.{ 'md' if fmt == 'md' else fmt}"
        if fmt == "html":
            ruta.write_text(html_status(datos, hoy, fecha_prev, cambios_l, varias), encoding="utf-8")
        elif fmt == "md":
            ruta.write_text("\n".join(texto_chat(datos, hoy, fecha_prev, cambios_l, varias=varias)) + "\n", encoding="utf-8")
        elif fmt == "xlsx":
            xlsx_status(datos, hoy, ruta, varias)
        elif fmt == "json":
            ruta.write_text(json.dumps(datos, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        rutas.append(ruta)
    if abrir:
        import webbrowser
        html_r = next((r for r in rutas if r.suffix == ".html"), rutas[0] if rutas else None)
        if html_r:
            webbrowser.open(html_r.resolve().as_uri())
    return rutas


# ---------------------------------------------------------------------------------------------- panel
def crear_panel(carpeta, empresas, nombre):
    carpeta = Path(carpeta).resolve()
    carpeta.mkdir(parents=True, exist_ok=True)
    rel = []
    for e in empresas:
        ep = Path(e).resolve()
        if not (ep / "proyectos").exists() and not (ep / ".harness").exists():
            sys.exit(f"{ep} no parece una carpeta de empresa del harness (no tiene proyectos/ ni .harness/).")
        try:  # rutas relativas al panel: así sirve en otro computador con la misma organización de carpetas
            rel.append(Path(os.path.relpath(ep, carpeta)).as_posix())
        except ValueError:  # otra unidad de disco en Windows
            rel.append(ep.as_posix())
    (carpeta / "panel.yaml").write_text(yaml.safe_dump({"nombre": nombre, "empresas": [str(r) for r in rel]}, allow_unicode=True, sort_keys=False),
                                        encoding="utf-8")
    (carpeta / "LEEME.md").write_text(
        f"# {nombre}\n\nTablero único con el status de varias empresas. Lee las carpetas listadas en `panel.yaml` (rutas relativas a\n"
        "esta carpeta) y deja aquí los reportes y su historial. **No contiene los datos de las empresas**: solo sus resúmenes.\n\n"
        "- Privacidad: esta carpeta debe tener su propio repositorio **privado**, independiente de los de cada empresa, para que\n"
        "  solo tú veas la vista unificada.\n"
        "- Actualizar: abre esta carpeta con tu asistente de IA y pide «dame el status general», o ejecuta\n"
        "  `python <carpeta de una empresa>/herramientas/status.py --panel . --formato html --abrir`.\n", encoding="utf-8")
    (carpeta / "AGENTS.md").write_text(
        "# Panel de status (PMO Harness)\n\nEsta carpeta es un tablero único de varias empresas. No es una carpeta de empresa: aquí no se crean\n"
        "proyectos. Para el status general ejecuta, con la ruta de cualquier empresa de `panel.yaml`:\n"
        "`python <empresa>/herramientas/status.py --panel .` (en el chat) o con `--formato html --abrir` (reporte web).\n"
        "Responde en español (Chile). Los reportes son de uso del usuario: no los subas a repositorios de las empresas.\n", encoding="utf-8")
    (carpeta / ".gitignore").write_text("# Lanzador del status semanal (rutas propias de cada computador)\nabrir-status.cmd\n", encoding="utf-8")
    subprocess.run(["git", "init", "-q", "-b", "main"], cwd=carpeta)
    return carpeta, rel


def empresas_panel(carpeta):
    datos = yaml.safe_load((Path(carpeta) / "panel.yaml").read_text(encoding="utf-8")) or {}
    return [(Path(carpeta) / e).resolve() for e in datos.get("empresas") or []], datos.get("nombre") or "Tablero general"


# ----------------------------------------------------------------------------------------------- main
def resumen_breve(raiz=RAIZ, hoy=None):
    """Para el estado de inicio de sesión: una línea por proyecto."""
    hoy = hoy or date.today()
    datos = [evaluar_empresa(Path(raiz), hoy)]
    return texto_chat(datos, hoy, None, [], breve=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--hoy", help="fecha de corte AAAA-MM-DD (por defecto, hoy)")
    ap.add_argument("--breve", action="store_true", help="una línea por proyecto")
    ap.add_argument("--formato", help="html, md, xlsx o json (varios separados por coma); sin esto, respuesta para el chat")
    ap.add_argument("--abrir", action="store_true", help="abrir el reporte en el navegador")
    ap.add_argument("--guardar", action="store_true", help="registrar el status de hoy en el historial aunque no se genere un reporte")
    ap.add_argument("--panel", help="carpeta de un panel (tablero único de varias empresas)")
    ap.add_argument("--crear-panel", help="crear la carpeta de un panel")
    ap.add_argument("--empresas", nargs="+", help="carpetas de empresa del panel (con --crear-panel)")
    ap.add_argument("--nombre", default="Tablero general", help="nombre del panel (con --crear-panel)")
    a = ap.parse_args()
    hoy = fecha(a.hoy) if a.hoy else date.today()
    if hoy is None:
        sys.exit("Fecha inválida (usar AAAA-MM-DD).")

    if a.crear_panel:
        if not a.empresas or len(a.empresas) < 2:
            sys.exit("Indica al menos dos carpetas de empresa con --empresas (un panel junta varias empresas).")
        carpeta, rel = crear_panel(a.crear_panel, a.empresas, a.nombre)
        print(f"Panel creado: {carpeta}\n- Empresas: {', '.join(map(str, rel))}\n- Siguiente paso (privacidad): crea su repositorio PRIVADO, "
              f"independiente de los de cada empresa:\n  gh repo create {carpeta.name} --private --source \"{carpeta}\" --push && "
              f"gh repo edit --add-topic pmo-harness-panel")
        return 0

    if a.panel:
        raices, titulo = empresas_panel(a.panel)
        destino, varias = Path(a.panel).resolve(), True
    else:
        raices, titulo = [RAIZ], None
        destino, varias = RAIZ / "status", False
    faltan = [r for r in raices if not r.exists()]
    if faltan:
        print("Carpetas de empresa no encontradas en este computador (se omiten): " + ", ".join(map(str, faltan)))
    datos = [evaluar_empresa(r, hoy) for r in raices if r.exists()]
    fecha_prev, previo = anterior(destino, hoy)
    cambios_l = cambios(datos, previo)

    if a.breve:
        print("\n".join(texto_chat(datos, hoy, fecha_prev, cambios_l, breve=True)) or "Sin proyectos.")
        return 0
    if a.formato:
        destino.mkdir(parents=True, exist_ok=True)
        formatos = [f.strip().lower() for f in a.formato.split(",") if f.strip()]
        malos = [f for f in formatos if f not in ("html", "md", "xlsx", "json")]
        if malos:
            sys.exit(f"Formato desconocido: {', '.join(malos)} (usar html, md, xlsx o json).")
        rutas = escribir_salidas(datos, hoy, destino, formatos, fecha_prev, cambios_l, varias, a.abrir)
        guardar_historial(destino, datos)
        print("\n".join(texto_chat(datos, hoy, fecha_prev, cambios_l, varias=varias)))
        print("\nReporte: " + ", ".join(r.as_posix() for r in rutas))
        return 0
    if a.guardar:
        guardar_historial(destino, datos)
    print("\n".join(texto_chat(datos, hoy, fecha_prev, cambios_l, varias=varias)))
    return 0


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    sys.exit(main())

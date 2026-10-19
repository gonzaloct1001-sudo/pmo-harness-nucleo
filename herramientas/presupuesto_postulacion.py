#!/usr/bin/env python3
"""Presupuesto de una postulación en Excel con fórmulas, fuentes de financiamiento y controles contra las bases.

Uso:
    python herramientas/presupuesto_postulacion.py <presupuesto.yaml> [--salida presupuesto.xlsx] [--proyecto proyectos/<slug>]

Hojas (todo con fórmulas; celdas de entrada en azul, modificables por el usuario):
- Parametros: reglas de las bases (tope de la entidad en % y monto, aporte propio mínimo, parte pecuniaria
  mínima del aporte propio) con su fuente, y los ítems de gasto con su tope.
- Presupuesto: partidas (cantidad × costo unitario) y reparto por fuente: entidad, propio pecuniario, propio
  no pecuniario y terceros.
- Resumen: totales por ítem y por fuente, % del total y cumplimiento de topes.
- Controles: cada regla de las bases con valor, límite y estado CUMPLE / NO CUMPLE / SIN REGLA.
- Cronograma (si `meses` > 0): distribución mensual de cada partida.

También calcula los controles en Python, los imprime y los guarda en `controles.csv` junto al YAML (lo lee
postulacion.py). Las reglas vienen de las bases: si un parámetro está vacío, el control queda "SIN REGLA".

**La fuente es el YAML.** El Excel es de salida: sirve para revisar y simular (sus fórmulas se recalculan), pero
los cambios definitivos se hacen en presupuesto.yaml y se vuelve a generar; cada ejecución reemplaza el Excel.
Números: se aceptan 0.7, "0,7", "70%" y "1.500.000". Un valor no numérico o un ítem no declarado detienen el script.
"""
import argparse
import csv
import re
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import autoria  # noqa: E402  (módulo del harness en herramientas/)

AZUL = Font(name="Arial", color="0000FF", size=10)
NORMAL = Font(name="Arial", size=10)
NEGRITA = Font(name="Arial", size=10, bold=True)
TITULO = Font(name="Arial", size=13, bold=True, color="1F3864")
ENCAB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONDO_ENC = PatternFill("solid", fgColor="1F3864")
VERDE = PatternFill("solid", fgColor="C6EFCE")
ROJO = PatternFill("solid", fgColor="FFC7CE")
MONEDA = '#,##0;[Red]-#,##0'
PCT = '0.0%'
FUENTES = [("entidad", "Entidad"), ("pecuniario", "Propio pecuniario"), ("no_pecuniario", "Propio no pecuniario"), ("terceros", "Terceros")]
TOL = 1e-9


def num(v, campo=""):
    """Número desde el YAML: acepta 0.7, "0,7", "70%", "1.500.000". Vacío o null = None. Otro texto = error."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    pct = s.endswith("%")
    s = s.rstrip("%")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "")
    s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        sys.exit(f"Valor no numérico en presupuesto.yaml{(' (' + campo + ')') if campo else ''}: {v!r}")
    return x / 100 if pct else x


def encabezado(ws, fila, textos, anchos=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t)
        c.font, c.fill = ENCAB, FONDO_ENC
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(anchos or [], 1):
        ws.column_dimensions[L(j)].width = w


def calcular(s):
    """Controles en Python (mismo criterio que las fórmulas del Excel)."""
    p = s.get("parametros") or {}
    partidas = s.get("partidas") or []
    ids_items = {it.get("id") for it in (s.get("items") or [])}
    for i, x in enumerate(partidas, 1):
        if x.get("item") not in ids_items:
            sys.exit(f"La partida {i} ({x.get('descripcion', '')}) usa el ítem '{x.get('item')}', que no está en `items` "
                     f"({', '.join(str(k) for k in ids_items)}). Corrige el ítem o decláralo.")
    total = sum((num(x.get("cantidad")) or 0) * (num(x.get("costo_unitario")) or 0) for x in partidas)
    por_fuente = {k: 0.0 for k, _ in FUENTES}
    malas = 0
    por_item = {}
    for x in partidas:
        t = (num(x.get("cantidad")) or 0) * (num(x.get("costo_unitario")) or 0)
        suma = sum(num(x.get(k)) or 0 for k, _ in FUENTES)
        if abs(suma - 1) > 1e-4:
            malas += 1
        for k, _ in FUENTES:
            por_fuente[k] += t * (num(x.get(k)) or 0)
        por_item[x.get("item")] = por_item.get(x.get("item"), 0) + t
    propio = por_fuente["pecuniario"] + por_fuente["no_pecuniario"]
    ctrl = []

    def agregar(cid, texto, valor, limite, cumple, fuente=p.get("fuente", "")):
        if limite is None:
            estado = "SIN REGLA"
        elif total <= 0:
            estado = "SIN DATOS"
        else:
            estado = "CUMPLE" if cumple else "NO CUMPLE"
        ctrl.append({"id": cid, "control": texto, "valor": valor, "limite": limite, "estado": estado, "fuente": fuente})

    ctrl.append({"id": "CTL-00", "control": "Costo total del presupuesto mayor que cero", "valor": total, "limite": 0,
                 "estado": "CUMPLE" if total > 0 else "NO CUMPLE", "fuente": "coherencia interna"})
    ent_pct = por_fuente["entidad"] / total if total else 0
    lim = num(p.get("financiamiento_entidad_max_pct"))
    agregar("CTL-01", "Aporte de la entidad / costo total ≤ máximo", ent_pct, lim, lim is not None and ent_pct <= lim + TOL)
    lim = num(p.get("financiamiento_entidad_max_monto"))
    agregar("CTL-02", "Aporte de la entidad ≤ monto máximo", por_fuente["entidad"], lim, lim is not None and por_fuente["entidad"] <= lim + TOL)
    pro_pct = propio / total if total else 0
    lim = num(p.get("aporte_propio_min_pct"))
    agregar("CTL-03", "Aporte propio / costo total ≥ mínimo", pro_pct, lim, lim is not None and pro_pct >= lim - TOL)
    pec = por_fuente["pecuniario"] / propio if propio else 0
    lim = num(p.get("aporte_pecuniario_min_pct_del_propio"))
    agregar("CTL-04", "Aporte pecuniario / aporte propio ≥ mínimo", pec, lim, lim is not None and pec >= lim - TOL)
    ctrl.append({"id": "CTL-05", "control": "Partidas cuyas fuentes no suman 100%", "valor": malas, "limite": 0,
                 "estado": "CUMPLE" if malas == 0 else "NO CUMPLE", "fuente": "coherencia interna"})
    for i, it in enumerate(s.get("items") or [], 6):
        tope = num(it.get("tope_pct"))
        v = por_item.get(it.get("id"), 0) / total if total else 0
        agregar(f"CTL-{i:02d}", f"Ítem {it.get('id')} ({it.get('nombre', '')}) / costo total ≤ tope", v, tope,
                tope is not None and v <= tope + TOL)
    return total, por_fuente, ctrl


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("presupuesto")
    ap.add_argument("--salida")
    ap.add_argument("--proyecto", help="carpeta del proyecto (autoría del archivo); por defecto se busca hacia arriba")
    a = ap.parse_args()
    ruta = Path(a.presupuesto)
    s = yaml.safe_load(ruta.read_text(encoding="utf-8")) or {}
    salida = Path(a.salida) if a.salida else ruta.with_name("presupuesto.xlsx")
    p = s.get("parametros") or {}
    items = s.get("items") or []
    partidas = s.get("partidas") or []
    n = len(partidas)

    wb = Workbook()
    # ---------------- Parametros ----------------
    wp = wb.active
    wp.title = "Parametros"
    wp["A1"] = f"Parámetros de las bases — {s.get('nombre', '')}"
    wp["A1"].font = TITULO
    wp["A2"] = ("Azul = entrada tomada de las bases (vacío = las bases no fijan esa regla). La fuente es presupuesto.yaml: "
                "este Excel sirve para revisar y simular; los cambios definitivos van al YAML.")
    encabezado(wp, 3, ["Parámetro", "Valor", "Fuente"], [58, 16, 50])
    filas_p = [("Financiamiento máximo de la entidad (% del costo total)", "financiamiento_entidad_max_pct", PCT),
               ("Financiamiento máximo de la entidad (monto)", "financiamiento_entidad_max_monto", MONEDA),
               ("Aporte propio mínimo (% del costo total)", "aporte_propio_min_pct", PCT),
               ("Aporte pecuniario mínimo (% del aporte propio)", "aporte_pecuniario_min_pct_del_propio", PCT)]
    for i, (texto, clave, formato) in enumerate(filas_p, 4):
        wp.cell(i, 1, texto).font = NORMAL
        c = wp.cell(i, 2, num(p.get(clave), clave))
        c.font, c.number_format = AZUL, formato
        wp.cell(i, 3, p.get("fuente", "")).font = NORMAL
    wp.cell(9, 1, "Ítems de gasto").font = NEGRITA
    encabezado(wp, 10, ["Ítem", "Nombre", "Tope (% del costo total)"])
    fila_item = {}
    for i, it in enumerate(items, 11):
        wp.cell(i, 1, it.get("id")).font = AZUL
        wp.cell(i, 2, it.get("nombre", "")).font = AZUL
        c = wp.cell(i, 3, num(it.get("tope_pct")))
        c.font, c.number_format = AZUL, PCT
        fila_item[it.get("id")] = i

    # ---------------- Presupuesto ----------------
    wsb = wb.create_sheet("Presupuesto")
    wsb["A1"] = f"Presupuesto — {s.get('nombre', '')} ({s.get('moneda', 'CLP')})"
    wsb["A1"].font = TITULO
    wsb["A2"] = "Azul = entrada. Las fuentes de cada partida son fracciones que deben sumar 100%."
    encabezado(wsb, 3, ["Ítem", "Descripción", "Unidad", "Cantidad", "Costo unitario", "Costo total",
                        "% Entidad", "% Propio pecuniario", "% Propio no pecuniario", "% Terceros", "Suma %",
                        "Entidad", "Propio pecuniario", "Propio no pecuniario", "Terceros"],
               [8, 42, 10, 10, 15, 16, 10, 12, 13, 10, 9, 15, 15, 15, 14])
    ini, fin = 4, 3 + max(n, 1)
    for r, x in enumerate(partidas, ini):
        for col, clave in ((1, "item"), (2, "descripcion"), (3, "unidad")):
            wsb.cell(r, col, x.get(clave, "")).font = AZUL
        for col, clave in ((4, "cantidad"), (5, "costo_unitario")):
            c = wsb.cell(r, col, num(x.get(clave)) or 0)
            c.font, c.number_format = AZUL, MONEDA if col == 5 else "#,##0.##"
        c = wsb.cell(r, 6, f"=D{r}*E{r}")
        c.font, c.number_format = NORMAL, MONEDA
        for j, (clave, _) in enumerate(FUENTES):
            c = wsb.cell(r, 7 + j, num(x.get(clave)) or 0)
            c.font, c.number_format = AZUL, PCT
            m = wsb.cell(r, 12 + j, f"=$F{r}*{L(7 + j)}{r}")
            m.font, m.number_format = NORMAL, MONEDA
        c = wsb.cell(r, 11, f"=SUM(G{r}:J{r})")
        c.font, c.number_format = NORMAL, PCT
    tot = fin + 1
    wsb.cell(tot, 2, "Total").font = NEGRITA
    for col in [6] + list(range(12, 16)):
        c = wsb.cell(tot, col, f"=SUM({L(col)}{ini}:{L(col)}{fin})")
        c.font, c.number_format = NEGRITA, MONEDA
    wsb.conditional_formatting.add(f"K{ini}:K{fin}", CellIsRule(operator="notBetween", formula=["0.9999", "1.0001"], fill=ROJO))
    rango = lambda col: f"Presupuesto!${col}${ini}:${col}${fin}"  # noqa: E731
    TOTAL = f"Presupuesto!$F${tot}"

    # ---------------- Resumen ----------------
    wr = wb.create_sheet("Resumen")
    wr["A1"] = "Resumen por ítem y por fuente"
    wr["A1"].font = TITULO
    encabezado(wr, 3, ["Ítem", "Nombre", "Costo total", "% del total", "Tope", "Estado", "Entidad", "Propio pecuniario",
                       "Propio no pecuniario", "Terceros"], [8, 30, 16, 11, 10, 12, 15, 15, 15, 14])
    fila_res = {}
    for r, it in enumerate(items, 4):
        wr.cell(r, 1, f"=Parametros!A{fila_item[it.get('id')]}").font = NORMAL
        wr.cell(r, 2, f"=Parametros!B{fila_item[it.get('id')]}").font = NORMAL
        c = wr.cell(r, 3, f"=SUMIF({rango('A')},A{r},{rango('F')})")
        c.font, c.number_format = NORMAL, MONEDA
        c = wr.cell(r, 4, f"=IFERROR(C{r}/{TOTAL},0)")
        c.font, c.number_format = NORMAL, PCT
        c = wr.cell(r, 5, f"=IF(Parametros!C{fila_item[it.get('id')]}=\"\",\"\",Parametros!C{fila_item[it.get('id')]})")
        c.font, c.number_format = NORMAL, PCT
        wr.cell(r, 6, f"=IF(E{r}=\"\",\"SIN TOPE\",IF(D{r}<=E{r}+0.0000001,\"CUMPLE\",\"NO CUMPLE\"))").font = NORMAL
        for j, col in enumerate("LMNO"):
            c = wr.cell(r, 7 + j, f"=SUMIF({rango('A')},A{r},{rango(col)})")
            c.font, c.number_format = NORMAL, MONEDA
        fila_res[it.get("id")] = r
    base = 5 + len(items)
    wr.cell(base, 1, "Por fuente").font = NEGRITA
    encabezado(wr, base + 1, ["Fuente", "", "Monto", "% del total"])
    fuentes_f = {}
    for k, (clave, nombre) in enumerate(FUENTES):
        r = base + 2 + k
        wr.cell(r, 1, nombre).font = NORMAL
        c = wr.cell(r, 3, f"=Presupuesto!{L(12 + k)}{tot}")
        c.font, c.number_format = NORMAL, MONEDA
        c = wr.cell(r, 4, f"=IFERROR(C{r}/{TOTAL},0)")
        c.font, c.number_format = NORMAL, PCT
        fuentes_f[clave] = r
    r_prop = base + 6
    wr.cell(r_prop, 1, "Aporte propio total").font = NEGRITA
    c = wr.cell(r_prop, 3, f"=C{fuentes_f['pecuniario']}+C{fuentes_f['no_pecuniario']}")
    c.font, c.number_format = NEGRITA, MONEDA
    c = wr.cell(r_prop, 4, f"=IFERROR(C{r_prop}/{TOTAL},0)")
    c.font, c.number_format = NEGRITA, PCT
    wr.cell(r_prop + 1, 1, "Costo total").font = NEGRITA
    c = wr.cell(r_prop + 1, 3, f"={TOTAL}")
    c.font, c.number_format = NEGRITA, MONEDA

    # ---------------- Controles ----------------
    wc = wb.create_sheet("Controles")
    wc["A1"] = "Controles contra las bases"
    wc["A1"].font = TITULO
    wc["A2"] = "Se recalculan solos al cambiar parámetros o partidas. NO CUMPLE debe quedar en cero antes de postular."
    encabezado(wc, 3, ["Id", "Control", "Valor", "Límite", "Estado", "Fuente"], [8, 52, 15, 15, 13, 40])

    def estado(cond, limite):
        return f"=IF({limite}=\"\",\"SIN REGLA\",IF({TOTAL}<=0,\"SIN DATOS\",IF({cond},\"CUMPLE\",\"NO CUMPLE\")))"

    reglas = [
        ("CTL-01", "Aporte de la entidad / costo total ≤ máximo", f"=IFERROR(Resumen!C{fuentes_f['entidad']}/{TOTAL},0)", "=IF(Parametros!B4=\"\",\"\",Parametros!B4)", PCT,
         "C{r}<=D{r}+0.0000001"),
        ("CTL-02", "Aporte de la entidad ≤ monto máximo", f"=Resumen!C{fuentes_f['entidad']}", "=IF(Parametros!B5=\"\",\"\",Parametros!B5)", MONEDA,
         "C{r}<=D{r}+0.0000001"),
        ("CTL-03", "Aporte propio / costo total ≥ mínimo", f"=Resumen!D{r_prop}", "=IF(Parametros!B6=\"\",\"\",Parametros!B6)", PCT,
         "C{r}>=D{r}-0.0000001"),
        ("CTL-04", "Aporte pecuniario / aporte propio ≥ mínimo", f"=IFERROR(Resumen!C{fuentes_f['pecuniario']}/Resumen!C{r_prop},0)",
         "=IF(Parametros!B7=\"\",\"\",Parametros!B7)", PCT, "C{r}>=D{r}-0.0000001"),
    ]
    wc.cell(4, 1, "CTL-00").font = NORMAL
    wc.cell(4, 2, "Costo total del presupuesto mayor que cero").font = NORMAL
    c = wc.cell(4, 3, f"={TOTAL}")
    c.font, c.number_format = NORMAL, MONEDA
    wc.cell(4, 4, 0).font = NORMAL
    wc.cell(4, 5, f"=IF(C4>0,\"CUMPLE\",\"NO CUMPLE\")").font = NEGRITA
    wc.cell(4, 6, "Coherencia interna").font = NORMAL
    r = 5
    for cid, texto, valor, limite, formato, cond in reglas:
        wc.cell(r, 1, cid).font = NORMAL
        wc.cell(r, 2, texto).font = NORMAL
        c = wc.cell(r, 3, valor)
        c.font, c.number_format = NORMAL, formato
        c = wc.cell(r, 4, limite)
        c.font, c.number_format = NORMAL, formato
        wc.cell(r, 5, estado(cond.format(r=r), f"D{r}")).font = NEGRITA
        wc.cell(r, 6, "=Parametros!C4").font = NORMAL
        r += 1
    wc.cell(r, 1, "CTL-05").font = NORMAL
    wc.cell(r, 2, "Partidas cuyas fuentes no suman 100%").font = NORMAL
    wc.cell(r, 3, f"=SUMPRODUCT(--(ABS(Presupuesto!K{ini}:K{fin}-1)>0.0001))*({TOTAL}>0)").font = NORMAL
    wc.cell(r, 4, 0).font = NORMAL
    wc.cell(r, 5, f"=IF(C{r}=0,\"CUMPLE\",\"NO CUMPLE\")").font = NEGRITA
    wc.cell(r, 6, "Coherencia interna").font = NORMAL
    r += 1
    for k, it in enumerate(items, 6):
        fr = fila_res[it.get("id")]
        wc.cell(r, 1, f"CTL-{k:02d}").font = NORMAL
        wc.cell(r, 2, f"=\"Ítem \"&Resumen!A{fr}&\" (\"&Resumen!B{fr}&\") / costo total ≤ tope\"").font = NORMAL
        c = wc.cell(r, 3, f"=Resumen!D{fr}")
        c.font, c.number_format = NORMAL, PCT
        c = wc.cell(r, 4, f"=Resumen!E{fr}")
        c.font, c.number_format = NORMAL, PCT
        wc.cell(r, 5, estado(f"C{r}<=D{r}+0.0000001", f"D{r}")).font = NEGRITA
        wc.cell(r, 6, "=Parametros!C4").font = NORMAL
        r += 1
    wc.cell(r + 1, 2, "Controles NO CUMPLE").font = NEGRITA
    wc.cell(r + 1, 3, f"=COUNTIF(E4:E{r - 1},\"NO CUMPLE\")").font = NEGRITA
    wc.conditional_formatting.add(f"E4:E{r - 1}", CellIsRule(operator="equal", formula=['"NO CUMPLE"'], fill=ROJO))
    wc.conditional_formatting.add(f"E4:E{r - 1}", CellIsRule(operator="equal", formula=['"CUMPLE"'], fill=VERDE))

    # ---------------- Cronograma ----------------
    meses = int(s.get("meses") or 0)
    if meses > 0 and n:
        wm = wb.create_sheet("Cronograma")
        wm["A1"] = "Cronograma financiero (distribución mensual de cada partida)"
        wm["A1"].font = TITULO
        wm["A2"] = "Azul = % del costo de la partida que se ejecuta cada mes (debe sumar 100%)."
        encabezado(wm, 3, ["Partida", "Costo total"] + [f"Mes {m}" for m in range(1, meses + 1)] + ["Suma %"],
                   [42, 15] + [9] * meses + [9])
        for i, x in enumerate(partidas):
            rr, rb = 4 + i, ini + i
            wm.cell(rr, 1, f"=Presupuesto!B{rb}").font = NORMAL
            c = wm.cell(rr, 2, f"=Presupuesto!F{rb}")
            c.font, c.number_format = NORMAL, MONEDA
            dist = x.get("distribucion") or [round(1 / meses, 6)] * meses
            for m in range(meses):
                c = wm.cell(rr, 3 + m, num(dist[m]) if m < len(dist) else 0)
                c.font, c.number_format = AZUL, PCT
            c = wm.cell(rr, 3 + meses, f"=SUM({L(3)}{rr}:{L(2 + meses)}{rr})")
            c.font, c.number_format = NORMAL, PCT
        base_m = 6 + n
        wm.cell(base_m, 1, "Montos por mes").font = NEGRITA
        for i in range(n):
            rr, rp = base_m + 1 + i, 4 + i
            wm.cell(rr, 1, f"=A{rp}").font = NORMAL
            for m in range(meses):
                c = wm.cell(rr, 3 + m, f"=$B{rp}*{L(3 + m)}{rp}")
                c.font, c.number_format = NORMAL, MONEDA
        rt = base_m + 1 + n
        wm.cell(rt, 1, "Total mes").font = NEGRITA
        for m in range(meses):
            c = wm.cell(rt, 3 + m, f"=SUM({L(3 + m)}{base_m + 1}:{L(3 + m)}{rt - 1})")
            c.font, c.number_format = NEGRITA, MONEDA

    wb.calculation.fullCalcOnLoad = True
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)
    raiz = Path(a.proyecto) if a.proyecto else next((d for d in ruta.resolve().parents if (d / "proyecto.yaml").exists()), None)
    aut = autoria.cargar_autoria(raiz)
    autoria.limpiar_ooxml(salida, aut.get("nombre", ""), f"Presupuesto {s.get('nombre', '')}".strip(), aut.get("organizacion", ""))

    total, por_fuente, ctrl = calcular(s)
    ruta_ctrl = ruta.with_name("controles.csv")
    with open(ruta_ctrl, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["id", "control", "valor", "limite", "estado", "fuente"])
        w.writeheader()
        for c in ctrl:
            w.writerow({**c, "valor": f"{c['valor']:.10g}" if isinstance(c["valor"], float) else c["valor"],
                        "limite": "" if c["limite"] is None else f"{c['limite']:.10g}"})
    mon = s.get("moneda", "CLP")
    print(f"Presupuesto generado: {salida.as_posix()}")
    print(f"Costo total: {total:,.0f} {mon}".replace(",", "."))
    for k, nombre in FUENTES:
        print(f"- {nombre}: {por_fuente[k]:,.0f} ({por_fuente[k] / total if total else 0:.1%})".replace(",", "."))
    print("\nControles:")
    for c in ctrl:
        v = c["valor"]
        vt = f"{v:.1%}" if isinstance(v, float) and v <= 1.5 and c["id"] not in ("CTL-02",) else (f"{v:,.0f}".replace(",", ".") if isinstance(v, float) else str(v))
        print(f"- {c['id']} {c['estado']}: {c['control']} (valor {vt})")
    malos = sum(1 for c in ctrl if c["estado"] == "NO CUMPLE")
    print(f"\nControles NO CUMPLE: {malos} · detalle en {ruta_ctrl.as_posix()}")


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    main()

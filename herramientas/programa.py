#!/usr/bin/env python3
"""Consolidación de un programa (arquetipo programa): subproyectos, presupuesto, personas, enmiendas y Gantt.

Uso:
    python herramientas/programa.py proyectos/<programa> [--hoy AAAA-MM-DD] [--corte-actualizacion] [--excel]
                                    [--registrar] [--sin-gantt] [--salida estado.md]

Lee programa/subproyectos.csv y, de cada subproyecto, su proyecto.yaml (línea base propia) y su cronograma; lee el
presupuesto del programa (presupuesto/, con cuentas CONTINGENCIA), las personas y asignaciones (programa/recursos/)
y las afectaciones de enmiendas (contractual/afectaciones.csv). Informa:
- avance planificado y real por subproyecto y del programa (ponderado según programa_info.ponderacion_avance),
  término planificado frente al término vigente de cada subproyecto, hitos vencidos y próximos;
- presupuesto por subproyecto y total: base, modificaciones, vigente, comprometido, real, disponible, estimado al
  término y desviación; contingencia (saldo y usos); cuadratura con el contrato y con cada subproyecto;
- personas sobreasignadas (suma de dedicaciones de un mes sobre su capacidad) y dotación por subproyecto;
- afectaciones de enmiendas sin aplicar o que no coinciden con el subproyecto.
Genera (salvo --sin-gantt): cronograma/consolidado.csv (tareas de todos los subproyectos con ids SP.T y dependencias
entre subproyectos "SP2:T05" resueltas), cronograma/gantt-programa.md (macro), cronograma/gantt-detalle.md y el
gantt.md de cada subproyecto. Con --excel: presupuesto/control-presupuestario.xlsx con fórmulas. Con --registrar
guarda las métricas en metricas/mediciones.csv (entidad PROYECTO = programa; entidad <id> = subproyecto).
No modifica los datos de origen.
"""
import argparse
import csv
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import _config  # noqa: E402
import _mediciones  # noqa: E402
import autoria  # noqa: E402

HERR = Path(__file__).resolve().parent
FUENTE = "herramientas/programa.py"
CONT = "CONTINGENCIA"
CAMPOS_CRON = ["id", "subproyecto", "tarea", "responsable", "inicio_plan", "fin_plan", "inicio_real", "fin_real", "avance_pct",
               "predecesoras", "es_hito", "estado", "comentario", "costo_plan", "costo_real"]


def leer(ruta):
    if not Path(ruta).exists():
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return [r for r in csv.DictReader(fh) if any((v or "").strip() for v in r.values() if isinstance(v, str))]


def num(v):
    s = str(v if v is not None else "").strip().replace("$", "").replace(" ", "")
    if not s:
        return 0.0
    pct = s.endswith("%")
    s = s.rstrip("%")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "")
    try:
        x = float(s.replace(",", "."))
    except ValueError:
        return 0.0
    return x / 100 if pct else x


def fch(v):
    try:
        return datetime.strptime(str(v or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def es_si(v):
    return str(v or "").strip().lower() in ("si", "sí", "true", "1", "x", "yes")


def avance(tareas, corte):
    """Avance plan y real ponderado por duración (los hitos no pesan), fechas extremas."""
    peso = plan = real = 0.0
    inicios, fines = [], []
    for t in tareas:
        ini, fin = fch(t.get("inicio_plan")), fch(t.get("fin_plan"))
        if ini:
            inicios.append(ini)
        if fin:
            fines.append(fin)
        if es_si(t.get("es_hito")) or not ini or not fin:
            continue
        d = max((fin - ini).days, 1)
        ap = 0.0 if corte < ini else 100.0 if corte >= fin else 100.0 * (corte - ini).days / max((fin - ini).days, 1)
        peso += d
        plan += d * ap
        real += d * num(t.get("avance_pct"))
    return (plan / peso if peso else 0.0), (real / peso if peso else 0.0), (min(inicios) if inicios else None), (max(fines) if fines else None)


def meses(desde, hasta):
    m = date(desde.year, desde.month, 1)
    while m <= hasta:
        yield m
        m = date(m.year + (m.month // 12), m.month % 12 + 1, 1)


def fin_de_mes(m):
    return date(m.year + (m.month // 12), m.month % 12 + 1, 1) - timedelta(days=1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("programa")
    ap.add_argument("--hoy")
    ap.add_argument("--corte-actualizacion", action="store_true", help="avance plan de cada subproyecto a su cronograma_actualizado")
    ap.add_argument("--excel", action="store_true", help="generar presupuesto/control-presupuestario.xlsx")
    ap.add_argument("--registrar", action="store_true")
    ap.add_argument("--sin-gantt", action="store_true")
    ap.add_argument("--salida")
    a = ap.parse_args()
    raiz = Path(a.programa)
    hoy = fch(a.hoy) if a.hoy else date.today()
    if hoy is None:
        sys.exit(f"Fecha de corte inválida: {a.hoy} (usar AAAA-MM-DD).")
    cfg = _config.cargar(raiz)
    if "programa" not in (cfg.get("arquetipos") or []):
        sys.exit(f"{raiz.as_posix()} no es un programa (falta el arquetipo programa en proyecto.yaml).")
    pinfo = cfg.get("programa_info") or {}
    moneda = pinfo.get("moneda") or (cfg.get("linea_base") or {}).get("moneda") or "CLP"
    um = cfg.get("umbrales") or {}
    amarillo, rojo = num(um.get("desvio_amarillo_pct", 5)) or 5, num(um.get("desvio_rojo_pct", 10)) or 10
    crit, adv, info = [], [], []
    m = {}

    # ---------------- Subproyectos ----------------
    subs = []
    for s in leer(raiz / "programa" / "subproyectos.csv"):
        carpeta = raiz / (s.get("carpeta") or f"subproyectos/{s.get('id')}")
        if not (carpeta / "proyecto.yaml").exists():
            crit.append(f"Subproyecto {s.get('id')}: no existe {carpeta.as_posix()}/proyecto.yaml.")
            continue
        scfg = _config.cargar(carpeta)
        tareas = leer(carpeta / "cronograma" / "cronograma.csv")
        nom_cfg = str((scfg.get("proyecto") or {}).get("nombre") or "")
        nom = nom_cfg if nom_cfg and "PENDIENTE" not in nom_cfg.upper() else (s.get("nombre") or s["id"])
        subs.append({"id": s["id"], "nombre": nom,
                     "carpeta": carpeta, "cfg": scfg, "tareas": tareas, "peso": num(s.get("peso")), "estado": s.get("estado", "")})
    if not subs:
        adv.append("El programa no tiene subproyectos registrados en programa/subproyectos.csv.")
    ids_sub = {s["id"] for s in subs}

    # ---------------- Presupuesto ----------------
    base = leer(raiz / "presupuesto" / "presupuesto.csv")
    mods = leer(raiz / "presupuesto" / "modificaciones.csv")
    gastos = leer(raiz / "presupuesto" / "gastos.csv")
    pron = leer(raiz / "presupuesto" / "pronostico.csv")
    claves = []
    datos = {}

    def reg(sp, cta):
        k = (sp.strip(), cta.strip().upper() if cta.strip().upper() == CONT else cta.strip())
        if k not in datos:
            datos[k] = {"base": 0.0, "mods": 0.0, "mods_enm": 0.0, "comp": 0.0, "real": 0.0, "por_gastar": None, "desc": ""}
            claves.append(k)
        return datos[k]

    for r in base:
        d = reg(r.get("subproyecto", ""), r.get("cuenta", ""))
        d["base"] += num(r.get("monto_base"))
        d["desc"] = d["desc"] or r.get("descripcion", "")
    for r in mods:
        d = reg(r.get("subproyecto", ""), r.get("cuenta", ""))
        d["mods"] += num(r.get("monto"))
        if (r.get("tipo") or "").strip().lower() == "enmienda":
            d["mods_enm"] += num(r.get("monto"))
    for r in gastos:
        k_sp, k_cta = r.get("subproyecto", "").strip(), r.get("cuenta", "").strip()
        if (k_sp, k_cta) not in datos and (k_sp, k_cta.upper()) not in datos:
            adv.append(f"Gasto de {r.get('fecha')} ({r.get('documento') or 'sin documento'}) en {k_sp}/{k_cta}, cuenta sin presupuesto.")
        tipo = (r.get("tipo") or "").strip().lower()
        d = reg(k_sp, k_cta)
        if tipo.startswith("compromet"):
            d["comp"] += num(r.get("monto"))
        elif tipo == "real":
            d["real"] += num(r.get("monto"))
        else:
            adv.append(f"Gasto de {r.get('fecha')} con tipo '{r.get('tipo')}': usar comprometido o real.")
    ultimo_pron = {}
    for r in sorted(pron, key=lambda x: x.get("fecha") or ""):
        ultimo_pron[(r.get("subproyecto", "").strip(), r.get("cuenta", "").strip())] = num(r.get("por_gastar"))
    for k, v in ultimo_pron.items():
        reg(*k)["por_gastar"] = v
    # coherencia de modificaciones
    por_id = {}
    for r in mods:
        tipo = (r.get("tipo") or "").strip().lower()
        if tipo in ("reasignacion", "reasignación", "contingencia"):
            por_id.setdefault(r.get("id"), 0.0)
            por_id[r.get("id")] += num(r.get("monto"))
        elif tipo == "enmienda" and not (r.get("documento") or "").strip():
            adv.append(f"Modificación {r.get('id')} tipo enmienda sin documento.")
        if not (r.get("aprobado_por") or "").strip():
            adv.append(f"Modificación {r.get('id')} ({r.get('tipo')}) sin aprobación registrada.")
    for mid, suma in por_id.items():
        if abs(suma) > 0.5:
            adv.append(f"Movimiento {mid} (reasignación o uso de contingencia) no suma cero ({suma:,.0f}): revisa sus filas.".replace(",", "."))
    for (sp, _), _d in datos.items():
        if sp and sp != "PROGRAMA" and sp not in ids_sub:
            adv.append(f"Presupuesto con subproyecto desconocido '{sp}'.")
    # por clave
    for k, d in datos.items():
        d["vig"] = d["base"] + d["mods"]
        d["disp"] = d["vig"] - d["comp"] - d["real"]
        d["eac"] = (d["real"] + d["comp"] + d["por_gastar"]) if d["por_gastar"] is not None else max(d["vig"], d["real"] + d["comp"])
        if k[1] == CONT:
            d["eac"] = d["real"] + d["comp"]  # la contingencia no se "gasta" por sí sola
        d["vac"] = d["vig"] - d["eac"]
        if d["disp"] < -0.5 and k[1] != CONT:
            crit.append(f"Sobregiro en {k[0]}/{k[1]}: comprometido + real supera el vigente en {-d['disp']:,.0f} {moneda}.".replace(",", "."))
    grupos = ["PROGRAMA"] + [s["id"] for s in subs]
    tot = {}
    for g in grupos + ["TOTAL"]:
        sel = [d for k, d in datos.items() if g == "TOTAL" or k[0] == g]
        tot[g] = {c: sum(d[c] for d in sel) for c in ("base", "mods", "vig", "comp", "real", "disp", "eac", "vac")}
        tot[g]["cont_base"] = sum(d["base"] + d["mods_enm"] for k, d in datos.items() if (g == "TOTAL" or k[0] == g) and k[1] == CONT)
        tot[g]["cont_vig"] = sum(d["vig"] for k, d in datos.items() if (g == "TOTAL" or k[0] == g) and k[1] == CONT)
    T = tot["TOTAL"]
    m["presupuesto_vigente_programa"] = T["vig"]
    m["ejecucion_programa"] = (T["comp"] + T["real"]) / T["vig"] if T["vig"] else 0.0
    m["vac_programa"] = T["vac"]
    m["contingencia_saldo_pct"] = T["cont_vig"] / T["cont_base"] if T["cont_base"] else None
    alerta_cont = num(pinfo.get("contingencia_alerta_pct", 0.25)) or 0.25
    if m["contingencia_saldo_pct"] is not None and m["contingencia_saldo_pct"] < alerta_cont:
        crit.append(f"Contingencia del programa bajo la alerta: queda {m['contingencia_saldo_pct']:.0%} de la inicial.")
    if T["vac"] < -0.5:
        crit.append(f"El estimado al término supera el presupuesto vigente en {-T['vac']:,.0f} {moneda}.".replace(",", "."))
    monto_ctto = num((cfg.get("linea_base") or {}).get("monto_vigente"))
    if base and monto_ctto and abs(monto_ctto - T["vig"]) > 0.5:
        adv.append(f"El presupuesto vigente del programa ({T['vig']:,.0f}) no cuadra con el monto vigente del contrato "
                   f"({monto_ctto:,.0f}).".replace(",", "."))
    for s in subs:
        lb = num((s["cfg"].get("linea_base") or {}).get("monto_vigente"))
        if base and lb and abs(lb - tot[s["id"]]["vig"]) > 0.5:
            adv.append(f"{s['id']}: el monto vigente de su línea base ({lb:,.0f}) no cuadra con su presupuesto vigente "
                       f"({tot[s['id']]['vig']:,.0f}).".replace(",", "."))

    # ---------------- Avance y términos ----------------
    tabla = []
    atrasados = 0
    hitos_prox = []
    for s in subs:
        corte = hoy
        act = fch(s["cfg"].get("cronograma_actualizado"))
        if act and act < hoy:
            if a.corte_actualizacion:
                corte = act
            else:
                adv.append(f"{s['id']}: avance real no actualizado desde el {act} (usa --corte-actualizacion para comparar a esa fecha).")
        plan, real, ini, fin = avance(s["tareas"], corte)
        termino = fch((s["cfg"].get("linea_base") or {}).get("fecha_termino_vigente"))
        desv = real - plan
        estado = "Rojo" if desv <= -rojo else "Amarillo" if desv <= -amarillo else "Verde"
        if termino and fin and fin > termino:
            atrasados += 1
            estado = "Rojo"
            crit.append(f"{s['id']} ({s['nombre']}) termina el {fin}, después de su término vigente ({termino}).")
        for t in s["tareas"]:
            ff = fch(t.get("fin_plan"))
            if es_si(t.get("es_hito")) and ff and num(t.get("avance_pct")) < 100:
                if ff < hoy:
                    crit.append(f"Hito vencido {s['id']}:{t.get('id')} '{t.get('tarea')}' ({ff}).")
                elif (ff - hoy).days <= 60:
                    hitos_prox.append((ff, f"{s['id']}: {t.get('tarea')}"))
        s.update(plan=plan, real=real, desv=desv, ini=ini, fin=fin, termino=termino, estado=estado)
        tabla.append(s)
    m["subproyectos_atrasados"] = atrasados
    pond = str(pinfo.get("ponderacion_avance") or "presupuesto").lower()

    def peso(s):
        if pond == "peso" and s["peso"]:
            return s["peso"]
        if pond == "duracion" and s["ini"] and s["fin"]:
            return max((s["fin"] - s["ini"]).days, 1)
        return tot[s["id"]]["vig"] if tot.get(s["id"], {}).get("vig") else 1.0

    pesos = {s["id"]: peso(s) for s in tabla}
    ptot = sum(pesos.values()) or 1
    m["avance_plan_programa"] = sum(s["plan"] * pesos[s["id"]] for s in tabla) / ptot / 100 if tabla else 0.0
    m["avance_real_programa"] = sum(s["real"] * pesos[s["id"]] for s in tabla) / ptot / 100 if tabla else 0.0
    m["desviacion_avance_programa"] = m["avance_real_programa"] - m["avance_plan_programa"]

    # ---------------- Cronograma consolidado ----------------
    consolidado = []
    prefijos = {}
    prog_tareas = leer(raiz / "cronograma" / "cronograma.csv")
    fuentes = [("PRG", "Programa", prog_tareas)] + [(s["id"], f"{s['id']} · {s['nombre']}", s["tareas"]) for s in subs]
    for pref, nombre, tareas in fuentes:
        prefijos[pref] = {t.get("id") for t in tareas}
    for pref, nombre, tareas in fuentes:
        for t in tareas:
            fila = {k: t.get(k, "") for k in CAMPOS_CRON}
            fila["id"] = f"{pref}.{t.get('id')}"
            fila["subproyecto"] = nombre
            preds = []
            for p in [x.strip() for x in re.split(r"[;,]", t.get("predecesoras") or "") if x.strip()]:
                if ":" in p:
                    sp, tid = p.split(":", 1)
                    sp = "PRG" if sp.upper() in ("PRG", "PROGRAMA") else sp
                    if tid not in prefijos.get(sp, set()):
                        crit.append(f"{pref}:{t.get('id')} depende de '{p}', que no existe.")
                        continue
                    preds.append(f"{sp}.{tid}")
                else:
                    preds.append(f"{pref}.{p}")
            fila["predecesoras"] = ";".join(preds)
            consolidado.append(fila)
    por_cid = {f["id"]: f for f in consolidado}
    for f_ in consolidado:
        ini = fch(f_["inicio_plan"])
        for p in [x for x in f_["predecesoras"].split(";") if x]:
            if p.split(".")[0] != f_["id"].split(".")[0] and p in por_cid:
                fin_p = fch(por_cid[p]["fin_plan"])
                if ini and fin_p and ini < fin_p:
                    adv.append(f"Dependencia entre subproyectos: {f_['id']} comienza el {ini}, antes de que termine {p} ({fin_p}).")

    if not a.sin_gantt:
        cron_dir = raiz / "cronograma"
        cron_dir.mkdir(exist_ok=True)
        with open(cron_dir / "consolidado.csv", "w", encoding="utf-8", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=CAMPOS_CRON)
            w.writeheader()
            w.writerows(consolidado)
        lineas = ["# Carta Gantt macro del programa", "",
                  f"Generada el {hoy.isoformat()} con `programa.py`. No editar a mano: una barra por subproyecto y los hitos contractuales.",
                  "", "```mermaid", "gantt", "    dateFormat YYYY-MM-DD", "    axisFormat %b %y", "    todayMarker on",
                  "    section Subproyectos"]
        for s in tabla:
            if s["ini"] and s["fin"]:
                marcas = ["done"] if s["real"] >= 99.5 else (["crit"] if s["estado"] == "Rojo" else (["active"] if s["ini"] <= hoy else []))
                nombre = re.sub(r"[:#;]", " ", f"{s['id']} {s['nombre']} ({s['real']:.0f}%)")
                cid = re.sub(r"\W", "", s["id"])
                lineas.append(f"    {nombre} :{', '.join(marcas + [cid])}, {s['ini']}, {s['fin']}")
        lineas.append("    section Términos vigentes")
        for s in tabla:
            if s["termino"]:
                lineas.append(f"    Término {re.sub(r'[:#;]', ' ', s['id'])} :milestone, {s['termino']}, 0d")
        termino_prog = fch((cfg.get("linea_base") or {}).get("fecha_termino_vigente"))
        if termino_prog:
            lineas.append(f"    Término del programa :milestone, crit, {termino_prog}, 0d")
        hitos_prg = [t for t in prog_tareas if es_si(t.get("es_hito")) and fch(t.get("fin_plan"))]
        if hitos_prg:
            lineas.append("    section Hitos del programa")
            lineas += [f"    {re.sub(r'[:#;]', ' ', t.get('tarea', ''))} :milestone, {t.get('fin_plan')}, 0d" for t in hitos_prg]
        lineas += ["```", ""]
        (cron_dir / "gantt-programa.md").write_text("\n".join(lineas), encoding="utf-8")
        py = sys.executable
        subprocess.run([py, str(HERR / "gantt_mermaid.py"), str(raiz), "--hoy", hoy.isoformat(), "--cronograma", str(cron_dir / "consolidado.csv"),
                        "--salida", str(cron_dir / "gantt-detalle.md"), "--titulo", "Carta Gantt detallada del programa"],
                       capture_output=True, text=True)
        for s in subs:
            if s["tareas"]:
                subprocess.run([py, str(HERR / "gantt_mermaid.py"), str(s["carpeta"]), "--hoy", hoy.isoformat(),
                                "--titulo", f"Carta Gantt {s['id']} · {s['nombre']}"], capture_output=True, text=True)
        info.append("Generados: cronograma/consolidado.csv, gantt-programa.md (macro), gantt-detalle.md y el gantt.md de cada subproyecto.")
        r = subprocess.run([py, str(HERR / "evm_cpm.py"), str(raiz), "--fecha", hoy.isoformat(), "--cronograma",
                            str(cron_dir / "consolidado.csv")], capture_output=True, text=True, encoding="utf-8", errors="replace")
        pron = re.search(r"Término pronosticado con el avance real: (\S+)", r.stdout or "")
        desv = re.search(r"desviación pronosticada: ([+-]?\d+) días", r.stdout or "")
        ruta_c = re.search(r"## Ruta crítica \(plan\)\s+(.+)", r.stdout or "")
        if pron:
            texto_rc = f"Programa completo: término pronosticado {pron.group(1)}"
            if desv:
                texto_rc += f" ({int(desv.group(1)):+d} días frente al término del contrato)"
                if int(desv.group(1)) > 0:
                    crit.append(texto_rc + ".")
            if ruta_c:
                texto_rc += f"; ruta crítica: {ruta_c.group(1).strip()}"
            info.append(texto_rc + ".")

    # ---------------- Personas ----------------
    personas = {p.get("id"): p for p in leer(raiz / "programa" / "recursos" / "personas.csv")}
    asig = leer(raiz / "programa" / "recursos" / "asignaciones.csv")
    termino_prog = fch((cfg.get("linea_base") or {}).get("fecha_termino_vigente")) or (hoy + timedelta(days=365))
    carga = {}
    dotacion = {s["id"]: 0.0 for s in subs}
    for r in asig:
        pid, sp = r.get("persona"), r.get("subproyecto")
        if pid not in personas:
            adv.append(f"Asignación de '{pid}' a {sp}: la persona no está en personas.csv.")
        if sp not in ids_sub and sp != "PROGRAMA":
            adv.append(f"Asignación de '{pid}' a subproyecto desconocido '{sp}'.")
        desde, hasta = fch(r.get("desde")) or hoy, fch(r.get("hasta")) or termino_prog
        ded = num(r.get("dedicacion_pct"))  # 50, "50%" o 0.5 significan media jornada
        ded = ded * 100 if 0 < ded <= 1 else ded
        for mm in meses(desde, hasta):
            carga.setdefault(pid, {}).setdefault(mm, 0.0)
            carga[pid][mm] += ded
        if desde <= hoy <= hasta and sp in dotacion:
            dotacion[sp] += ded / 100
    sobre = []
    for pid, porm in carga.items():
        cap = num((personas.get(pid) or {}).get("capacidad_pct")) or 100
        exceso = [(mm, v) for mm, v in sorted(porm.items()) if v > cap + 0.01 and fin_de_mes(mm) >= hoy]
        if exceso:
            nombre = (personas.get(pid) or {}).get("nombre", pid)
            desde_m, hasta_m = exceso[0][0], exceso[-1][0]
            pico = max(v for _, v in exceso)
            sobre.append(f"{nombre}: hasta {pico:.0f}% (capacidad {cap:.0f}%) entre {desde_m:%Y-%m} y {hasta_m:%Y-%m}")
    m["personas_sobreasignadas"] = len(sobre)
    for s_ in sobre:
        adv.append(f"Sobreasignación: {s_}.")

    # ---------------- Afectaciones de enmiendas ----------------
    afect = leer(raiz / "contractual" / "afectaciones.csv")
    pend = [r for r in afect if not es_si(r.get("aplicado"))]
    m["afectaciones_pendientes"] = len(pend)
    for r in pend:
        adv.append(f"Enmienda {r.get('documento')} ({r.get('tipo')}) sin aplicar en {r.get('subproyecto')}: {r.get('descripcion') or r.get('despues')}.")
    for r in afect:
        if es_si(r.get("aplicado")) and (r.get("tipo") or "").strip().lower() == "plazo" and r.get("subproyecto") in ids_sub:
            s = next(x for x in subs if x["id"] == r["subproyecto"])
            vig = str((s["cfg"].get("linea_base") or {}).get("fecha_termino_vigente") or "")
            ultima = max((x for x in afect if x.get("subproyecto") == r["subproyecto"] and (x.get("tipo") or "").lower() == "plazo"),
                         key=lambda x: x.get("fecha_vigencia") or "")
            if ultima is r and vig != str(r.get("despues")):
                adv.append(f"{r['subproyecto']}: la enmienda {r.get('documento')} fija término {r.get('despues')}, pero su línea base dice {vig}.")
        if es_si(r.get("aplicado")) and (r.get("tipo") or "").strip().lower() == "monto":
            if not any((x.get("documento") or "") == r.get("documento") and x.get("subproyecto") == r.get("subproyecto") for x in mods):
                adv.append(f"Enmienda {r.get('documento')} cambia el monto de {r.get('subproyecto')}, pero no hay modificación de presupuesto con ese documento.")

    # ---------------- Informe ----------------
    fm = lambda v: f"{v:,.0f}".replace(",", ".")  # noqa: E731
    nombre = (cfg.get("proyecto") or {}).get("nombre", raiz.name)
    out = [f"# Estado del programa — {nombre}", "", f"Fecha de corte: {hoy} · subproyectos: {len(subs)} · ponderación del avance: {pond}", "",
           "## Avance y plazos por subproyecto", "",
           "| Subproyecto | Plan % | Real % | Desv. (pts) | Fin planificado | Término vigente | Estado |", "|---|---|---|---|---|---|---|"]
    for s in tabla:
        out.append(f"| {s['id']} · {s['nombre']} | {s['plan']:.0f} | {s['real']:.0f} | {s['desv']:+.0f} | {s['fin'] or '—'} | {s['termino'] or '—'} | {s['estado']} |")
    out.append(f"| **Programa** | {m['avance_plan_programa']:.0%} | {m['avance_real_programa']:.0%} | "
               f"{m['desviacion_avance_programa'] * 100:+.0f} | | {(cfg.get('linea_base') or {}).get('fecha_termino_vigente', '—')} | |")
    if datos:
        out += ["", f"## Presupuesto ({moneda})", "", "| Grupo | Vigente | Comprometido | Real | Disponible | Ejecutado | Estimado al término | Desviación | Contingencia (saldo/base) |",
                "|---|---|---|---|---|---|---|---|---|"]
        for g in grupos + ["TOTAL"]:
            x = tot[g]
            if not x["vig"] and g != "TOTAL":
                continue
            ej = (x["comp"] + x["real"]) / x["vig"] if x["vig"] else 0
            cont = f"{fm(x['cont_vig'])} / {fm(x['cont_base'])}" if x["cont_base"] else "—"
            etiqueta = "**Total programa**" if g == "TOTAL" else g
            out.append(f"| {etiqueta} | {fm(x['vig'])} | {fm(x['comp'])} | {fm(x['real'])} | {fm(x['disp'])} | {ej:.0%} | {fm(x['eac'])} | {fm(x['vac'])} | {cont} |")
    if subs:
        out += ["", "## Equipo", "", "| Subproyecto | Dotación hoy (jornadas completas equivalentes) |", "|---|---|"]
        out += [f"| {sid} | {v:.1f} |" for sid, v in dotacion.items()]
        if sobre:
            out += ["", "Personas sobreasignadas:"] + [f"- {x}" for x in sobre]
    if hitos_prox:
        out += ["", "## Hitos en los próximos 60 días", ""] + [f"- {d}: {t}" for d, t in sorted(hitos_prox)]
    for titulo, lista in (("Crítico", crit), ("Advertencias", adv), ("Información", info)):
        out += ["", f"## {titulo} ({len(lista)})", ""] + ([f"- {x}" for x in dict.fromkeys(lista)] or ["- Sin hallazgos."])
    texto = "\n".join(out) + "\n"

    if a.excel and datos:
        ruta = excel(raiz, claves, datos, subs, base, mods, gastos, moneda, cfg)
        texto += f"\nControl presupuestario: {ruta.as_posix()}\n"
    if a.registrar:
        nuevas = [{"fecha": hoy.isoformat(), "entidad": "PROYECTO", "metrica": k, "valor": f"{v:.10g}", "fuente": FUENTE, "comentario": ""}
                  for k, v in m.items() if v is not None]
        for s in tabla:
            x = tot[s["id"]]
            for k, v in (("avance_plan", s["plan"] / 100), ("avance_real", s["real"] / 100), ("desviacion_avance", s["desv"] / 100),
                         ("presupuesto_vigente", x["vig"]), ("ejecucion_presupuesto", (x["comp"] + x["real"]) / x["vig"] if x["vig"] else 0),
                         ("vac", x["vac"])):
                nuevas.append({"fecha": hoy.isoformat(), "entidad": s["id"], "metrica": k, "valor": f"{v:.10g}", "fuente": FUENTE, "comentario": ""})
        _mediciones.registrar(raiz / "metricas" / "mediciones.csv", nuevas, FUENTE, hoy.isoformat())
    else:
        texto += "\n(Métricas no registradas: usar --registrar para guardarlas en metricas/mediciones.csv.)\n"
    if a.salida:
        Path(a.salida).write_text(texto, encoding="utf-8")
    print(texto)


def excel(raiz, claves, datos, subs, base, mods, gastos, moneda, cfg):
    """Control presupuestario con fórmulas: hojas de datos (copia de los CSV) y hojas de cálculo que las resumen."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    N, NB, AZ = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10, color="0000FF")
    TIT, ENC = Font(name="Arial", size=13, bold=True, color="1F3864"), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    FON, ROJO = PatternFill("solid", fgColor="1F3864"), PatternFill("solid", fgColor="FFC7CE")
    MON, PCT = "#,##0;[Red]-#,##0", "0.0%"
    wb = Workbook()

    def hoja_datos(nombre, filas, campos, montos):
        ws = wb.create_sheet(nombre)
        for j, c in enumerate(campos, 1):
            x = ws.cell(1, j, c)
            x.font, x.fill = ENC, FON
            ws.column_dimensions[L(j)].width = 16 if c not in ("descripcion", "motivo") else 40
        for i, r in enumerate(filas, 2):
            for j, c in enumerate(campos, 1):
                v = num(r.get(c)) if c in montos else r.get(c, "")
                x = ws.cell(i, j, v)
                x.font = N
                if c in montos:
                    x.number_format = MON
        return max(len(filas) + 1, 2)

    wr = wb.active
    wr.title = "Resumen"
    np_ = hoja_datos("Presupuesto", base, ["subproyecto", "cuenta", "descripcion", "monto_base", "documento"], {"monto_base"})
    nm = hoja_datos("Modificaciones", mods, ["id", "fecha", "tipo", "documento", "subproyecto", "cuenta", "monto", "motivo", "aprobado_por"], {"monto"})
    ng = hoja_datos("Gastos", gastos, ["fecha", "subproyecto", "cuenta", "tipo", "monto", "documento", "proveedor", "descripcion"], {"monto"})

    wd = wb.create_sheet("Detalle", 1)
    wd["A1"] = f"Detalle por subproyecto y cuenta ({moneda})"
    wd["A1"].font = TIT
    wd["A2"] = "Fórmulas sobre las hojas de datos. Azul = por gastar estimado (editable para simular). La fuente son los CSV de presupuesto/."
    cab = ["Subproyecto", "Cuenta", "Base", "Modificaciones", "Vigente", "Comprometido", "Real", "Disponible", "Ejecutado", "Por gastar", "Estimado al término", "Desviación"]
    for j, c in enumerate(cab, 1):
        x = wd.cell(3, j, c)
        x.font, x.fill = ENC, FON
        wd.column_dimensions[L(j)].width = 15
    orden = sorted(claves, key=lambda k: (k[0] != "PROGRAMA", k[0], k[1] == CONT, k[1]))
    for i, (sp, cta) in enumerate(orden, 4):
        d = datos[(sp, cta)]
        wd.cell(i, 1, sp).font = N
        wd.cell(i, 2, cta).font = N
        f = {3: f"=SUMIFS(Presupuesto!$D$2:$D${np_},Presupuesto!$A$2:$A${np_},A{i},Presupuesto!$B$2:$B${np_},B{i})",
             4: f"=SUMIFS(Modificaciones!$G$2:$G${nm},Modificaciones!$E$2:$E${nm},A{i},Modificaciones!$F$2:$F${nm},B{i})",
             5: f"=C{i}+D{i}",
             6: f"=SUMIFS(Gastos!$E$2:$E${ng},Gastos!$B$2:$B${ng},A{i},Gastos!$C$2:$C${ng},B{i},Gastos!$D$2:$D${ng},\"comprometido\")",
             7: f"=SUMIFS(Gastos!$E$2:$E${ng},Gastos!$B$2:$B${ng},A{i},Gastos!$C$2:$C${ng},B{i},Gastos!$D$2:$D${ng},\"real\")",
             8: f"=E{i}-F{i}-G{i}", 9: f"=IFERROR((F{i}+G{i})/E{i},0)",
             11: (f"=F{i}+G{i}" if cta == CONT else f"=IF(J{i}=\"\",MAX(E{i},F{i}+G{i}),F{i}+G{i}+J{i})"), 12: f"=E{i}-K{i}"}
        for col, formula in f.items():
            x = wd.cell(i, col, formula)
            x.font, x.number_format = N, (PCT if col == 9 else MON)
        x = wd.cell(i, 10, d["por_gastar"] if d["por_gastar"] is not None else None)
        x.font, x.number_format = AZ, MON
    ult = 3 + len(orden)
    wd.conditional_formatting.add(f"H4:H{ult}", CellIsRule(operator="lessThan", formula=["0"], fill=ROJO))
    wd.conditional_formatting.add(f"L4:L{ult}", CellIsRule(operator="lessThan", formula=["0"], fill=ROJO))

    wr["A1"] = f"Control presupuestario — {(cfg.get('proyecto') or {}).get('nombre', '')} ({moneda})"
    wr["A1"].font = TIT
    wr["A2"] = "Resumen por subproyecto calculado con fórmulas desde la hoja Detalle."
    cab = ["Subproyecto", "Vigente", "Comprometido", "Real", "Disponible", "Ejecutado", "Estimado al término", "Desviación",
           "Contingencia vigente", "Contingencia base"]
    for j, c in enumerate(cab, 1):
        x = wr.cell(3, j, c)
        x.font, x.fill = ENC, FON
        wr.column_dimensions[L(j)].width = 17
    grupos = ["PROGRAMA"] + [s["id"] for s in subs]
    for i, g in enumerate(grupos, 4):
        wr.cell(i, 1, g).font = N
        rng = lambda col: f"Detalle!${col}$4:${col}${ult}"  # noqa: E731
        f = {2: f"=SUMIFS({rng('E')},{rng('A')},A{i})", 3: f"=SUMIFS({rng('F')},{rng('A')},A{i})", 4: f"=SUMIFS({rng('G')},{rng('A')},A{i})",
             5: f"=SUMIFS({rng('H')},{rng('A')},A{i})", 6: f"=IFERROR((C{i}+D{i})/B{i},0)", 7: f"=SUMIFS({rng('K')},{rng('A')},A{i})",
             8: f"=B{i}-G{i}", 9: f"=SUMIFS({rng('E')},{rng('A')},A{i},{rng('B')},\"{CONT}\")",
             10: (f"=SUMIFS({rng('C')},{rng('A')},A{i},{rng('B')},\"{CONT}\")"
                  f"+SUMIFS(Modificaciones!$G$2:$G${nm},Modificaciones!$E$2:$E${nm},A{i},Modificaciones!$F$2:$F${nm},\"{CONT}\","
                  f"Modificaciones!$C$2:$C${nm},\"enmienda\")")}
        for col, formula in f.items():
            x = wr.cell(i, col, formula)
            x.font, x.number_format = N, (PCT if col == 6 else MON)
    t = 4 + len(grupos)
    wr.cell(t, 1, "Total programa").font = NB
    for col in range(2, 11):
        if col == 6:
            x = wr.cell(t, col, f"=IFERROR((C{t}+D{t})/B{t},0)")
            x.number_format = PCT
        else:
            x = wr.cell(t, col, f"=SUM({L(col)}4:{L(col)}{t - 1})")
            x.number_format = MON
        x.font = NB
    wr.conditional_formatting.add(f"E4:E{t}", CellIsRule(operator="lessThan", formula=["0"], fill=ROJO))
    wr.conditional_formatting.add(f"H4:H{t}", CellIsRule(operator="lessThan", formula=["0"], fill=ROJO))
    wb.calculation.fullCalcOnLoad = True
    ruta = raiz / "presupuesto" / "control-presupuestario.xlsx"
    ruta.parent.mkdir(exist_ok=True)
    wb.save(ruta)
    aut = autoria.cargar_autoria(raiz)
    autoria.limpiar_ooxml(ruta, aut.get("nombre", ""), f"Control presupuestario {(cfg.get('proyecto') or {}).get('nombre', '')}".strip(),
                          aut.get("organizacion", ""))
    return ruta


if __name__ == "__main__":
    import _consola
    _consola.preparar()
    main()

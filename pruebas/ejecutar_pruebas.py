#!/usr/bin/env python3
"""Batería de pruebas del harness. Ejecutar después de cualquier cambio en skills, scripts o arquetipos.

Uso:
    python pruebas/ejecutar_pruebas.py [--detalle]

Trabaja sobre una copia temporal del repositorio (no modifica nada) y verifica que cada herramienta
produzca los resultados esperados sobre los ejemplos ficticios (incluidos los errores sembrados a propósito).
Termina con código 0 si todo pasa y 1 si algo falla.
"""
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
PY = sys.executable
for _flujo in (sys.stdout, sys.stderr):  # la consola de Windows (cp1252) no debe caerse al mostrar el detalle de una falla
    try:
        _flujo.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass
resultados = []
DETALLE = "--detalle" in sys.argv


def correr(args, cwd, entrada=None, env=None):
    r = subprocess.run([PY] + args, cwd=cwd, capture_output=True, text=True, input=entrada, timeout=300,
                       encoding="utf-8", errors="replace", env={**os.environ, **(env or {})})
    return r.returncode, (r.stdout or "") + (r.stderr or "")


def prueba(nombre, args, cwd, esperados=(), ausentes=(), codigo=0, entrada=None, env=None):
    rc, salida = correr(args, cwd, entrada, env)
    fallas = []
    if codigo is not None and rc != codigo:
        fallas.append(f"código {rc} (esperado {codigo})")
    fallas += [f"falta: {e!r}" for e in esperados if e not in salida]
    fallas += [f"sobra: {e!r}" for e in ausentes if e in salida]
    resultados.append((nombre, not fallas, fallas))
    if DETALLE or fallas:
        print(f"\n--- {nombre} ---\n{salida[:3000]}")
    return salida


def main():
    tmp = Path(tempfile.mkdtemp(prefix="pmo-pruebas-"))
    repo = tmp / "repo"
    largo = (lambda p: "\\\\?\\" + str(Path(p).resolve())) if os.name == "nt" else str  # rutas de más de 260 caracteres en Windows
    comunes = shutil.ignore_patterns(".git", "__pycache__", "_procesado", ".venv", "respaldo-nucleo")

    def solo_nucleo(carpeta, nombres):
        """En una carpeta de empresa, las pruebas usan solo el núcleo: sin su empresa, su registro ni sus proyectos."""
        ignorar = set(comunes(carpeta, nombres))
        c = Path(carpeta.replace("\\\\?\\", "", 1)).resolve()
        if c == RAIZ.resolve():
            ignorar |= {"empresa", ".harness"} & set(nombres)
        elif c == (RAIZ / "proyectos").resolve():
            ignorar |= set(nombres)
        return ignorar

    shutil.copytree(largo(RAIZ), largo(repo), ignore=solo_nucleo)
    ficha0 = repo / "FICHA-DE-INICIO.md"
    if ficha0.exists():  # la ficha de la empresa puede tener respuestas: las pruebas parten de una en blanco
        ficha0.write_text(re.sub(r"^(\s*-\s+[^:\n]+:)[ \t]*\S.*$", r"\1", ficha0.read_text(encoding="utf-8"), flags=re.M), encoding="utf-8")
    H = "herramientas/"
    try:
        prueba("Compatibilidad Python 3.9+", ["pruebas/compatibilidad_python.py"], repo, ["OK"], codigo=None)

        prueba("validar_proyecto: demo ejecución", [H + "validar_proyecto.py", "ejemplos/demo-andes", "--hoy", "2026-09-25"], repo,
               ["T13", "después del término contractual", "R-004", "no tiene plan de acción", "A-001 vencido", "| SP1 ERP |"])

        prueba("evm_cpm: valor ganado y ruta crítica", [H + "evm_cpm.py", "ejemplos/demo-andes", "--fecha", "2026-09-25"], repo,
               ["| SPI (EV/PV) | 0.79 |", "| CPI (EV/AC) | 0.87 |", "desviación pronosticada: +29 días", "T12 → T13"])

        prueba("calcular_cartera: script propio del portafolio",
               ["ejemplos/portafolio-acciones/scripts/calcular_cartera.py", "ejemplos/portafolio-acciones", "--fecha", "2026-08-31"],
               repo, ["61.589.330", "TCOR 24.9%"])

        prueba("evaluar_metricas: portafolio", [H + "evaluar_metricas.py", "ejemplos/portafolio-acciones", "--hoy", "2026-08-31"], repo,
               ["POL-01 incumplida", "TCOR tiene peso = 24.9%", "POL-04 incumplida", "PNOR", "falta items/fichas/RAUS.md",
                "SALI", "Ficha BPAC: revisión vencida", "FARA (Forestal Araucanía) cumple los requisitos para pasar a 'Preselección'"],
               ausentes=["'En revisión'."])

        # Evaluación de propuestas con datos de prueba
        fondo = repo / "ejemplos" / "fondo-innovacion" / "evaluacion"
        with open(fondo / "admisibilidad.csv", "a", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            for p in ("P-001", "P-002", "P-003"):
                for req in ("ADM-01", "ADM-02", "ADM-03", "ADM-04", "ADM-05"):
                    cumple = "no" if (p == "P-003" and req == "ADM-02") or (p == "P-002" and req == "ADM-05") else "si"
                    w.writerow([p, req, cumple, "prueba", "Prueba", "2026-09-20"])
        puntajes = {"P-001": [5, 4, 4, 4, 4], "P-002": [3, 2, 2, 3, 3], "P-003": [5, 4, 4, 3, 4]}
        with open(fondo / "puntajes.csv", "a", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            for p, ps in puntajes.items():
                for i, v in enumerate(ps, 1):
                    w.writerow([p, f"C{i}", v, "Evaluador A", "2026-09-21", ""])
            w.writerow(["P-001", "C2", 2, "Evaluador B", "2026-09-21", "discrepancia de prueba"])
        prueba("puntaje_pauta: admisibilidad, umbrales, ranking y discrepancias",
               [H + "puntaje_pauta.py", "ejemplos/fondo-innovacion", "--fecha", "2026-09-22"], repo,
               ["| 1 | P-001 |", "no admisible: incumple ADM-02", "puntaje 2.55 bajo el mínimo 3.5", "P-002: ADM-05",
                "discrepancia entre evaluadores"])

        prueba("metricas_agil: sprints, DORA, deuda y trazabilidad", [H + "metricas_agil.py", "ejemplos/equipo-software", "--fecha", "2026-09-28"],
               repo, ["Velocidad promedio: 18.7", "Tasa de falla de cambios: 33%", "sin prueba: 3", "RQ-07 sin historia", "prioridad alta: 3"])

        csv_jira = tmp / "jira.csv"
        csv_jira.write_text("Issue key,Summary,Issue Type,Status,Priority,Sprint,Story Points\nAPP-1,Nueva,Story,In Progress,High,S5,3\n", encoding="utf-8")
        prueba("importar_backlog: CSV tipo Jira", [H + "importar_backlog.py", "ejemplos/equipo-software", "--desde", "csv", "--archivo", str(csv_jira)],
               repo, ["nuevos: 1", "Backlog actualizado"])

        prueba("modelo_economico: VAN y TIR con control", [H + "modelo_economico.py", "ejemplos/estudio-solar/estudio/modelo/supuestos-modelo.yaml",
                                                           "--salida", str(tmp / "modelo.xlsx")], repo, ["26.908.200", "TIR: 12.5%", "Recuperación: 7 años"])

        a, b = tmp / "a.md", tmp / "b.md"
        a.write_text("Plazo de entrega: 30 días.\nEl proveedor entrega informes mensuales.\n", encoding="utf-8")
        b.write_text("Plazo de entrega: 45 días.\nEl proveedor entrega informes mensuales.\nEl cliente deberá aprobar en 5 días.\n", encoding="utf-8")
        prueba("comparar_documentos: cambios sustantivos", [H + "comparar_documentos.py", str(a), str(b)], repo,
               ["Cambios: 2", "POSIBLE CAMBIO SUSTANTIVO", "[-30-] {+45+}"])

        # Alta documental: inventario del paquete de prueba
        salida = prueba("crear_proyecto: combinación ejecucion + obra + agil",
                        [H + "crear_proyecto.py", "prueba-combinada", "--arquetipos", "ejecucion,obra,agil", "--nombre", "Prueba"], repo,
                        ["Proyecto creado", "ejecucion, obra, agil"])
        proy = repo / "proyectos" / "prueba-combinada"
        faltan = [p for p in ("cronograma/cronograma.csv", "obra/rfi.csv", "agil/backlog.csv", "metricas/metricas-ejecucion.yaml",
                              "metricas/metricas-agil.yaml", "metricas/metricas.yaml", "contractual/linea-base-vigente.md") if not (proy / p).exists()]
        resultados.append(("crear_proyecto: archivos de los tres arquetipos", not faltan, [f"falta {x}" for x in faltan]))
        shutil.copytree(repo / "ejemplos" / "paquete-alta-andes", proy / "_entrada", dirs_exist_ok=True)
        prueba("extraer_entrada: inventario del paquete de alta", [H + "extraer_entrada.py", "proyectos/prueba-combinada"], repo,
               ["| contrato | 2 |", "| cronograma | 1 |", "| correo | 2 |", "| chat-ia | 1 |", "| planilla-control | 1 |"])
        with open(proy / "obra" / "rfi.csv", "a", encoding="utf-8", newline="") as f:
            csv.writer(f).writerow(["RFI-001", "2026-09-01", "Contratista", "ITO", "Fundaciones", "¿Cota?", "2026-09-05", "", "", "", "abierta", ""])
        prueba("validar_proyecto: registros de obra", [H + "validar_proyecto.py", "proyectos/prueba-combinada", "--hoy", "2026-09-25"], repo,
               ["RFI RFI-001 sin respuesta"])
        prueba("evaluar_metricas: proyecto combinado recién creado", [H + "evaluar_metricas.py", "proyectos/prueba-combinada"], repo,
               ["Evaluación de métricas"])

        # Hooks
        subprocess.run(["git", "init", "-q"], cwd=repo)
        (repo / "proyectos" / "prueba-combinada" / "contractual" / "secreto.pdf").write_text("x")
        evento = json.dumps({"tool_name": "Bash", "tool_input": {"command": "git add -f proyectos/prueba-combinada/contractual/secreto.pdf"}})
        prueba("hook: bloquea archivos confidenciales", [H + "hooks/proteger_confidenciales.py"], repo, ["BLOQUEADO"], codigo=2, entrada=evento)
        malo = repo / "proyectos" / "prueba-combinada" / "riesgos" / "riesgos.csv"
        with open(malo, "a", encoding="utf-8") as f:
            f.write("R-001,25/09/2026,Riesgo,desc,cat,3,3,9,resp,Mitigar,plan,2026-10-01,abierto\n")
            f.write("R-002,2026-09-25,Riesgo,desc,cat,3,3,9,resp,Mitigar,plan,2026-10-01,abierto,COLUMNA_EXTRA\n")
        prueba("hook: detecta CSV mal formado", [H + "hooks/validar_edicion.py"], repo, ["columnas", "formato AAAA-MM-DD"], codigo=2,
               entrada=json.dumps({"tool_input": {"file_path": str(malo)}}))
        prueba("hook: estado al iniciar sesión", [H + "hooks/inicio_sesion.py"], repo, ["[Estado del PMO Harness]", "prueba-combinada"])

        # --- v1.6 · Fase 1: correcciones del núcleo ---
        avance = proy / "_entrada" / "_procesado" / "avance-alta.md"
        avance.parent.mkdir(parents=True, exist_ok=True)
        avance.write_text("# Avance del alta\nestado: en_curso\n", encoding="utf-8")
        prueba("hook: ofrece continuar un alta en curso", [H + "hooks/inicio_sesion.py"], repo, ["prueba-combinada: hay un alta en curso"])
        avance.write_text("# Avance del alta\nestado: terminada\n", encoding="utf-8")
        prueba("hook: no avisa de un alta terminada", [H + "hooks/inicio_sesion.py"], repo, ["[Estado del PMO Harness]"],
               ausentes=["hay un alta en curso"])

        prueba("consola: salida UTF-8 aunque la consola sea cp1252",
               [H + "evaluar_metricas.py", "ejemplos/portafolio-acciones", "--hoy", "2026-08-31"], repo,
               ["Evaluación de métricas", "falta items/fichas/RAUS.md"], env={"PYTHONIOENCODING": "cp1252"})

        prueba("validar_proyecto: tarea que comienza antes que su predecesora",
               [H + "validar_proyecto.py", "ejemplos/demo-andes", "--hoy", "2026-09-25"], repo,
               ["T13 comienza el 2027-04-16, antes de que termine su predecesora T12 (2027-04-30)"])
        demo_yaml = repo / "ejemplos" / "demo-andes" / "proyecto.yaml"
        demo_yaml.write_text(demo_yaml.read_text(encoding="utf-8") + "\ncronograma_actualizado: 2026-09-18\n", encoding="utf-8")
        prueba("validar_proyecto: advierte avance real desactualizado",
               [H + "validar_proyecto.py", "ejemplos/demo-andes", "--hoy", "2026-09-25"], repo,
               ["no se actualiza desde el 2026-09-18 (7 días)"])
        prueba("validar_proyecto: --corte-actualizacion calcula el plan a esa fecha",
               [H + "validar_proyecto.py", "ejemplos/demo-andes", "--hoy", "2026-09-25", "--corte-actualizacion"], repo,
               ["plan al 2026-09-18", "Avance planificado calculado al 2026-09-18"], ausentes=["no se actualiza desde"])

        mediciones = repo / "ejemplos" / "demo-andes" / "metricas" / "mediciones.csv"
        antes = mediciones.read_text(encoding="utf-8")
        prueba("evm_cpm: sin --registrar no modifica mediciones", [H + "evm_cpm.py", "ejemplos/demo-andes", "--fecha", "2026-09-26"], repo,
               ["Resultados no registrados"])
        resultados.append(("evm_cpm: mediciones intactas sin --registrar", mediciones.read_text(encoding="utf-8") == antes,
                           ["mediciones.csv cambió sin --registrar"]))
        correr([H + "evm_cpm.py", "ejemplos/demo-andes", "--fecha", "2026-09-26", "--registrar"], repo)
        resultados.append(("evm_cpm: --registrar guarda mediciones", "2026-09-26,PROYECTO,spi" in mediciones.read_text(encoding="utf-8"),
                           ["no se registró spi del 2026-09-26"]))

        prueba("extraer_entrada: --archivar conserva subcarpetas de carga", [H + "extraer_entrada.py", "proyectos/prueba-combinada", "--archivar"],
               repo, ["Archivados: 11 archivo(s)"])
        ent = proy / "_entrada"
        quedan = [x for x in ent.rglob("*") if x.is_file() and "_archivo" not in x.parts and "_procesado" not in x.parts
                  and x.name not in ("LEEME.md", ".gitkeep") and not x.name.startswith("_LEEME")]
        faltan_sub = [d for d in ("01-contratos-y-anexos", "03-cronogramas", "09-otros") if not (ent / d).is_dir()]
        archivado = list((ent / "_archivo").rglob("Contrato_PAD_firmado.pdf"))
        resultados.append(("extraer_entrada: archivo ordenado y carga vacía", not quedan and not faltan_sub and bool(archivado)
                           and "01-contratos-y-anexos" in archivado[0].as_posix(),
                           [f"quedan: {quedan}", f"faltan subcarpetas: {faltan_sub}", f"archivado: {archivado}"]))

        (proy / "riesgos" / "riesgos.csv").write_text(
            "id,fecha_registro,titulo\nR-9,30/09/2026,Riesgo\n", encoding="utf-8")
        prueba("hook: valida CSV modificado por consola (Bash)", [H + "hooks/validar_edicion.py"], repo,
               ["riesgos.csv", "formato AAAA-MM-DD"], codigo=2, entrada=json.dumps({"tool_name": "Bash", "tool_input": {"command": "cat >> x"}}))


        # --- v1.6 · Fase 2: autoría limpia de los entregables ---
        py_yaml = proy / "proyecto.yaml"
        py_yaml.write_text(py_yaml.read_text(encoding="utf-8").replace(
            '  nombre: ""                   # autor', '  nombre: "María Pérez"       # autor', 1), encoding="utf-8")
        md_ok = tmp / "informe.md"
        md_ok.write_text("# Informe mensual de avance\n\nEl proyecto avanza según lo planificado.\n\n"
                         "| Frente | Estado |\n|---|---|\n| Obras | Verde |\n", encoding="utf-8")
        prueba("md_a_docx: genera Word con autoría limpia", [H + "md_a_docx.py", str(md_ok), str(tmp / "informe.docx"),
               "--proyecto", "proyectos/prueba-combinada"], repo, ["Documento generado"])
        prueba("verificar_autoria: Word generado sin rastros", [H + "verificar_autoria.py", str(tmp / "informe.docx"),
               "--proyecto", "proyectos/prueba-combinada"], repo, ["Resultado: OK"], ausentes=["ERROR"])
        from docx import Document as _Doc
        props = _Doc(str(tmp / "informe.docx")).core_properties
        resultados.append(("md_a_docx: autor y fecha actuales", props.author == "María Pérez" and props.created.year >= 2026
                           and "python-docx" not in (props.comments or ""), [f"autor={props.author!r} creado={props.created}"]))

        md_pend = tmp / "pendiente.md"
        md_pend.write_text("# Informe\n\nMonto: [PENDIENTE: confirmar monto]\n", encoding="utf-8")
        prueba("md_a_docx: no emite con [PENDIENTE]", [H + "md_a_docx.py", str(md_pend), str(tmp / "p.docx")], repo,
               ["marcador(es) [PENDIENTE]"], codigo=1)

        md_ia = tmp / "rastros.md"
        md_ia.write_text("# Informe\n\nEste informe fue generado automáticamente con Claude en el harness.​\n", encoding="utf-8")
        prueba("verificar_autoria: detecta menciones y caracteres invisibles", [H + "verificar_autoria.py", str(md_ia)], repo,
               ["término «Claude»", "generado automáticamente", "harness", "carácter(es) invisible(s)"], codigo=1)

        from docx import Document as _Doc2
        crudo = tmp / "crudo.docx"
        d = _Doc2()
        d.add_paragraph("Texto normal del informe.")
        d.save(str(crudo))
        prueba("verificar_autoria: detecta metadatos de python-docx", [H + "verificar_autoria.py", str(crudo)], repo,
               ["python-docx", "customXml", "fecha de creación antigua"], codigo=1)
        prueba("verificar_autoria: --corregir limpia metadatos y partes heredadas",
               [H + "verificar_autoria.py", str(crudo), "--proyecto", "proyectos/prueba-combinada", "--corregir"], repo,
               ["corregido en crudo.docx", "Resultado: OK"])

        from pptx import Presentation as _Prs
        ppt = tmp / "presentacion.pptx"
        prs = _Prs()
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Estado del proyecto"
        prs.save(str(ppt))
        prueba("verificar_autoria: detecta metadatos de python-pptx", [H + "verificar_autoria.py", str(ppt)], repo,
               ["python-pptx"], codigo=1)
        prueba("verificar_autoria: --corregir limpia PowerPoint",
               [H + "verificar_autoria.py", str(ppt), "--proyecto", "proyectos/prueba-combinada", "--corregir"], repo, ["Resultado: OK"])

        prueba("modelo_economico: Excel con autoría limpia", [H + "verificar_autoria.py", str(tmp / "modelo.xlsx")], repo,
               ["Resultado: OK"], ausentes=["openpyxl"])


        # --- v1.6 · Fase 3: arquetipo postulacion ---
        prueba("crear_proyecto: arquetipo postulacion", [H + "crear_proyecto.py", "prueba-postulacion", "--arquetipos", "postulacion",
               "--perfil", "biblioteca/perfiles/fondo-corfo"], repo, ["Proyecto creado"])
        pp = repo / "proyectos" / "prueba-postulacion"
        faltan = [p for p in ("postulacion/etapas.yaml", "postulacion/requisitos.csv", "postulacion/pauta-inversa.yaml",
                              "redteam/perspectivas.yaml", "economico/presupuesto.yaml", "entrega/checklist.md",
                              "metricas/metricas-postulacion.yaml", "reportes/catalogo/estado-postulacion.md",
                              "_entrada/00-bases-y-aclaraciones") if not (pp / p).exists()]
        texto_yaml = (pp / "proyecto.yaml").read_text(encoding="utf-8")
        resultados.append(("crear_proyecto: estructura y bloque postulacion", not faltan and "decision_go:" in texto_yaml
                           and "biblioteca/perfiles/fondo-corfo" in texto_yaml, [f"falta {x}" for x in faltan]))
        paquete = repo / "ejemplos" / "paquete-postulacion-demo"
        for sub in ("00-bases-y-aclaraciones", "06-correos"):
            shutil.copytree(paquete / sub, pp / "_entrada" / sub, dirs_exist_ok=True)
        prueba("extraer_entrada: reconoce bases de concurso y el formulario", [H + "extraer_entrada.py", "proyectos/prueba-postulacion"], repo,
               ["| bases | 2 |", "| formulario | 1 |", "| correo | 1 |"])
        prueba("extraer_requisitos: barrido detecta excluyente escondido en anexo y mención de IA",
               [H + "extraer_requisitos.py", "proyectos/prueba-postulacion"], repo,
               ["carta de compromiso", "causal de inadmisibilidad", "ATENCIÓN: 1 candidato(s) mencionan inteligencia artificial",
                "nuevos en esta pasada: 24"])
        barrido = (pp / "postulacion" / "barrido-requisitos.csv").read_text(encoding="utf-8")
        claves = ["Podrán postular personas jurídicas", "no podrá exceder 12 páginas", "al menos el 50% deberá ser pecuniario",
                  "hasta el 20 de octubre de 2026", "nota final igual o superior a 3,5", "No son financiables"]
        resultados.append(("extraer_requisitos: cobertura de requisitos clave", all(c in barrido for c in claves),
                           [f"no encontrado: {c}" for c in claves if c not in barrido]))
        prueba("extraer_requisitos: incremental sin duplicar", [H + "extraer_requisitos.py", "proyectos/prueba-postulacion"], repo,
               ["nuevos en esta pasada: 0"])
        (pp / "postulacion" / "bases" / "AC02-aclaracion-2.md").write_text(
            "# Aclaración N°2\n\nLos postulantes deberán adjuntar una carta Gantt del proyecto en formato PDF.\n", encoding="utf-8")
        prueba("extraer_requisitos: una aclaración nueva agrega solo lo nuevo", [H + "extraer_requisitos.py", "proyectos/prueba-postulacion"],
               repo, ["nuevos en esta pasada: 1"])

        # v2.1 · Réplicas fieles: las bases quedan en fuentes/ y el barrido funciona sin los originales (otro computador)
        _rep = sorted((pp / "fuentes").glob("F*.md"))
        _txt_rep = "\n".join(x.read_text(encoding="utf-8") for x in _rep)
        _inv_pp = (pp / "_entrada" / "_procesado" / "inventario.md").read_text(encoding="utf-8")
        resultados.append(("extraer_entrada: réplicas fieles de las bases en fuentes/, el correo solo con ficha",
                           len(_rep) == 3 and "categoria: bases" in _txt_rep and "calidad: completa" in _txt_rep and "huella:" in _txt_rep
                           and "Podrán postular personas jurídicas" in _txt_rep and "categoria: correo" not in _txt_rep
                           and "Réplicas en fuentes/" in _inv_pp and "Solo ficha" in _inv_pp,
                           [f"réplicas: {[x.name for x in _rep]}", _inv_pp[-700:]]))
        shutil.move(str(pp / "_entrada" / "_procesado"), str(tmp / "procesado-respaldo"))
        prueba("extraer_requisitos: sin los originales, el barrido usa las réplicas y no duplica", [H + "extraer_requisitos.py",
               "proyectos/prueba-postulacion"], repo, ["nuevos en esta pasada: 0"])
        shutil.move(str(tmp / "procesado-respaldo"), str(pp / "_entrada" / "_procesado"))
        prueba("postulacion: alerta barrido sin revisar, IA y compuerta de recepción",
               [H + "postulacion.py", "proyectos/prueba-postulacion", "--hoy", "2026-10-10"], repo,
               ["Candidato con lenguaje de exclusión sin revisar", "Las bases mencionan inteligencia artificial",
                "calendario_con_cierre == 1 (actual: 0)", "← actual"])

        demo = "ejemplos/postulacion-demo"
        prueba("presupuesto_postulacion: controles detectan aporte pecuniario bajo el mínimo",
               [H + "presupuesto_postulacion.py", demo + "/economico/presupuesto.yaml", "--salida", str(tmp / "presupuesto.xlsx")], repo,
               ["CTL-04 NO CUMPLE", "CTL-01 NO CUMPLE", "CTL-09 CUMPLE", "Controles NO CUMPLE: 3"])
        prueba("presupuesto_postulacion: Excel con autoría limpia", [H + "verificar_autoria.py", str(tmp / "presupuesto.xlsx")], repo,
               ["Resultado: OK"])
        from openpyxl import load_workbook as _lw
        _wb = _lw(str(tmp / "presupuesto.xlsx"))
        formula_ok = str(_wb["Controles"]["E4"].value).startswith("=IF(") and str(_wb["Presupuesto"]["F4"].value) == "=D4*E4"
        resultados.append(("presupuesto_postulacion: hojas y fórmulas", _wb.sheetnames == ["Parametros", "Presupuesto", "Resumen", "Controles", "Cronograma"]
                           and formula_ok, [f"hojas: {_wb.sheetnames}"]))
        ok_yaml = tmp / "presupuesto-ok.yaml"
        ok_yaml.write_text((repo / demo / "economico" / "presupuesto.yaml").read_text(encoding="utf-8")
                           .replace("cantidad: 1, costo_unitario: 4000000, entidad: 1, pecuniario: 0",
                                    "cantidad: 1, costo_unitario: 4000000, entidad: 0, pecuniario: 1")
                           .replace("cantidad: 6, costo_unitario: 1000000, entidad: 1, pecuniario: 0",
                                    "cantidad: 6, costo_unitario: 1000000, entidad: 0.5, pecuniario: 0.5")
                           .replace("cantidad: 12, costo_unitario: 1000000, entidad: 0, pecuniario: 0, no_pecuniario: 1",
                                    "cantidad: 12, costo_unitario: 1000000, entidad: 0, pecuniario: 0.5, no_pecuniario: 0.5"), encoding="utf-8")
        prueba("presupuesto_postulacion: presupuesto que cumple", [H + "presupuesto_postulacion.py", str(ok_yaml), "--salida", str(tmp / "ok.xlsx")],
               repo, ["Controles NO CUMPLE: 0"])

        prueba("postulacion: estado del demo (etapa red team)", [H + "postulacion.py", demo, "--hoy", "2026-11-03", "--registrar"], repo,
               ["Presupuesto: CTL-04 NO CUMPLE", "Falla fatal abierta RT-01", "criterios bajo el mínimo por criterio de las bases: C4",
                "Documento vencido al cierre: Certificado de vigencia", "| 8 | Red team (revisión roja) ← actual | Pendiente",
                "| Puntaje estimado red team (mín. / objetivo) | 3.40 (3.5 / 4) |", "Excluyente pendiente REQ-03",
                "declaración son del postulante"])
        prueba("evaluar_metricas: reglas del arquetipo postulacion", [H + "evaluar_metricas.py", demo, "--hoy", "2026-11-03"], repo,
               ["POS-01 incumplida", "POS-02 incumplida", "POS-03 incumplida", "POS-04 incumplida"])

        (pp / "contenidos" / "archivos" / "cv.pdf").write_text("x")
        evento = json.dumps({"tool_name": "Bash", "tool_input": {"command": "git add -f proyectos/prueba-postulacion/contenidos/archivos/cv.pdf"}})
        prueba("hook: bloquea documentos del banco de contenidos", [H + "hooks/proteger_confidenciales.py"], repo,
               ["BLOQUEADO", "banco de contenidos"], codigo=2, entrada=evento)

        evento = json.dumps({"tool_name": "PowerShell", "tool_input": {"command": "git add -f proyectos/prueba-postulacion/_entrada/00-bases-y-aclaraciones/.gitkeep"}})
        prueba("hook: permite marcadores .gitkeep (también desde PowerShell)", [H + "hooks/proteger_confidenciales.py"], repo,
               [], ausentes=["BLOQUEADO"], codigo=0, entrada=evento)


        # --- v1.6 · Fase 4: guía del usuario ---
        prueba("guia_proyecto: guía de una postulación", [H + "guia_proyecto.py", "ejemplos/postulacion-demo"], repo,
               ["Postulación (oferente)", "«prepara la postulación»", "`matriz-cumplimiento`", "Controles automáticos",
                "Estado de la postulación", "No hace falta usar el nombre exacto"], ausentes=["`gestionar-sprint`"])
        resultados.append(("guia_proyecto: escribe GUIA.md", (repo / "ejemplos" / "postulacion-demo" / "GUIA.md").exists(), ["no se escribió GUIA.md"]))
        prueba("guia_proyecto: guía de un proyecto de ejecución", [H + "guia_proyecto.py", "ejemplos/demo-andes", "--no-escribir"], repo,
               ["`actualizar-gantt`", "`control-cambios`"], ausentes=["`preparar-postulacion`", "`gestionar-sprint`"])
        prueba("guia_proyecto: guía general", [H + "guia_proyecto.py", "--general"], repo, ["`instalar`", "`nuevo-proyecto`", "Postulación (oferente)"])
        import yaml as _yaml
        _cat = _yaml.safe_load((repo / ".claude" / "skills" / "ayuda" / "catalogo.yaml").read_text(encoding="utf-8"))
        _ids = {s["id"] for s in _cat["skills"]}
        _dirs = {p.name for p in (repo / ".claude" / "skills").iterdir() if p.is_dir()}
        _manif = RAIZ / ".harness" / "nucleo-instalado.csv"
        if _manif.exists():  # carpeta de empresa: las skills propias de la empresa no son del núcleo
            _del_nucleo = _manif.read_text(encoding="utf-8")
            _dirs = {d for d in _dirs if f".claude/skills/{d}/SKILL.md" in _del_nucleo}
        resultados.append(("ayuda: el catálogo cubre todas las skills", _ids == _dirs,
                           [f"faltan en el catálogo: {sorted(_dirs - _ids)}", f"sobran: {sorted(_ids - _dirs)}"]))


        # --- v1.6 · Correcciones de la revisión independiente ---
        borde = tmp / "borde"
        (borde / "postulacion").mkdir(parents=True)
        (borde / "bases.md").write_text(
            "# Bases\n\n## ARTÍCULO 5. Requisitos\n\n1. Adjuntar certificado de vigencia de la sociedad.\n"
            "2. Presentar la oferta en la plataforma.\n3. Acompañar carta de compromiso firmada, so pena de inadmisibilidad.\n\n"
            "ARTÍCULO 6: Los postulantes deberán declarar el uso de Inteligencia Artificial.\n\n"
            "Se prohíbe el uso de ChatGPT u otras herramientas generativas en la redacción.\n\n"
            "Según el art. 8 de las bases, el aporte mínimo será de 30%.\n", encoding="utf-8")
        prueba("extraer_requisitos: listas numeradas, artículos, IA en mayúsculas y abreviaturas",
               [H + "extraer_requisitos.py", str(borde), "--origen", str(borde / "bases.md")], repo,
               ["candidatos: 6", "so pena de inadmisibilidad", "ATENCIÓN: 2 candidato(s)"])
        texto_b = (borde / "postulacion" / "barrido-requisitos.csv").read_text(encoding="utf-8")
        resultados.append(("extraer_requisitos: no corta «art. 8»", "Según el art. 8 de las bases, el aporte mínimo será de 30%." in texto_b,
                           ["la frase con «art. 8» quedó cortada"]))

        med = repo / demo / "metricas" / "mediciones.csv"
        med.write_text("fecha,entidad,metrica,valor,fuente,comentario,revisado_por\n2026-10-01,PROYECTO,manual,1,usuario,,Paula\n", encoding="utf-8")
        correr([H + "postulacion.py", demo, "--hoy", "2026-11-03", "--registrar"], repo)
        contenido = med.read_text(encoding="utf-8")
        resultados.append(("postulacion --registrar: conserva columnas extra y filas previas",
                           contenido.startswith("fecha,entidad,metrica,valor,fuente,comentario,revisado_por")
                           and "2026-10-01,PROYECTO,manual,1,usuario,,Paula" in contenido and "cumplimiento_pct" in contenido,
                           [contenido[:300]]))

        decl = tmp / "declaracion-uso-ia.md"
        decl.write_text("# Declaración\n\nDeclaro que utilicé Claude (Anthropic) para ordenar ideas.\n", encoding="utf-8")
        pyml = repo / demo / "proyecto.yaml"
        pyml.write_text(pyml.read_text(encoding="utf-8").replace("  terminos_permitidos: []\n",
                        "  terminos_permitidos: []\n  excluir: [\"declaracion-uso-ia.*\"]\n", 1), encoding="utf-8")
        prueba("verificar_autoria: una declaración exigida y excluida no bloquea", [H + "verificar_autoria.py", str(decl), "--proyecto", demo],
               repo, ["excluido de la verificación", "Resultado: OK"])
        legitimo = tmp / "propuesta.md"
        legitimo.write_text("# Propuesta\n\nUsaremos datos del Observatorio Gemini Sur; las alertas son generadas automáticamente.\n", encoding="utf-8")
        prueba("verificar_autoria: términos ambiguos son advertencia, no error", [H + "verificar_autoria.py", str(legitimo)], repo,
               ["revisar: término «Gemini»", "Resultado: OK"])
        resto = tmp / "resto.md"
        resto.write_text("# Informe\n\n¡Claro! Aquí tienes el informe solicitado.\n", encoding="utf-8")
        prueba("verificar_autoria: restos de conversación con un asistente son error", [H + "verificar_autoria.py", str(resto)], repo,
               ["ERROR: término «Aquí tienes»"], codigo=1)

        prueba("presupuesto_postulacion: plantilla vacía no pasa la compuerta (CTL-00)",
               [H + "presupuesto_postulacion.py", "proyectos/prueba-postulacion/economico/presupuesto.yaml"], repo, ["CTL-00 NO CUMPLE"])
        malo = tmp / "presupuesto-malo.yaml"
        malo.write_text("nombre: X\nparametros: {financiamiento_entidad_max_pct: '70%'}\nitems: [{id: ADM, nombre: Adm, tope_pct: '10%'}]\n"
                        "partidas: [{item: GAD, descripcion: Gastos, cantidad: 1, costo_unitario: '1.500.000', entidad: 1}]\n", encoding="utf-8")
        prueba("presupuesto_postulacion: ítem no declarado detiene el cálculo", [H + "presupuesto_postulacion.py", str(malo), "--salida", str(tmp / "m.xlsx")],
               repo, ["usa el ítem 'GAD'"], codigo=1)
        malo.write_text(malo.read_text(encoding="utf-8").replace("item: GAD", "item: ADM"), encoding="utf-8")
        prueba("presupuesto_postulacion: acepta «70%» y «1.500.000»", [H + "presupuesto_postulacion.py", str(malo), "--salida", str(tmp / "m.xlsx")],
               repo, ["Costo total: 1.500.000", "CTL-01 NO CUMPLE", "CTL-06 NO CUMPLE"])

        obs = repo / demo / "redteam" / "observaciones.csv"
        obs.write_text(obs.read_text(encoding="utf-8").replace(",abierta\n", ",Abierta \n"), encoding="utf-8")
        with open(repo / demo / "redteam" / "puntajes.csv", "a", encoding="utf-8", newline="") as f:
            f.write("roja-2,financiero,C4,4,Corrigió sostenibilidad,2026-11-04\n")
        prueba("postulacion: estados con mayúsculas y ronda parcial roja-2", [H + "postulacion.py", demo, "--hoy", "2026-11-04"], repo,
               ["Falla fatal abierta RT-01", "Falla fatal abierta RT-02", "| financiero | 3.55 | sí |"])
        cont = repo / demo / "contenidos" / "registro.csv"
        cont.write_text(cont.read_text(encoding="utf-8").replace("2026-12-31", "31-12-2026"), encoding="utf-8")
        prueba("postulacion: vigencia con fecha inválida se trata como vencida", [H + "postulacion.py", demo, "--hoy", "2026-11-04"], repo,
               ["vigencia '31-12-2026' inválida"])
        prueba("postulacion: fecha de corte inválida", [H + "postulacion.py", demo, "--hoy", "2026-13-45"], repo, ["Fecha de corte inválida"], codigo=1)

        # Carpeta del harness dentro de otro repositorio (el repo de prueba aún no tiene .git propio en su padre)
        subprocess.run(["git", "init", "-q"], cwd=tmp)
        shutil.rmtree(repo / ".git", ignore_errors=True)
        prueba("verificar_entorno: detecta repositorio padre", [H + "verificar_entorno.py"], repo,
               ["DENTRO de otro repositorio"], codigo=None)
        subprocess.run(["git", "init", "-q"], cwd=repo)
        prueba("verificar_entorno: repositorio propio", [H + "verificar_entorno.py"], repo,
               ["Repositorio Git: sí"], ausentes=["DENTRO de otro repositorio"], codigo=None)


        # --- v1.6 · Programas con subproyectos ---
        prueba("crear_proyecto: programa", [H + "crear_proyecto.py", "prog-x", "--arquetipos", "programa,ejecucion", "--nombre", "Programa X"],
               repo, ["Programa creado"])
        prueba("crear_proyecto: subproyecto dentro del programa", [H + "crear_proyecto.py", "erp", "--programa", "prog-x", "--nombre", "ERP"],
               repo, ["Subproyecto creado: proyectos/prog-x/subproyectos/erp", "Registrado en proyectos/prog-x/programa/subproyectos.csv"])
        prueba("crear_proyecto: subproyecto ágil", [H + "crear_proyecto.py", "app", "--programa", "prog-x", "--arquetipos", "ejecucion,agil",
               "--nombre", "App"], repo, ["Subproyecto creado"])
        px = repo / "proyectos" / "prog-x"
        lista = (px / "programa" / "subproyectos.csv").read_text(encoding="utf-8")
        ok_estructura = ("erp,ERP,subproyectos/erp" in lista and "app,App,subproyectos/app" in lista
                         and not (px / "subproyectos" / "erp" / "contractual").exists() and not (px / "subproyectos" / "erp" / "_entrada").exists()
                         and (px / "subproyectos" / "app" / "agil" / "backlog.csv").exists()
                         and 'programa: "../.."' in (px / "subproyectos" / "erp" / "proyecto.yaml").read_text(encoding="utf-8"))
        resultados.append(("crear_proyecto: subproyectos registrados, sin contrato ni entrada propios", ok_estructura, [lista]))
        prueba("crear_proyecto: un subproyecto exige un programa", [H + "crear_proyecto.py", "zz", "--programa", "prueba-combinada"], repo,
               ["no tiene el arquetipo programa"], codigo=1)
        prueba("guia_proyecto: subproyecto indica su programa", [H + "guia_proyecto.py", "proyectos/prog-x/subproyectos/erp", "--no-escribir"],
               repo, ["Es un subproyecto del programa «Programa X»"])

        pd = "ejemplos/programa-demo"
        prueba("programa: estado consolidado del demo", [H + "programa.py", pd, "--hoy", "2026-10-15", "--excel", "--registrar"], repo,
               ["sp2 (App de despacho) termina el 2027-03-20, después de su término vigente (2027-02-28)",
                "Luis Mora: hasta 130% (capacidad 100%) entre 2027-01 y 2027-01", "Enmienda E02 (plazo) sin aplicar en sp2",
                "El estimado al término supera el presupuesto vigente en 35.000.000 CLP",
                "| **Total programa** | 1.090.000.000 | 130.000.000 | 625.000.000 | 335.000.000 | 69% | 1.125.000.000 | -35.000.000 | 45.000.000 / 110.000.000 |",
                "| sp1 · ERP | 64 | 61 | -3 |", "ruta crítica: sp3.T03 → PRG.T02"])
        cons = (repo / pd / "cronograma" / "consolidado.csv").read_text(encoding="utf-8")
        generados = all((repo / pd / r).exists() for r in ("cronograma/gantt-programa.md", "cronograma/gantt-detalle.md",
                                                         "subproyectos/sp2/cronograma/gantt.md", "presupuesto/control-presupuestario.xlsx"))
        resultados.append(("programa: consolidado con dependencias entre subproyectos y Gantt generados",
                           "sp2.T03,sp2,Piloto" not in cons and "sp2.T02;sp1.T04" in cons and generados, [cons[:400]]))
        prueba("programa: control presupuestario con autoría limpia", [H + "verificar_autoria.py", pd + "/presupuesto/control-presupuestario.xlsx",
               "--proyecto", pd], repo, ["Resultado: OK"])
        from openpyxl import load_workbook as _lw2
        _wbp = _lw2(str(repo / pd / "presupuesto" / "control-presupuestario.xlsx"))
        resultados.append(("programa: Excel con fórmulas por subproyecto", _wbp.sheetnames[:2] == ["Resumen", "Detalle"]
                           and str(_wbp["Resumen"]["B4"].value).startswith("=SUMIFS(") and str(_wbp["Detalle"]["E4"].value) == "=C4+D4",
                           [str(_wbp.sheetnames)]))
        med_p = (repo / pd / "metricas" / "mediciones.csv").read_text(encoding="utf-8")
        resultados.append(("programa --registrar: métricas del programa y por subproyecto",
                           ",PROYECTO,subproyectos_atrasados,1," in med_p and ",sp2,vac,-80000000," in med_p, [med_p[:300]]))
        prueba("validar_proyecto: subproyecto contra su propia línea base", [H + "validar_proyecto.py", pd + "/subproyectos/sp2", "--hoy", "2026-10-15"],
               repo, ["después del término contractual vigente (2027-02-28)", "depende de 'sp1:T04' (otro subproyecto)"],
               ausentes=["predecesora 'sp1:T04' no existe"])
        prueba("evaluar_metricas: reglas del programa", [H + "evaluar_metricas.py", pd, "--hoy", "2026-10-15"], repo,
               ["PRG-01 incumplida", "PRG-02 incumplida"])
        doc_sp = tmp / "sp2.md"
        doc_sp.write_text("# Informe del subproyecto\n\nAvance del período.\n", encoding="utf-8")
        correr([H + "md_a_docx.py", str(doc_sp), str(tmp / "sp2.docx"), "--proyecto", pd + "/subproyectos/sp2"], repo)
        from docx import Document as _Doc3
        _d3 = _Doc3(str(tmp / "sp2.docx"))
        resultados.append(("md_a_docx: subproyecto usa la marca y el cliente del programa",
                           "Programa Logística Integrada · App de despacho" in _d3.sections[0].header.paragraphs[0].text
                           and "Distribuidora Austral" in _d3.sections[0].footer.paragraphs[0].text,
                           [_d3.sections[0].header.paragraphs[0].text, _d3.sections[0].footer.paragraphs[0].text]))

        # Aplicar la Enmienda N°2 limpia las alertas; errores de dependencias y contingencia se detectan
        sp2y = repo / pd / "subproyectos" / "sp2" / "proyecto.yaml"
        sp2y.write_text(sp2y.read_text(encoding="utf-8").replace("fecha_termino_vigente: 2027-02-28", "fecha_termino_vigente: 2027-03-31"), encoding="utf-8")
        afe = repo / pd / "contractual" / "afectaciones.csv"
        afe.write_text(afe.read_text(encoding="utf-8").replace("modo sin conexión,no", "modo sin conexión,si"), encoding="utf-8")
        prueba("programa: aplicar la enmienda limpia las alertas", [H + "programa.py", pd, "--hoy", "2026-10-15", "--sin-gantt"], repo,
               ["| sp2 · App de despacho | 45 | 47 | +3 | 2027-03-20 | 2027-03-31 | Verde |"], ausentes=["E02 (plazo) sin aplicar", "sp2 (App de despacho) termina"])
        c3 = repo / pd / "subproyectos" / "sp3" / "cronograma" / "cronograma.csv"
        c3.write_text(c3.read_text(encoding="utf-8").replace(",sp1:T01,", ",sp1:T04,").replace(",T01,no,no iniciada", ",T01;sp9:T01,no,no iniciada"), encoding="utf-8")
        mo = repo / pd / "presupuesto" / "modificaciones.csv"
        mo.write_text(mo.read_text(encoding="utf-8").replace("SP2,RRHH,40000000", "SP2,RRHH,40000000").replace("sp2,RRHH,40000000", "sp2,RRHH,30000000"), encoding="utf-8")
        prueba("programa: detecta dependencias violadas, inexistentes y contingencia que no cuadra", [H + "programa.py", pd, "--hoy", "2026-10-15"],
               repo, ["Dependencia entre subproyectos: sp3.T01 comienza el 2026-09-01, antes de que termine sp1.T04",
                      "depende de 'sp9:T01', que no existe", "Movimiento M-04 (reasignación o uso de contingencia) no suma cero"])
        prueba("hook: muestra el programa con sus subproyectos", [H + "hooks/inicio_sesion.py"], repo, ["prog-x (programa con 2 subproyecto(s)"])


        # --- v2.0 · Importación desde MS Project (XML) ---
        xml_msp = "ejemplos/msproject/cronograma-programa.xml"
        correr([H + "crear_proyecto.py", "prog-msp", "--arquetipos", "programa,ejecucion"], repo)
        for sid in ("erp", "app", "portal"):
            correr([H + "crear_proyecto.py", sid, "--programa", "prog-msp"], repo)
        mapa = "ERP=erp,App de despacho=app,Portal de clientes=portal"
        pm = repo / "proyectos" / "prog-msp"
        antes_app = (pm / "subproyectos" / "app" / "cronograma" / "cronograma.csv").read_text(encoding="utf-8")
        prueba("importar_msproject: vista previa no reemplaza nada", [H + "importar_msproject.py", xml_msp, "proyectos/prog-msp", "--mapa", mapa],
               repo, ["Tareas importadas: 11 · hitos: 3 · recursos: 4 · costo planificado total: 565.000.000",
                      "2 dependencia(s) no fin-comienzo", "1 dependencia(s) con desfase", "Vista previa: nada se reemplazó"])
        resultados.append(("importar_msproject: cronograma vigente intacto en la vista previa",
                           (pm / "subproyectos" / "app" / "cronograma" / "cronograma.csv").read_text(encoding="utf-8") == antes_app
                           and (pm / "subproyectos" / "app" / "cronograma" / "importado-msproject.csv").exists(), ["se modificó cronograma.csv"]))
        prueba("importar_msproject: --reemplazar con respaldo y línea base", [H + "importar_msproject.py", xml_msp, "proyectos/prog-msp",
               "--mapa", mapa, "--reemplazar", "--linea-base"], repo, ["línea base 0 de MS Project"])
        app_csv = (pm / "subproyectos" / "app" / "cronograma" / "cronograma.csv").read_text(encoding="utf-8")
        erp_csv = (pm / "subproyectos" / "erp" / "cronograma" / "cronograma.csv").read_text(encoding="utf-8")
        resultados.append(("importar_msproject: ramas por subproyecto, dependencias cruzadas, CC con desfase, costos y recursos",
                           "T11,App de despacho,Piloto con choferes" in app_csv and "T10;erp:T6" in app_csv and "MSP: T9 CC+5d" in app_csv
                           and "MSP: T9 FF" in app_csv and "Carla Vidal; Luis Mora" in app_csv and ",PRG:T1," in app_csv
                           and "T3,ERP,Levantamiento de procesos,Pedro Soto,2026-06-02,2026-07-15,2026-06-02,2026-07-15,100,PRG:T1,no,terminada,,60000000,65000000" in erp_csv
                           and list((pm / "subproyectos" / "erp" / "cronograma" / "linea-base").glob("*-msproject.csv"))
                           and list((pm / "subproyectos" / "app" / "cronograma" / "respaldos").glob("cronograma-*.csv")),
                           [app_csv[:600], erp_csv[:300]]))
        prueba("programa: consolida lo importado de MS Project", [H + "programa.py", "proyectos/prog-msp", "--hoy", "2026-10-15"], repo,
               ["| erp · erp | 64 | 61 |", "ruta crítica"])
        (tmp / "plan.mpp").write_bytes(b"\xd0\xcf\x11\xe0binario")
        prueba("importar_msproject: .mpp pide exportar a XML", [H + "importar_msproject.py", str(tmp / "plan.mpp"), "proyectos/prog-msp"], repo,
               ["Guardar como"], codigo=1)
        prueba("importar_msproject: proyecto independiente, todo a su cronograma", [H + "importar_msproject.py", xml_msp, "proyectos/prueba-combinada"],
               repo, ["11 tarea(s) → proyectos/prueba-combinada/cronograma/importado-msproject.csv"])
        ent_msp = pm / "_entrada" / "03-cronogramas"
        ent_msp.mkdir(parents=True, exist_ok=True)
        shutil.copy2(repo / xml_msp, ent_msp / "Plan maestro.xml")
        (ent_msp / "Plan maestro v2.mpp").write_bytes(b"\xd0\xcf\x11\xe0otro")
        prueba("extraer_entrada: reconoce XML y .mpp de MS Project", [H + "extraer_entrada.py", "proyectos/prog-msp"], repo,
               ["| cronograma | 2 |"])
        inv = (pm / "_entrada" / "_procesado" / "inventario.md").read_text(encoding="utf-8")
        resultados.append(("extraer_entrada: resumen de MS Project e instrucción para .mpp",
                           "cronograma de MS Project (XML): 11 tareas" in inv and "Guardar como > tipo XML" in inv, [inv[-800:]]))

        # --- v2.1 · Cualquier asistente de IA: estado, traspaso, control de Git y lanzador de Python ---
        _agents = (repo / "AGENTS.md").read_text(encoding="utf-8")
        _sin_fila = sorted(d for d in _dirs if f".claude/skills/{d}/SKILL.md" not in _agents)
        resultados.append(("AGENTS.md: la tabla de procedimientos cubre todas las skills", not _sin_fila, [f"faltan: {_sin_fila}"]))
        _entradas = {n: (repo / n).read_text(encoding="utf-8") for n in ("CLAUDE.md", "GEMINI.md")}
        resultados.append(("CLAUDE.md y GEMINI.md remiten a AGENTS.md sin duplicar reglas",
                           all("@AGENTS.md" in t and "## Principios de trabajo" not in t for t in _entradas.values()),
                           [f"{n}: {t[:200]}" for n, t in _entradas.items()]))

        trasp = pp / "seguimiento" / "traspaso.md"
        resultados.append(("crear_proyecto: el proyecto nuevo trae su traspaso", trasp.exists(), ["falta seguimiento/traspaso.md"]))
        trasp.write_text(trasp.read_text(encoding="utf-8").replace("fecha:", "fecha: 2026-10-19", 1)
                         .replace("asistente:  ", "asistente: codex  ", 1).replace("estado: sin_pendientes", "estado: en_curso", 1)
                         .replace("tarea_en_curso:", "tarea_en_curso: matriz de cumplimiento", 1)
                         .replace("proximo_paso:", "proximo_paso: revisar los candidatos del barrido", 1), encoding="utf-8")
        prueba("estado: muestra el traspaso pendiente y los controles de Git inactivos", [H + "estado.py"], repo,
               ["[Estado del PMO Harness]", "prueba-postulacion: traspaso (2026-10-19, codex): tarea en curso «matriz de cumplimiento»; "
                "próximo paso «revisar los candidatos del barrido»", "Controles de Git INACTIVOS"])
        subprocess.run(["git", "config", "core.hooksPath", ".githooks"], cwd=repo)
        (repo / ".harness").mkdir(exist_ok=True)
        (repo / ".harness" / "estado.yaml").write_text('instalado: true\nversion_harness: "2.1"\n', encoding="utf-8")
        prueba("estado: instalado sin asistentes registrados pide configurar-asistente", [H + "estado.py"], repo,
               ["Ningún asistente registrado", "configurar-asistente"], ausentes=["Controles de Git INACTIVOS"])
        (repo / ".harness" / "asistentes.yaml").write_text("asistentes:\n  - id: claude-code\n  - id: codex\n", encoding="utf-8")
        prueba("estado: asistentes registrados y archivos sin guardar por proyecto", [H + "estado.py"], repo,
               ["Asistentes configurados: claude-code, codex", "prueba-postulacion:", "modificados sin guardar"])
        shutil.rmtree(repo / ".harness")

        def commit_con_control(archivos, forzar=False):
            subprocess.run(["git", "reset", "-q"], cwd=repo)
            subprocess.run(["git", "add"] + (["-f"] if forzar else []) + archivos, cwd=repo)
            r = subprocess.run(["git", "-c", "user.name=Prueba", "-c", "user.email=prueba@ejemplo.cl", "commit", "-q", "-m", "prueba"],
                               cwd=repo, capture_output=True, text=True, encoding="utf-8", errors="replace")
            return r.returncode, (r.stdout or "") + (r.stderr or "")

        (pp / "_entrada" / "nota.txt").write_text("material del cliente", encoding="utf-8")
        rc_c, sal_c = commit_con_control(["proyectos/prueba-postulacion/_entrada/nota.txt"], forzar=True)
        resultados.append(("control de Git: detiene un commit con material de _entrada", rc_c != 0 and "commit detenido" in sal_c
                           and "_entrada/nota.txt" in sal_c, [f"rc={rc_c}", sal_c[:600]]))
        rc_c, sal_c = commit_con_control(["proyectos/prueba-combinada/riesgos/riesgos.csv"])
        resultados.append(("control de Git: detiene un commit con un CSV mal formado", rc_c != 0 and "formato AAAA-MM-DD" in sal_c,
                           [f"rc={rc_c}", sal_c[:600]]))
        (pp / "fuentes").mkdir(exist_ok=True)
        (pp / "fuentes" / "CV01-curriculum.md").write_text("---\nid: CV01\nconfidencialidad: sensible\n---\nDatos personales\n", encoding="utf-8")
        rc_c, sal_c = commit_con_control(["proyectos/prueba-postulacion/fuentes/CV01-curriculum.md"])
        resultados.append(("control de Git: detiene una réplica marcada sensible", rc_c != 0 and "réplica marcada" in sal_c,
                           [f"rc={rc_c}", sal_c[:600]]))
        rc_c, sal_c = commit_con_control(["proyectos/prueba-postulacion/seguimiento/traspaso.md"])
        resultados.append(("control de Git: deja pasar un commit limpio", rc_c == 0, [f"rc={rc_c}", sal_c[:600]]))
        subprocess.run(["git", "reset", "-q"], cwd=repo)
        (pp / "fuentes" / "CV01-curriculum.md").unlink()
        prueba("precommit --archivos: valida un CSV recién editado", [H + "hooks/precommit.py", "--archivos",
               "proyectos/prueba-combinada/riesgos/riesgos.csv"], repo, ["formato AAAA-MM-DD"], codigo=1)
        prueba("precommit --archivos: archivo correcto", [H + "hooks/precommit.py", "--archivos", "ejemplos/demo-andes/riesgos/riesgos.csv"],
               repo, ["OK: 1 archivo(s)"])

        sh = shutil.which("sh")
        if not sh:  # Windows desde PowerShell: el sh que trae Git
            _git = Path(subprocess.run(["git", "--exec-path"], capture_output=True, text=True).stdout.strip())
            sh = next((str(x) for x in (_git.parents[2] / "usr" / "bin" / "sh.exe", _git.parents[2] / "bin" / "sh.exe") if x.exists()), None)
        if sh:
            _hooks = json.loads((repo / ".claude" / "settings.json").read_text(encoding="utf-8"))["hooks"]
            _env = {**os.environ, "CLAUDE_PROJECT_DIR": str(repo)}

            def hook_claude(tipo, entrada):
                cmd = _hooks[tipo][0]["hooks"][0]["command"]
                r = subprocess.run([sh, "-c", cmd], cwd=repo, input=entrada, capture_output=True, text=True,
                                   encoding="utf-8", errors="replace", env=_env, timeout=120)
                return r.returncode, (r.stdout or "") + (r.stderr or "")

            rc_h, sal_h = hook_claude("SessionStart", "{}")
            resultados.append(("settings.json: el hook de inicio corre a través de py.sh", rc_h == 0 and "[Estado del PMO Harness]" in sal_h,
                               [f"rc={rc_h}", sal_h[:400]]))
            rc_h, sal_h = hook_claude("PreToolUse", json.dumps({"tool_name": "Bash", "tool_input":
                                      {"command": "git add -f proyectos/prueba-postulacion/_entrada/nota.txt"}}))
            resultados.append(("settings.json: el bloqueo de confidenciales conserva el código 2 a través de py.sh",
                               rc_h == 2 and "BLOQUEADO" in sal_h, [f"rc={rc_h}", sal_h[:400]]))
        else:
            resultados.append(("settings.json: hooks a través de py.sh (omitida: no hay sh en este computador)", True, []))


        # --- v2.1 · Una empresa por carpeta: herencia, banco común, carpeta nueva y actualización del núcleo ---
        sys.path.insert(0, str(repo / "herramientas"))
        import _nucleo as _nuc
        _oficial = _nuc.origen_oficial(repo)
        resultados.append(("NUCLEO.txt: origen público y comparación de URL (https, ssh, .git)", _oficial.startswith("https://github.com/")
                           and _nuc.misma_url(_oficial + ".git", _oficial.upper() + "/")
                           and _nuc.misma_url("git@github.com:" + _oficial.split("github.com/")[1] + ".git", _oficial)
                           and not _nuc.misma_url(_oficial + "-otro", _oficial), [_oficial]))
        (tmp / "lf.md").write_bytes(b"linea uno\nlinea dos\n")
        (tmp / "crlf.md").write_bytes(b"linea uno\r\nlinea dos\r\n")
        (tmp / "bin.dat").write_bytes(b"\0\r\n")
        (tmp / "bin2.dat").write_bytes(b"\0\n")
        resultados.append(("_nucleo: la huella ignora CRLF/LF en texto, no en binarios (zip de GitHub vs copia en Windows)",
                           _nuc.huella(tmp / "lf.md") == _nuc.huella(tmp / "crlf.md") and _nuc.huella(tmp / "bin.dat") != _nuc.huella(tmp / "bin2.dat"),
                           ["huellas distintas por saltos de línea, o binario alterado"]))
        subprocess.run(["git", "remote", "add", "origin", _oficial + ".git"], cwd=repo)
        prueba("estado: reconoce una copia clonada del núcleo público", [H + "estado.py"], repo, ["Copia clonada del núcleo público", "PRIVADO propio"])
        prueba("nueva_empresa --aqui: configura la empresa de esta carpeta y la desconecta del núcleo",
               [H + "nueva_empresa.py", "Consultora Sur", "--aqui"], repo, ["Empresa «Consultora Sur» configurada", "empresa/", "se desconectó"])
        _remotos = subprocess.run(["git", "remote"], cwd=repo, capture_output=True, text=True).stdout.strip()
        _est_aqui = (repo / ".harness" / "estado.yaml").read_text(encoding="utf-8")
        resultados.append(("nueva_empresa --aqui: sin remoto al núcleo y con el origen oficial para actualizar",
                           _remotos == "" and _oficial in _est_aqui, [f"remotos: {_remotos!r}", _est_aqui]))
        subprocess.run(["git", "remote", "add", "origin", _oficial], cwd=repo)
        prueba("estado: advierte si una carpeta instalada apunta al núcleo público", [H + "estado.py"], repo, ["ATENCIÓN", "núcleo PÚBLICO"])
        subprocess.run(["git", "remote", "remove", "origin"], cwd=repo)
        emp = repo / "empresa"
        ey = emp / "empresa.yaml"
        ey.write_text(ey.read_text(encoding="utf-8").replace('  nombre: ""\n  cargo: ""', '  nombre: "Ana Rojas"\n  cargo: "Gerenta"', 1)
                      .replace('color_primario: "1F3864"', 'color_primario: "7A1F1F"', 1), encoding="utf-8")
        shutil.copy2(repo / "ejemplos" / "demo-andes" / "marca" / "logo.png", emp / "marca" / "logo.png")
        (emp / "reportes" / "catalogo" / "informe-avance.md").write_text(
            "---\ntipo: informe-avance\nnombre: Informe de avance al directorio\nformato: docx\nmodo: generar\nplantilla: \"\"\n"
            "cadencia: mensual\nperiodo_archivo: mes\ndestinatarios: [Directorio]\nfuentes: [seguimiento/]\n---\n\n# Propósito\nAvance mensual.\n",
            encoding="utf-8")
        with open(emp / "contenidos" / "registro.csv", "a", encoding="utf-8", newline="") as f:
            f.write("EMP-01,certificado,Certificado de vigencia de la sociedad,,2025-01-01,2026-01-31,Consultora Sur,\n")
        prueba("crear_proyecto: hereda de la empresa", [H + "crear_proyecto.py", "postu-empresa", "--arquetipos", "postulacion"], repo,
               ["Hereda de la empresa «Consultora Sur»", "informe-avance"])
        pe = repo / "proyectos" / "postu-empresa"
        ty = (pe / "proyecto.yaml").read_text(encoding="utf-8")
        resultados.append(("crear_proyecto: postulante, formato y formatos de reporte heredados de la empresa",
                           "postulante: Consultora Sur" in ty and 'color_primario: ""' in ty and "logo: marca/logo.png" in ty
                           and (pe / "reportes" / "catalogo" / "informe-avance.md").exists(), [ty[:900]]))
        md_emp = tmp / "emp.md"
        md_emp.write_text("# Informe de avance\n\nAvance del período.\n", encoding="utf-8")
        correr([H + "md_a_docx.py", str(md_emp), str(tmp / "emp.docx"), "--proyecto", "proyectos/postu-empresa"], repo)
        import zipfile as _zip
        from docx import Document as _Doc4
        _media = [n for n in _zip.ZipFile(tmp / "emp.docx").namelist() if n.startswith("word/media/")]
        _autor = _Doc4(str(tmp / "emp.docx")).core_properties.author
        resultados.append(("md_a_docx: el proyecto sale con la autoría y el logo de la empresa", _autor == "Ana Rojas" and bool(_media),
                           [f"autor={_autor!r} media={_media}"]))
        prueba("postulacion: vigencias del banco común de la empresa", [H + "postulacion.py", "proyectos/postu-empresa", "--hoy", "2026-10-19"],
               repo, ["Documento vencido al cierre: Certificado de vigencia de la sociedad"])
        prueba("estado: muestra la empresa de la carpeta", [H + "estado.py"], repo, ["Empresa: Consultora Sur"])
        prueba("nueva_empresa --aqui: no mezcla dos empresas en una carpeta", [H + "nueva_empresa.py", "Otra Empresa", "--aqui"], repo,
               ["ya es de la empresa «Consultora Sur»"], codigo=1)

        env_git = {"GIT_AUTHOR_NAME": "Prueba", "GIT_AUTHOR_EMAIL": "prueba@ejemplo.cl",
                   "GIT_COMMITTER_NAME": "Prueba", "GIT_COMMITTER_EMAIL": "prueba@ejemplo.cl"}
        prueba("nueva_empresa: carpeta y repositorio propios para otra empresa",
               [H + "nueva_empresa.py", "Empresa Dos", "--destino", str(tmp / "empresas"), "--sin-github"], repo,
               ["Carpeta de la empresa «Empresa Dos»", "Repositorio Git creado"], env=env_git)
        e2 = tmp / "empresas" / "empresa-dos"
        _git_e2 = lambda *a: subprocess.run(["git", *a], cwd=e2, capture_output=True, text=True).stdout.strip()
        _ey2 = (e2 / "empresa" / "empresa.yaml").read_text(encoding="utf-8") if (e2 / "empresa" / "empresa.yaml").exists() else ""
        resultados.append(("nueva_empresa: solo el núcleo, sin proyectos ni datos de otra empresa; primer commit con el control activo",
                           (e2 / "AGENTS.md").exists() and "Empresa Dos" in _ey2 and "Ana Rojas" not in _ey2
                           and (e2 / "proyectos").is_dir() and not any((e2 / "proyectos").iterdir())
                           and len(_git_e2("log", "--oneline").splitlines()) == 1 and _git_e2("config", "--get", "core.hooksPath") == ".githooks"
                           and (e2 / ".harness" / "nucleo-instalado.csv").exists(),
                           [_git_e2("log", "--oneline"), _git_e2("config", "--get", "core.hooksPath"), _ey2[:300]]))

        fuente = tmp / "nucleo-nuevo"
        correr([H + "nueva_empresa.py", "Fuente", "--destino", str(tmp), "--slug", "nucleo-nuevo", "--sin-github"], repo, env=env_git)
        _log = fuente / "CHANGELOG.md"
        _log.write_text("# Cambios\n\n## 9.9 — prueba\n- Novedad de prueba.\n\n" + _log.read_text(encoding="utf-8").split("\n", 1)[1], encoding="utf-8")
        (fuente / "herramientas" / "nuevo_de_prueba.py").write_text("print('hola')\n", encoding="utf-8")
        (fuente / "biblioteca" / "LEEME.md").unlink()
        correr([str(e2 / "herramientas" / "crear_proyecto.py"), "proy-e2", "--arquetipos", "estudio"], e2)
        (e2 / ".claude" / "skills" / "mi-skill").mkdir()
        (e2 / ".claude" / "skills" / "mi-skill" / "SKILL.md").write_text("---\nname: mi-skill\n---\nPropia.\n", encoding="utf-8")
        with open(e2 / "README.md", "a", encoding="utf-8") as f:
            f.write("\nnota local\n")
        prueba("actualizar_nucleo: vista previa sin cambios", [str(e2 / "herramientas" / "actualizar_nucleo.py"), "--desde", str(fuente)], e2,
               ["→ 9.9", "1 nuevos", "1 eliminados", "Modificados a mano", "README.md", ".claude/skills/mi-skill/SKILL.md",
                "Novedad de prueba", "Vista previa: nada cambió"])
        resultados.append(("actualizar_nucleo: la vista previa no modifica nada", not (e2 / "herramientas" / "nuevo_de_prueba.py").exists()
                           and (e2 / "biblioteca" / "LEEME.md").exists(), ["la vista previa cambió archivos"]))
        estado_e2 = (e2 / "proyectos" / "proy-e2" / "proyecto.yaml").read_text(encoding="utf-8")
        prueba("actualizar_nucleo: aplica solo el núcleo y hace el commit", [str(e2 / "herramientas" / "actualizar_nucleo.py"), "--desde",
               str(fuente), "--aplicar", "--sin-pruebas"], e2, ["Aplicado", "Commit: «Actualización del núcleo a 9.9»"], env=env_git)
        _resp = list((e2 / ".harness" / "respaldo-nucleo").glob("*/README.md"))
        resultados.append(("actualizar_nucleo: conserva proyectos, empresa y skills propias; respalda lo modificado a mano",
                           (e2 / "herramientas" / "nuevo_de_prueba.py").exists() and not (e2 / "biblioteca" / "LEEME.md").exists()
                           and (e2 / ".claude" / "skills" / "mi-skill" / "SKILL.md").exists()
                           and (e2 / "proyectos" / "proy-e2" / "proyecto.yaml").read_text(encoding="utf-8") == estado_e2
                           and "Empresa Dos" in (e2 / "empresa" / "empresa.yaml").read_text(encoding="utf-8")
                           and "nota local" not in (e2 / "README.md").read_text(encoding="utf-8")
                           and bool(_resp) and "nota local" in _resp[0].read_text(encoding="utf-8")
                           and 'version_harness: "9.9"' in (e2 / ".harness" / "estado.yaml").read_text(encoding="utf-8"),
                           [f"respaldo: {_resp}", (e2 / ".harness" / "estado.yaml").read_text(encoding="utf-8")[:300]]))
        prueba("estado: anuncia las skills propias de la empresa", [str(e2 / "herramientas" / "estado.py")], e2,
               ["Skills propias de esta carpeta", "mi-skill"])
        (fuente / "herramientas" / "tercero_de_prueba.py").write_text("print('3')\n", encoding="utf-8")
        prueba("actualizar_nucleo: una segunda actualización también se aplica", [str(e2 / "herramientas" / "actualizar_nucleo.py"),
               "--desde", str(fuente), "--aplicar", "--sin-pruebas"], e2, ["1 nuevos", "Aplicado"], ausentes=["mi-skill/SKILL.md\n  - "],
               env=env_git)
        resultados.append(("actualizar_nucleo: la skill propia sobrevive a actualizaciones sucesivas",
                           (e2 / ".claude" / "skills" / "mi-skill" / "SKILL.md").exists()
                           and (e2 / "herramientas" / "tercero_de_prueba.py").exists()
                           and ".claude/skills/mi-skill/" not in (e2 / ".harness" / "nucleo-instalado.csv").read_text(encoding="utf-8"),
                           ["se borró la skill propia o quedó registrada como del núcleo"]))
        _metricas_ok = (fuente / "herramientas" / "evaluar_metricas.py").read_text(encoding="utf-8")
        (fuente / "herramientas" / "evaluar_metricas.py").write_text("raise RuntimeError('formato de datos cambiado')\n", encoding="utf-8")
        prueba("actualizar_nucleo: si un proyecto real deja de funcionar, revierte", [str(e2 / "herramientas" / "actualizar_nucleo.py"),
               "--desde", str(fuente), "--aplicar", "--sin-pruebas"], e2,
               ["dejarían de funcionar", "proyectos/proy-e2 · evaluar_metricas.py", "se revirtió la actualización"], codigo=1, env=env_git)
        resultados.append(("actualizar_nucleo: tras revertir por un proyecto, el script sigue como antes",
                           "RuntimeError" not in (e2 / "herramientas" / "evaluar_metricas.py").read_text(encoding="utf-8"), ["quedó el script roto"]))
        (fuente / "herramientas" / "evaluar_metricas.py").write_text(_metricas_ok, encoding="utf-8")
        (fuente / "herramientas" / "otro_de_prueba.py").write_text("print('x')\n", encoding="utf-8")
        _bateria_antes = (e2 / "pruebas" / "ejecutar_pruebas.py").read_text(encoding="utf-8")
        (fuente / "pruebas" / "ejecutar_pruebas.py").write_text("import sys\nprint('0/1 pruebas correctas.')\nsys.exit(1)\n", encoding="utf-8")
        prueba("actualizar_nucleo: si las pruebas fallan, revierte todo", [str(e2 / "herramientas" / "actualizar_nucleo.py"), "--desde",
               str(fuente), "--aplicar"], e2, ["se revirtió la actualización"], codigo=1)
        resultados.append(("actualizar_nucleo: tras revertir, el núcleo queda como estaba",
                           not (e2 / "herramientas" / "otro_de_prueba.py").exists()
                           and (e2 / "pruebas" / "ejecutar_pruebas.py").read_text(encoding="utf-8") == _bateria_antes
                           and 'version_harness: "9.9"' in (e2 / ".harness" / "estado.yaml").read_text(encoding="utf-8"),
                           ["la reversión dejó cambios"]))
        shutil.rmtree(repo / ".harness", ignore_errors=True)
        shutil.rmtree(repo / "empresa", ignore_errors=True)

        # Carpeta del mantenedor del núcleo: no acepta empresas, proyectos ni actualizaciones por este medio
        subprocess.run(["git", "config", "--local", "pmo.mantenedor", "true"], cwd=repo)
        prueba("estado: reconoce la carpeta del núcleo del mantenedor", [H + "estado.py"], repo, ["Carpeta del NÚCLEO"],
               ausentes=["NO INSTALADO", "SIN PROYECTOS"])
        prueba("crear_proyecto: no crea proyectos en la carpeta del núcleo", [H + "crear_proyecto.py", "zz", "--arquetipos", "ejecucion"], repo,
               ["no lleva proyectos"], codigo=1)
        prueba("nueva_empresa --aqui: no instala una empresa en la carpeta del núcleo", [H + "nueva_empresa.py", "X", "--aqui"], repo,
               ["no lleva empresas"], codigo=1)
        prueba("actualizar_nucleo: la carpeta del núcleo se actualiza con Git", [H + "actualizar_nucleo.py"], repo, ["se actualiza con Git"], codigo=1)
        subprocess.run(["git", "config", "--local", "--unset", "pmo.mantenedor"], cwd=repo)

        # Réplicas: lineamientos sí; CV y estados financieros nunca (aunque parezcan informes); lo dudoso, por decidir
        correr([H + "crear_proyecto.py", "prueba-replicas", "--arquetipos", "ejecucion"], repo)
        pr = repo / "proyectos" / "prueba-replicas"
        (pr / "_entrada" / "05-actas-y-minutas").mkdir(parents=True, exist_ok=True)
        (pr / "_entrada" / "09-otros").mkdir(parents=True, exist_ok=True)
        (pr / "_entrada" / "05-actas-y-minutas" / "Lineamientos_Directorio_2026.md").write_text(
            "# Lineamientos del directorio 2026\n\nEl directorio define que el objetivo del proyecto se amplía a dos regiones.\n"
            "Los informes de avance deberán incluir el estado de esta definición.\n", encoding="utf-8")
        (pr / "_entrada" / "09-otros" / "Juan_Perez.md").write_text(
            "# Curriculum Vitae\n\nJuan Pérez, ingeniero. RUT 11.111.111-1.\n\nInforme de gestión 2024; informe de avance; informe final; "
            "reporte mensual del proyecto; informe técnico.\n", encoding="utf-8")
        (pr / "_entrada" / "09-otros" / "EEFF_2025.md").write_text(
            "# Estados financieros 2025\n\nBalance general y estado de resultados de la sociedad.\n", encoding="utf-8")
        (pr / "_entrada" / "09-otros" / "1.2026+Plan+de+Trabajo+y+presupuesto+CV.md").write_text(  # CV = Crea y Valida (CORFO)
            "# Plan de trabajo y presupuesto CV\n\nActividades por mes, montos por ítem y aportes del proyecto Crea y Valida.\n", encoding="utf-8")
        (pr / "_entrada" / "09-otros" / "Notas_varias.md").write_text(
            "# Notas varias\n\nIdeas sueltas sobre el proyecto, sin clasificar todavía, para conversar con el equipo.\n", encoding="utf-8")
        prueba("extraer_entrada: réplicas según la política (lineamientos sí; CV y estados financieros no)",
               [H + "extraer_entrada.py", "proyectos/prueba-replicas"], repo, ["Réplicas en fuentes/: 1 · solo ficha: 2 · por decidir en el alta: 2"])
        _rep_r = sorted(x.name for x in (pr / "fuentes").glob("F*.md"))
        resultados.append(("extraer_entrada: un CV que parece informe no se replica", _rep_r == [n for n in _rep_r if "Lineamientos" in n]
                           and len(_rep_r) == 1, [str(_rep_r)]))
        _inv_r = list(csv.DictReader(open(pr / "_entrada" / "_procesado" / "inventario.csv", encoding="utf-8")))
        _id_notas = next(r["id"] for r in _inv_r if r["archivo"].endswith("Notas_varias.md"))
        _cat_cv = next(r["categoria_sugerida"] for r in _inv_r if "Plan+de+Trabajo" in r["archivo"])
        resultados.append(("extraer_entrada: «CV» como sigla (Crea y Valida) sin datos personales no es un currículum", _cat_cv != "cv", [_cat_cv]))
        prueba("extraer_entrada --replicar: replica lo que el alta decide", [H + "extraer_entrada.py", "proyectos/prueba-replicas",
               "--replicar", _id_notas], repo, [f"{_id_notas} → fuentes/{_id_notas}-Notas_varias.md"])
        shutil.rmtree(pr)
        if os.name == "nt":  # un archivo con ruta de más de 260 caracteres no puede saltarse en silencio
            _hondo = tmp / ("carpeta-de-empresa-con-nombre-largo-" * 4)[:140] / "proy"
            _ent = Path("\\\\?\\" + str(_hondo.resolve())) / "_entrada" / "00-bases-y-aclaraciones"
            _ent.mkdir(parents=True)
            (_ent / (("Bases_con_nombre_muy_largo_para_superar_el_limite_de_Windows_" * 2) + ".md")).write_text(
                "# Bases\n\nLos postulantes deberán presentar la carta de compromiso firmada.\n", encoding="utf-8")
            prueba("extraer_entrada: no salta archivos con rutas de más de 260 caracteres (Windows)", [H + "extraer_entrada.py", str(_hondo)],
                   repo, ["| bases | 1 |", "Réplicas en fuentes/: 1"])

        # Instaladores: sintaxis (en Windows no se puede probar el de Mac en un Mac real)
        _bash = None
        if os.name == "nt":
            _gitx = Path(subprocess.run(["git", "--exec-path"], capture_output=True, text=True).stdout.strip())
            _bash = next((str(x) for x in (_gitx.parents[2] / "bin" / "bash.exe", _gitx.parents[2] / "usr" / "bin" / "bash.exe") if x.exists()), None)
        else:
            _bash = shutil.which("bash")
        if _bash:
            _r = subprocess.run([_bash, "-n", "herramientas/preparar_pc.sh"], cwd=repo, capture_output=True, text=True)
            _r2 = subprocess.run([_bash, "-n", "INSTALAR-MAC.command"], cwd=repo, capture_output=True, text=True)
            resultados.append(("instalador de Mac: sintaxis de bash", _r.returncode == 0 and _r2.returncode == 0, [_r.stderr, _r2.stderr]))
        _ps = shutil.which("powershell") or shutil.which("pwsh")
        if _ps:
            _cmd = ("$e=$null; [void][System.Management.Automation.Language.Parser]::ParseFile((Resolve-Path 'herramientas/preparar_pc.ps1'), "
                    "[ref]$null, [ref]$e); if ($e.Count) { $e | ForEach-Object { $_.Message }; exit 1 }")
            _r = subprocess.run([_ps, "-NoProfile", "-Command", _cmd], cwd=repo, capture_output=True, text=True)
            resultados.append(("instalador de Windows: sintaxis de PowerShell", _r.returncode == 0, [_r.stdout, _r.stderr]))


        # --- v2.2 · Barrido selectivo, OCR, planillas oficiales, formulario de plataforma y status ---
        correr([H + "crear_proyecto.py", "prueba-cat", "--arquetipos", "postulacion"], repo)
        pc = repo / "proyectos" / "prueba-cat"
        eb, ei, ea = pc / "_entrada" / "00-bases-y-aclaraciones", pc / "_entrada" / "04-reportes-y-plantillas", pc / "_entrada" / "05-actas-y-minutas"
        for d_ in (eb, ei, ea):
            d_.mkdir(parents=True, exist_ok=True)
        (eb / "Bases_concurso_2026.md").write_text("# Bases\n\nLos postulantes deberán presentar la carta de compromiso firmada.\n", encoding="utf-8")
        (eb / "10ley_21591_royalty.md").write_text("# Ley 21.591\n\nLos gobiernos regionales deberán destinar los recursos según esta ley.\n", encoding="utf-8")
        (ea / "Transcripcion charla explicativa.md").write_text("# Charla\n\nLos postulantes deberán adjuntar una carta del municipio, dijo el expositor.\n", encoding="utf-8")
        (ei / "Formato vacio FORM1.md").write_text(
            "Resumen de la Iniciativa\nNOMBRE OFICIAL DE LA INICIATIVA (*)\n80\n Indique el título oficial del proyecto.\n"
            "27/4/26, 19:32 Portal de postulación\nhttps://portal.ejemplo.cl/postulacion 1/4\n--- Página 2 ---\n"
            "RESUMEN EJECUTIVO (*)\n Describir el objeto principal; deberá incluir la duración.\n1,500\n"
            "¿BENEFICIARIOS CONSIDERAN PERSONAS ADULTOS MAYORES DE 60\nAÑOS?\n", encoding="utf-8")
        from pypdf import PdfWriter
        _pw = PdfWriter()
        for _ in range(3):
            _pw.add_blank_page(width=595, height=842)
        with open(eb / "Bases_tecnicas_escaneadas.pdf", "wb") as f:
            _pw.write(f)
        prueba("extraer_entrada: categorías por nombre (normativa, charla, formulario) y escaneo sin texto", [H + "extraer_entrada.py", "proyectos/prueba-cat"],
               repo, ["| normativa | 1 |", "| charla | 1 |", "| formulario | 1 |", "| bases | 2 |", "Escaneos sin texto suficiente: 1"])
        _inv_c = (pc / "_entrada" / "_procesado" / "inventario.md").read_text(encoding="utf-8")
        _rep_scan = list((pc / "fuentes").glob("*Bases_tecnicas_escaneadas*.md"))
        resultados.append(("extraer_entrada: el escaneo queda con instrucción de OCR y réplica marcada «ocr: pendiente»",
                           "pásalo por OCR" in _inv_c and bool(_rep_scan) and "ocr: pendiente" in _rep_scan[0].read_text(encoding="utf-8"),
                           [_inv_c[-900:]]))
        prueba("extraer_requisitos: por defecto solo bases y formularios; avisa lo no barrido", [H + "extraer_requisitos.py", "proyectos/prueba-cat"],
               repo, ["No barridos (otras categorías del inventario): charla 1, normativa 1"], ausentes=["vienen de charlas"])
        _barr = (pc / "postulacion" / "barrido-requisitos.csv").read_text(encoding="utf-8")
        resultados.append(("extraer_requisitos: candidatos de bases y formulario, sin ley ni charla",
                           "carta de compromiso" in _barr and "deberá incluir la duración" in _barr and "gobiernos regionales" not in _barr
                           and "municipio" not in _barr, [_barr[:600]]))
        prueba("extraer_requisitos: lo que viene de una charla no es vinculante", [H + "extraer_requisitos.py", "proyectos/prueba-cat",
               "--categorias", "bases,formulario,charla"], repo, ["1 candidato(s) vienen de charlas", "NO son vinculantes"])
        resultados.append(("extraer_requisitos: columna vinculante = no para la charla",
                           "no: confirmar por consulta" in (pc / "postulacion" / "barrido-requisitos.csv").read_text(encoding="utf-8"), ["sin marca"]))
        from docx import Document as _Doc5
        _d5 = _Doc5()
        _d5.add_paragraph("Bases técnicas del concurso. Los postulantes deberán acreditar experiencia regional de al menos dos años.")
        (eb / "ocr").mkdir()
        _d5.save(str(eb / "ocr" / "Bases_tecnicas_escaneadas.docx"))
        prueba("extraer_entrada: la versión OCR reemplaza al escaneo", [H + "extraer_entrada.py", "proyectos/prueba-cat"], repo,
               ["Réplicas en fuentes/"], ausentes=["Escaneos sin texto suficiente"])
        _reps = sorted(x.name for x in (pc / "fuentes").glob("*Bases_tecnicas_escaneadas*.md"))
        resultados.append(("extraer_entrada: una sola réplica, la de la versión con texto",
                           len(_reps) == 1 and "experiencia regional" in (pc / "fuentes" / _reps[0]).read_text(encoding="utf-8"), [str(_reps)]))

        # Planilla oficial con listas desplegables modernas y una forma: openpyxl las borra, llenar_xlsx.py no
        import zipfile as _zf
        from openpyxl import Workbook as _WB, load_workbook as _lw3
        fx = tmp / "planilla-oficial.xlsx"
        _wb = _WB()
        _ws = _wb.active
        _ws.title = "Datos"
        _ws["A1"], _ws["B1"], _ws["C1"] = "Concepto", "Monto", "Total"
        _ws["A2"], _ws["C2"] = "Honorarios", "=B2*2"
        _ws["D2"].number_format = "dd-mm-yyyy"
        _ws.merge_cells("E1:F1")
        _ws["E1"] = "Combinada"
        _ws2 = _wb.create_sheet("Resumen")
        _ws2["A1"] = "=Datos!C2"
        _ws2.protection.sheet = True
        _wb.create_sheet("Listas")["A1"] = "Sí"
        _wb.save(fx)
        _partes = {}
        with _zf.ZipFile(fx) as z:
            _infos = z.infolist()
            _partes = {i.filename: z.read(i.filename) for i in _infos}
        _sx = _partes["xl/worksheets/sheet1.xml"].decode("utf-8")
        if "xmlns:r=" not in _sx:
            _sx = _sx.replace("<worksheet ", '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
        _sx = _sx.replace("</worksheet>", '<drawing r:id="rIdDr1"/><extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
                          'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations count="1" '
                          'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidation type="list" allowBlank="1">'
                          "<x14:formula1><xm:f>Listas!$A$1:$A$2</xm:f></x14:formula1><xm:sqref>B2:B5</xm:sqref></x14:dataValidation>"
                          "</x14:dataValidations></ext></extLst></worksheet>")
        _partes["xl/worksheets/sheet1.xml"] = _sx.encode("utf-8")
        _partes["xl/worksheets/_rels/sheet1.xml.rels"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdDr1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
            'Target="../drawings/drawing1.xml"/></Relationships>').encode("utf-8")
        _partes["xl/drawings/drawing1.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:twoCellAnchor><xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff>'
            "<xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>8</xdr:row>"
            '<xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Instrucciones"/>'
            '<xdr:cNvSpPr txBox="1"/></xdr:nvSpPr><xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></a:xfrm><a:prstGeom prst="rect">'
            "<a:avLst/></a:prstGeom></xdr:spPr><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>Instrucciones</a:t></a:r></a:p></xdr:txBody></xdr:sp>"
            "<xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>").encode("utf-8")
        _partes["[Content_Types].xml"] = _partes["[Content_Types].xml"].decode("utf-8").replace(
            "</Types>", '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>').encode("utf-8")
        with _zf.ZipFile(fx, "w", _zf.ZIP_DEFLATED) as z:
            for nombre_p, contenido in _partes.items():
                z.writestr(nombre_p, contenido)
        import warnings as _w
        with _w.catch_warnings():
            _w.simplefilter("ignore")
            _rt = _lw3(fx)
            _rt.save(tmp / "openpyxl-guardado.xlsx")
        with _zf.ZipFile(tmp / "openpyxl-guardado.xlsx") as z:
            _pierde = "x14:dataValidation" not in z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        resultados.append(("llenar_xlsx: (motivo) guardar con openpyxl borra las listas desplegables modernas", _pierde, ["openpyxl las conservó"]))
        prueba("inspeccionar_xlsx: advierte que no se guarde con openpyxl", [H + "inspeccionar_xlsx.py", str(fx)], repo,
               ["no guardar con openpyxl", "lista(s) desplegable(s)", "forma(s)", "llenar_xlsx.py"])
        valores = tmp / "valores.yaml"
        valores.write_text("celdas:\n  - {hoja: Datos, celda: B2, valor: 1500000}\n  - {hoja: Datos, celda: A3, valor: \"Pasajes & viáticos\"}\n"
                           "  - {hoja: Datos, celda: D2, valor: 2026-01-15}\n  - {hoja: Datos, celda: B3, valor: \"55%\"}\n"
                           "  - {hoja: Datos, celda: C2, valor: 7}\n  - {hoja: Resumen, celda: B1, valor: 1}\n  - {hoja: Datos, celda: F1, valor: x}\n"
                           "  - {hoja: \"Datos \", celda: A9, valor: fila nueva}\n", encoding="utf-8")
        prueba("llenar_xlsx: escribe valores y rechaza fórmulas, hojas protegidas y celdas combinadas ocultas",
               [H + "llenar_xlsx.py", str(fx), "--valores", str(valores), "--salida", str(tmp / "planilla-llena.xlsx")], repo,
               ["Celdas a escribir: 5 · rechazadas: 3", "tiene una fórmula", "la hoja está protegida", "escribe en E1", "Verificación: OK",
                "x14:dataValidation 1", "xdr:sp 1"], codigo=1)
        with _w.catch_warnings():
            _w.simplefilter("ignore")
            _ll = _lw3(tmp / "planilla-llena.xlsx", read_only=True)
            _vals = (_ll["Datos"]["B2"].value, _ll["Datos"]["A3"].value, str(_ll["Datos"]["D2"].value)[:10], _ll["Datos"]["B3"].value,
                     _ll["Datos"]["A9"].value, _ll["Datos"]["C2"].value)
        with _zf.ZipFile(fx) as z1, _zf.ZipFile(tmp / "planilla-llena.xlsx") as z2:
            _iguales = all(z1.read(n) == z2.read(n) for n in z1.namelist() if n not in ("xl/worksheets/sheet1.xml", "xl/workbook.xml"))
        resultados.append(("llenar_xlsx: valores correctos y todo lo demás idéntico byte a byte",
                           _vals == (1500000, "Pasajes & viáticos", "2026-01-15", 0.55, "fila nueva", "=B2*2") and _iguales, [str(_vals), f"iguales={_iguales}"]))

        # Formulario de plataforma: reconstrucción (texto del portal y Word con instrucciones en azul), control y salidas
        prueba("formulario: reconstruye el formulario impreso del portal", [H + "formulario.py", "proyectos/prueba-cat", "--reconstruir",
               str(ei / "Formato vacio FORM1.md")], repo, ["campos: 3", "con límite de caracteres: 2", "obligatorios: 2"])
        import yaml as _yaml2
        _bor = _yaml2.safe_load((pc / "postulacion" / "formulario.borrador.yaml").read_text(encoding="utf-8"))
        _c = _bor["campos"]
        resultados.append(("formulario: preguntas, límites (también tras la ayuda), obligatorios y títulos en dos líneas",
                           [x["pregunta"] for x in _c] == ["NOMBRE OFICIAL DE LA INICIATIVA", "RESUMEN EJECUTIVO",
                                                            "¿BENEFICIARIOS CONSIDERAN PERSONAS ADULTOS MAYORES DE 60 AÑOS?"]
                           and [x["limite"] for x in _c] == [80, 1500, ""] and "Portal" not in str(_c), [str(_c)[:600]]))
        from docx.shared import RGBColor as _RGB
        _dw = _Doc5()
        _dw.add_heading("Proyecto", level=1)

        def _azul(t):
            _dw.add_paragraph().add_run(t).font.color.rgb = _RGB(0, 0, 0xFF)
        _azul("1. Título del Proyecto *\nNombre de identificación del proyecto.")
        _azul("0 / 250")
        _dw.add_paragraph("Mi proyecto de prueba")
        _azul("2. ¿Cuál es la tendencia tecnológica? *\nElija una.")
        _dw.add_paragraph("Inteligencia Artificial (IA)")
        _dw.add_paragraph("Robótica")
        _dw.add_paragraph("3. Región de impacto del proyecto *")
        _dw.add_paragraph("Antofagasta")
        _dw.save(str(tmp / "formulario-corfo.docx"))
        prueba("formulario: reconstruye un Word con las instrucciones en azul", [H + "formulario.py", "proyectos/prueba-cat", "--reconstruir",
               str(tmp / "formulario-corfo.docx")], repo, ["campos: 3", "con límite de caracteres: 1", "con respuesta previa: 1"])
        _c2 = _yaml2.safe_load((pc / "postulacion" / "formulario.borrador.yaml").read_text(encoding="utf-8"))["campos"]
        resultados.append(("formulario: pregunta, límite, respuesta previa, selección con opciones y pregunta sin color",
                           _c2[0]["limite"] == 250 and _c2[0]["respuesta"] == "Mi proyecto de prueba" and _c2[1]["tipo"] == "seleccion"
                           and _c2[1]["opciones"] == ["Inteligencia Artificial (IA)", "Robótica"] and _c2[2]["pregunta"].startswith("3. Región"),
                           [str(_c2)[:600]]))
        _bor["plataforma"] = "Portal de prueba"
        _bor["campos"][0]["respuesta"] = "Transferencia tecnológica para la agricultura del desierto en la región, con apoyo del Miniзterio regional"
        _bor["campos"][1]["archivo"] = "propuesta/campos/C02.md"
        (pc / "propuesta" / "campos").mkdir(parents=True, exist_ok=True)
        (pc / "propuesta" / "campos" / "C02.md").write_text("---\ncampo: C02\n---\n**Resumen** del proyecto.​ Dura 18 meses.\n", encoding="utf-8")
        (pc / "postulacion" / "formulario.yaml").write_text(_yaml2.safe_dump(_bor, allow_unicode=True, sort_keys=False), encoding="utf-8")
        prueba("formulario: control de límites, obligatorios y caracteres raros", [H + "formulario.py", "proyectos/prueba-cat"], repo,
               ["exceden su límite: 1", "Cyrillic Small Letter Ze", "Zero Width Space", "| C02 | RESUMEN EJECUTIVO * | 1500 | 37 | ok |"])
        prueba("formulario: --limpiar quita invisibles y deja los dudosos para revisar", [H + "formulario.py", "proyectos/prueba-cat", "--limpiar"],
               repo, ["Campos corregidos: 1", "caracteres raros: 1"])
        prueba("formulario: Word de trabajo y hoja de carga", [H + "formulario.py", "proyectos/prueba-cat", "--word", "--hoja-carga"], repo,
               ["Word de trabajo:", "Hoja de carga:"])
        _hc = list((pc / "propuesta" / "salida").glob("hoja-de-carga_*.html"))
        _ww = list((pc / "propuesta" / "salida").glob("formulario-trabajo_*.docx"))
        resultados.append(("formulario: la hoja de carga trae cada campo con su botón Copiar y el Word la instrucción en azul",
                           bool(_hc) and _hc[0].read_text(encoding="utf-8").count("Copiar</button>") == 3 and bool(_ww)
                           and any(r.font.color.rgb == _RGB(0, 0, 0xFF) for p in _Doc5(str(_ww[0])).paragraphs for r in p.runs if r.font.color.type),
                           [str(_hc), str(_ww)]))
        prueba("verificar_autoria: el Word de trabajo sale limpio", [H + "verificar_autoria.py", str(_ww[0]) if _ww else "x", "--proyecto",
               "proyectos/prueba-cat"], repo, ["Resultado: OK"])
        prueba("postulacion: muestra el formulario y lo exige en la compuerta de entrega", [H + "postulacion.py", "proyectos/prueba-cat"], repo,
               ["Formulario de la plataforma: listos / exceden el límite / obligatorios vacíos | 1 de 3 / 1 / 0", "formulario_campos_excedidos == 0"])

        # Salidas JSON para el status
        _js = subprocess.run([PY, H + "postulacion.py", "ejemplos/postulacion-demo", "--hoy", "2026-11-03", "--json"], cwd=repo, capture_output=True,
                             text=True, encoding="utf-8")
        _jv = subprocess.run([PY, H + "validar_proyecto.py", "ejemplos/demo-andes", "--hoy", "2026-09-25", "--json"], cwd=repo, capture_output=True,
                             text=True, encoding="utf-8")
        _je = subprocess.run([PY, H + "evaluar_metricas.py", "ejemplos/portafolio-acciones", "--hoy", "2026-08-31", "--json"], cwd=repo,
                             capture_output=True, text=True, encoding="utf-8")
        try:
            _pj, _vj, _ej = json.loads(_js.stdout), json.loads(_jv.stdout), json.loads(_je.stdout)
            _ok_json = (_pj["etapa"] == "red_team" and _pj["metricas"]["excluyentes_pendientes"] == 7 and _vj["conteos"]["riesgos_altos_sin_plan"] == 1
                        and _vj["conteos"]["acuerdos_vencidos"] == 1 and round(_vj["avance_total"]["real"]) == 27 and _ej["reglas_criticas"] >= 1)
        except (json.JSONDecodeError, KeyError, TypeError) as ex:
            _ok_json = False
            _pj = str(ex)
        resultados.append(("postulacion, validar_proyecto y evaluar_metricas: salida --json con los mismos cálculos", _ok_json, [str(_pj)[:300]]))

        # Status de la empresa: semáforos, vencimientos propios, reporte web y Excel, historial, inicio de sesión
        _vc = pc / "seguimiento" / "vencimientos.csv"
        _vc.write_text("id,fecha,compromiso,tipo,responsable,estado,comentario\nV-01,2026-10-24,Reporte SSM 1S26,informe,Responsable,pendiente,\n",
                       encoding="utf-8")
        prueba("status: semáforos, alerta de vencimiento propio y decisiones", [H + "status.py", "--hoy", "2026-10-20"], repo,
               ["# Status", "🟡", "Reporte SSM 1S26 (24-10, en 4 días)","## Lo que viene (14 días)", "Decidir si se postula (Go/No-Go)",
                "## Postulaciones"])
        prueba("status: el vencimiento pasado pone el proyecto en rojo", [H + "status.py", "--hoy", "2026-10-25", "--breve"], repo,
               ["🔴", "venció «Reporte SSM 1S26»"])
        prueba("status: reporte web, Markdown y Excel con historial", [H + "status.py", "--hoy", "2026-10-20", "--formato", "html,md,xlsx"], repo,
               ["Reporte:", "status-2026-10-20.html"])
        _st = repo / "status"
        resultados.append(("status: archivos del reporte e historial en status/",
                           all((_st / f"status-2026-10-20.{e}").exists() for e in ("html", "md", "xlsx")) and (_st / "historial" / "2026-10-20.json").exists()
                           and "Requiere tu decisión" in (_st / "status-2026-10-20.html").read_text(encoding="utf-8"), [str(list(_st.rglob("*")))]))
        prueba("verificar_autoria: el Excel del status sale limpio", [H + "verificar_autoria.py", str(_st / "status-2026-10-20.xlsx")], repo, ["Resultado: OK"])
        prueba("status: compara con el status anterior", [H + "status.py", "--hoy", "2026-10-25"], repo,
               ["## Cambios desde el 20-10-2026", "empeoró"])
        (repo / ".harness").mkdir(exist_ok=True)
        (repo / ".harness" / "estado.yaml").write_text('instalado: true\nversion_harness: "2.2"\n', encoding="utf-8")
        prueba("estado: al iniciar muestra el status breve", [H + "estado.py"], repo, ["Status de los proyectos", "Prueba"])
        shutil.rmtree(repo / ".harness")
        shutil.rmtree(_st)

        # Tablero único de varias empresas (carpeta y repositorio aparte) y agenda semanal (simulada)
        correr([H + "nueva_empresa.py", "Empresa Tres", "--destino", str(tmp / "empresas"), "--sin-github"], repo, env=env_git)
        panel = tmp / "empresas" / "Panel-General"
        prueba("status: crear el panel de varias empresas", [H + "status.py", "--crear-panel", str(panel), "--empresas", str(e2),
               str(tmp / "empresas" / "empresa-tres"), "--nombre", "Panel de prueba"], repo, ["Panel creado", "--private"])
        _py_panel = (panel / "panel.yaml").read_text(encoding="utf-8") if (panel / "panel.yaml").exists() else ""
        resultados.append(("status: el panel guarda rutas relativas y tiene su propio repositorio",
                           "../empresa-dos" in _py_panel and "../empresa-tres" in _py_panel and (panel / ".git").exists()
                           and (panel / "AGENTS.md").exists(), [_py_panel]))
        prueba("status: tablero único de varias empresas", [H + "status.py", "--panel", str(panel), "--formato", "html"], repo,
               ["# Tablero general", "| Empresa |", "Empresa Dos"])
        resultados.append(("status: el reporte del panel queda en la carpeta del panel, no en las empresas",
                           any(panel.glob("status-*.html")) and not (e2 / "status").exists(), [str(list(panel.iterdir()))]))
        prueba("status: un panel requiere al menos dos empresas", [H + "status.py", "--crear-panel", str(tmp / "p2"), "--empresas", str(e2)], repo,
               ["al menos dos carpetas"], codigo=1)
        _ag = prueba("agendar_status: simula la programación semanal", [H + "agendar_status.py", "--dia", "lunes", "--hora", "9:00", "--simular"],
                     repo, [], codigo=None)
        resultados.append(("agendar_status: comando del sistema correcto para este computador",
                           ("/SC WEEKLY /D MON /ST 09:00" in _ag) if os.name == "nt" else ("StartCalendarInterval" in _ag or "cron" in _ag), [_ag[:400]]))
        prueba("agendar_status: valida la hora", [H + "agendar_status.py", "--dia", "lunes", "--hora", "25:00", "--simular"], repo,
               ["--hora HH:MM"], codigo=1)

        # Instalado y sin proyectos: el hook de inicio debe guiar a crear o cargar uno (va al final: borra proyectos)
        (repo / ".harness").mkdir(exist_ok=True)
        (repo / ".harness" / "estado.yaml").write_text("instalado: true\n", encoding="utf-8")
        for d in (repo / "proyectos").iterdir():
            if d.is_dir():
                shutil.rmtree(d)
        prueba("hook: instalado sin proyectos ofrece crear o cargar uno", [H + "hooks/inicio_sesion.py"], repo,
               ["Instalado pero SIN PROYECTOS", "nuevo-proyecto", "alta-proyecto", "en marcha"])


        # Ficha de inicio: sin instalar, el hook la ofrece o avisa que ya tiene respuestas
        shutil.rmtree(repo / ".harness", ignore_errors=True)
        prueba("hook: sin instalar ofrece la ficha y el manual", [H + "hooks/inicio_sesion.py"], repo,
               ["NO INSTALADO", "FICHA-DE-INICIO.md", "MANUAL-DE-USUARIO.md"], ausentes=["tiene respuestas"])
        ficha = repo / "FICHA-DE-INICIO.md"
        ficha.write_text(ficha.read_text(encoding="utf-8").replace("- Nombre y apellido:", "- Nombre y apellido: María Pérez"), encoding="utf-8")
        prueba("hook: detecta la ficha completada", [H + "hooks/inicio_sesion.py"], repo, ["FICHA-DE-INICIO.md tiene respuestas"])

    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    ok = sum(1 for _, bien, _ in resultados if bien)
    print("\n# Resultado de las pruebas\n")
    for nombre, bien, fallas in resultados:
        print(f"{'OK   ' if bien else 'FALLA'} {nombre}" + ("" if bien else "\n      " + "\n      ".join(fallas)))
    print(f"\n{ok}/{len(resultados)} pruebas correctas.")
    return 0 if ok == len(resultados) else 1


if __name__ == "__main__":
    sys.exit(main())
